from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
LOGO_PATH = Path(r"C:\Users\1\Downloads\Qom_University_logo.png")
OUTPUT_PATH = PROJECT_ROOT / "docs" / "porseman_thesis_ali_ahmadi_draft.docx"
PERSIAN_FONT = "B Nazanin"
ENGLISH_FONT = "Times New Roman"


def set_rtl_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = alignment
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_after = Pt(8)

    paragraph_properties = paragraph._p.get_or_add_pPr()
    bidi = paragraph_properties.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        paragraph_properties.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_rtl_run(run, size, bold=False):
    run.font.name = PERSIAN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold

    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.append(fonts)
    fonts.set(qn("w:ascii"), PERSIAN_FONT)
    fonts.set(qn("w:hAnsi"), PERSIAN_FONT)
    fonts.set(qn("w:eastAsia"), PERSIAN_FONT)
    fonts.set(qn("w:cs"), PERSIAN_FONT)
    size_cs = run_properties.find(qn("w:szCs"))
    if size_cs is None:
        size_cs = OxmlElement("w:szCs")
        run_properties.append(size_cs)
    size_cs.set(qn("w:val"), str(round(size * 2)))

    rtl = run_properties.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        run_properties.append(rtl)
    rtl.set(qn("w:val"), "1")


def set_ltr_run(run, size=12, bold=False):
    run.font.name = ENGLISH_FONT
    run.font.size = Pt(size)
    run.font.bold = bold

    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.append(fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), ENGLISH_FONT)
    size_cs = run_properties.find(qn("w:szCs"))
    if size_cs is None:
        size_cs = OxmlElement("w:szCs")
        run_properties.append(size_cs)
    size_cs.set(qn("w:val"), str(round(size * 2)))

    rtl = run_properties.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        run_properties.append(rtl)
    rtl.set(qn("w:val"), "0")


def add_text(document, text, size=14, bold=False, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    heading_match = re.match(r"^[۰-۹0-9]+(?:-[۰-۹0-9]+)+", text)
    if re.match(r"^فصل [^:]+:", text):
        heading_level = 1
    elif heading_match:
        heading_level = min(text.split(".", 1)[0].count("-") + 1, 3)
    else:
        heading_level = None

    paragraph = document.add_paragraph(
        style=f"Heading {heading_level}" if heading_level else None
    )
    set_rtl_paragraph(paragraph, alignment)
    # Keep an entire Latin phrase, including its internal spaces, in one LTR run.
    # Splitting each word into a separate run makes Word reorder the words in RTL text.
    english_pattern = re.compile(
        r"[A-Za-z][A-Za-z0-9@._+/#:-]*(?:[ \t]+[A-Za-z0-9@._+/#:-]+)*"
    )
    cursor = 0
    for match in english_pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_rtl_run(run, size=size, bold=bold)
        run = paragraph.add_run(f"\u200e{match.group()}\u200e")
        set_ltr_run(run, size=12, bold=bold)
        cursor = match.end()
    if cursor < len(text) or not text:
        run = paragraph.add_run(text[cursor:])
        set_rtl_run(run, size=size, bold=bold)
    return paragraph


def math_run(text):
    run = OxmlElement("m:r")
    properties = OxmlElement("m:rPr")
    fonts = OxmlElement("w:rFonts")
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), "Cambria Math")
    properties.append(fonts)
    run.append(properties)
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    return run


def math_norm(symbol):
    """Create a native Office Math norm with double vertical delimiters."""
    delimiter = OxmlElement("m:d")
    properties = OxmlElement("m:dPr")
    begin = OxmlElement("m:begChr")
    begin.set(qn("m:val"), "‖")
    end = OxmlElement("m:endChr")
    end.set(qn("m:val"), "‖")
    properties.extend((begin, end))
    expression = OxmlElement("m:e")
    expression.append(math_run(symbol))
    delimiter.extend((properties, expression))
    return delimiter


def math_subscript(base, subscript):
    element = OxmlElement("m:sSub")
    base_element = OxmlElement("m:e")
    base_element.append(math_run(base))
    subscript_element = OxmlElement("m:sub")
    subscript_element.append(math_run(subscript))
    element.extend((base_element, subscript_element))
    return element


def math_sub_superscript(base, subscript, superscript):
    element = OxmlElement("m:sSubSup")
    base_element = OxmlElement("m:e")
    base_element.append(math_run(base))
    subscript_element = OxmlElement("m:sub")
    subscript_element.append(math_run(subscript))
    superscript_element = OxmlElement("m:sup")
    superscript_element.append(math_run(superscript))
    element.extend((base_element, subscript_element, superscript_element))
    return element


def math_superscript(base, superscript):
    element = OxmlElement("m:sSup")
    base_element = OxmlElement("m:e")
    base_element.append(math_run(base))
    superscript_element = OxmlElement("m:sup")
    superscript_element.append(math_run(superscript))
    element.extend((base_element, superscript_element))
    return element


def set_cell_width(cell, width):
    properties = cell._tc.get_or_add_tcPr()
    width_element = properties.find(qn("w:tcW"))
    if width_element is None:
        width_element = OxmlElement("w:tcW")
        properties.append(width_element)
    width_element.set(qn("w:w"), str(round(width.cm * 567)))
    width_element.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_cell_text(cell, text, size=14, bold=False, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph = cell.paragraphs[0]
    set_rtl_paragraph(paragraph, alignment)
    paragraph.paragraph_format.space_after = Pt(0)
    english_pattern = re.compile(
        r"[A-Za-z][A-Za-z0-9@._+/#-]*(?:[ \t]+[A-Za-z0-9@._+/#-]+)*"
    )
    cursor = 0
    for match in english_pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_rtl_run(run, size=size, bold=bold)
        run = paragraph.add_run(f"\u200e{match.group()}\u200e")
        set_ltr_run(run, size=12, bold=bold)
        cursor = match.end()
    if cursor < len(text) or not text:
        run = paragraph.add_run(text[cursor:])
        set_rtl_run(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_results_comparison_table(document):
    rows = (
        ("مدل پایه BAAI/bge-m3", "۷۴٫۳۳٪", "۹۱٫۹۴٪", "۸۱٫۹۷٪", False),
        ("مدل BGE-M3 ریزتنظیم‌شده با داده پرسمان", "۸۱٫۷۰٪", "۹۴٫۹۴٪", "۸۷٫۶۷٪", True),
        ("jinaai/jina-embeddings-v3", "۷۹٫۳۳٪", "۹۳٫۰۰٪", "۸۵٫۴۵٪", False),
        ("Snowflake/snowflake-arctic-embed-l-v2.0", "۷۹٫۰۱٪", "۹۳٫۳۲٪", "۸۵٫۳۱٪", False),
    )
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Cm(7), Cm(2.5), Cm(2.5), Cm(2.5))
    for index, width in enumerate(widths):
        table.columns[index].width = width
        set_cell_width(table.cell(0, index), width)

    headers = ("مدل", "Recall@1", "Recall@5", "MRR@10")
    for index, text in enumerate(headers):
        cell = table.cell(0, index)
        set_cell_shading(cell, "1F4E78")
        add_cell_text(cell, text, size=14, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = None
            color = run._element.get_or_add_rPr().find(qn("w:color"))
            if color is None:
                color = OxmlElement("w:color")
                run._element.get_or_add_rPr().append(color)
            color.set(qn("w:val"), "FFFFFF")

    for model, recall_1, recall_5, mrr_10, is_fine_tuned in rows:
        cells = table.add_row().cells
        values = (model, recall_1, recall_5, mrr_10)
        for index, value in enumerate(values):
            set_cell_width(cells[index], widths[index])
            if is_fine_tuned:
                set_cell_shading(cells[index], "D9EAF7")
            add_cell_text(
                cells[index],
                value,
                size=14,
                bold=is_fine_tuned,
                alignment=WD_ALIGN_PARAGRAPH.RIGHT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER,
            )


def math_sum(lower_parts, body_parts):
    element = OxmlElement("m:nary")
    properties = OxmlElement("m:naryPr")
    character = OxmlElement("m:chr")
    character.set(qn("m:val"), "∑")
    limit_location = OxmlElement("m:limLoc")
    limit_location.set(qn("m:val"), "undOvr")
    properties.extend((character, limit_location))
    lower_limit = OxmlElement("m:sub")
    for part in lower_parts:
        lower_limit.append(part)
    upper_limit = OxmlElement("m:sup")
    body = OxmlElement("m:e")
    for part in body_parts:
        body.append(part)
    element.extend((properties, lower_limit, upper_limit, body))
    return element


def add_equation(document, number, parts):
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    table_properties = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table_properties.append(layout)
    table.columns[0].width = Cm(13)
    table.columns[1].width = Cm(2)
    set_cell_width(table.cell(0, 0), Cm(13))
    set_cell_width(table.cell(0, 1), Cm(2))
    equation_paragraph = table.cell(0, 0).paragraphs[0]
    equation_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math = OxmlElement("m:oMath")
    for part in parts:
        math.append(part)
    equation_paragraph._p.append(math)
    number_paragraph = table.cell(0, 1).paragraphs[0]
    set_rtl_paragraph(number_paragraph, WD_ALIGN_PARAGRAPH.RIGHT)
    set_rtl_run(number_paragraph.add_run(f"({number})"), size=12)


def add_cosine_equation(document, number, normalized=False):
    parts = [math_run("cos(q, d) = ")]
    if normalized:
        parts.append(math_run("q · d"))
    else:
        fraction = OxmlElement("m:f")
        numerator = OxmlElement("m:num")
        numerator.append(math_run("q · d"))
        denominator = OxmlElement("m:den")
        denominator.append(math_norm("q"))
        denominator.append(math_run(" × "))
        denominator.append(math_norm("d"))
        fraction.extend((numerator, denominator))
        parts.append(fraction)
    add_equation(document, number, parts)


def add_ranking_objective_equation(document, number):
    add_equation(
        document,
        number,
        [
            math_run("s("),
            math_subscript("q", "i"),
            math_run(", "),
            math_sub_superscript("d", "i", "+"),
            math_run(") > s("),
            math_subscript("q", "i"),
            math_run(", "),
            math_superscript("d", "−"),
            math_run(")"),
        ],
    )


def add_contrastive_loss_equation(document, number):
    """Add the per-query contrastive loss as a single native Office equation."""
    tau = chr(0x03C4)
    minus = chr(0x2212)
    def positive_score():
        return [
            math_run("exp(s("),
            math_subscript("q", "i"),
            math_run(", "),
            math_sub_superscript("d", "i", "+"),
            math_run(f") / {tau})"),
        ]

    def negative_score():
        return [
            math_run("exp(s("),
            math_subscript("q", "i"),
            math_run(", "),
            math_sub_superscript("d", "ij", minus),
            math_run(f") / {tau})"),
        ]

    fraction = OxmlElement("m:f")
    numerator = OxmlElement("m:num")
    for part in positive_score():
        numerator.append(part)
    denominator = OxmlElement("m:den")
    for part in positive_score():
        denominator.append(part)
    denominator.append(math_run(" + "))
    summation = OxmlElement("m:nary")
    sum_properties = OxmlElement("m:naryPr")
    sum_character = OxmlElement("m:chr")
    sum_character.set(qn("m:val"), chr(0x2211))
    sum_properties.append(sum_character)
    lower_limit = OxmlElement("m:sub")
    lower_limit.append(math_run("j = 1"))
    upper_limit = OxmlElement("m:sup")
    upper_limit.append(math_run("7"))
    sum_body = OxmlElement("m:e")
    for part in negative_score():
        sum_body.append(part)
    summation.extend((sum_properties, lower_limit, upper_limit, sum_body))
    denominator.append(summation)
    fraction.extend((numerator, denominator))
    add_equation(
        document,
        number,
        [math_subscript("L", "i"), math_run(f" = {minus}log "), fraction],
    )


def add_cosine_similarity_equation(document, number):
    dot = chr(0x00B7)
    add_equation(
        document,
        number,
        [
            math_run("sim(q, d) = "),
            math_fraction(
                [math_run(f"q {dot} d")],
                [math_norm("q"), math_run(" "), math_norm("d")],
            ),
        ],
    )


def add_mrr_at_10_equation(document, number):
    leq = chr(0x2264)
    member = chr(0x2208)
    reciprocal_rank = math_fraction(
        [math_run(f"I[rank(q) {leq} 10]")],
        [math_run("rank(q)")],
    )
    add_equation(
        document,
        number,
        [
            math_run("MRR@10 = "),
            math_fraction([math_run("1")], [math_run("|Q|")]),
            math_subscript(chr(0x2211), f"q {member} Q"),
            reciprocal_rank,
        ],
    )


def math_fraction(numerator_parts, denominator_parts):
    fraction = OxmlElement("m:f")
    numerator = OxmlElement("m:num")
    for part in numerator_parts:
        numerator.append(part)
    denominator = OxmlElement("m:den")
    for part in denominator_parts:
        denominator.append(part)
    fraction.extend((numerator, denominator))
    return fraction


def add_recall_equation(document, number):
    add_equation(
        document,
        number,
        [
            math_run("Recall@k = "),
            math_fraction([math_run("1")], [math_run("|Q|")]),
            math_subscript("Σ", "q ∈ Q"),
            math_run(" I[rank(q) ≤ k]"),
        ],
    )


def add_mrr_equation(document, number):
    reciprocal_rank = math_fraction([math_run("1")], [math_run("rank(q)")])
    add_equation(
        document,
        number,
        [
            math_run("MRR@10 = "),
            math_fraction([math_run("1")], [math_run("|Q|")]),
            math_subscript("Σ", "q ∈ Q"),
            reciprocal_rank,
            math_run(" I[rank(q) ≤ 10]"),
        ],
    )


def add_title_page(document):
    picture_paragraph = document.add_paragraph()
    set_rtl_paragraph(picture_paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    picture_paragraph.add_run().add_picture(str(LOGO_PATH), width=Cm(3.2))

    add_text(document, "دانشگاه قم", size=18, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(
        document,
        "دانشکده فنی و مهندسی، گروه مهندسی کامپیوتر",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_text(document, "", size=8, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(
        document,
        "پایان‌نامه کارشناسی ارشد",
        size=16,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_text(document, "", size=10, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(
        document,
        "بهبود بازیابی معنایی پاسخ‌های فارسی در مجموعه پرسش‌وپاسخ پرسمان با "
        "ریزتنظیم مدل BGE-M3 و نمونه‌های منفی سخت بازبینی‌شده",
        size=18,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_text(document, "", size=14, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(document, "تهیه‌کننده: علی احمدی", size=16, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(
        document,
        "استاد راهنما: دکتر امیر جلالی",
        size=16,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_text(document, "", size=10, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(document, "تاریخ: شهریور ۱۴۰۵", alignment=WD_ALIGN_PARAGRAPH.CENTER)


def add_persian_abstract(document):
    document.add_page_break()
    add_text(document, "چکیده", size=18, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    paragraphs = (
        "بازیابی معنایی پاسخ‌های مرتبط از میان مجموعه‌های پرسش‌وپاسخ فارسی، یکی از "
        "اجزای کلیدی سامانه‌های جست‌وجوی معنایی و معماری‌های بازیابی‌افزوده (RAG) است. "
        "مدل‌های embedding چندزبانه، با وجود توانایی عمومی در نمایش معنایی متن، لزوماً "
        "با سبک نگارش، واژگان، تنوع موضوعی و ساختار پرسش‌وپاسخ corpusهای فارسی دامنه‌ای "
        "سازگار نیستند. هدف این پژوهش، بهبود بازیابی پاسخ‌های فارسی در مجموعۀ پرسمان از "
        "طریق ریزتنظیم مدل BAAI/bge-m3 بود.",
        "ابتدا دادۀ خام پرسمان پالایش شد. در این مرحله، متن‌های غیرپرسشی، پرسش‌های "
        "مقاله‌مانند، رکوردهای تکراری، پاسخ‌های خالی و پاسخ‌های بسیار بلند کنترل شدند. "
        "طول پاسخ‌ها با توکنایزر مدل BGE-M3 اندازه‌گیری شد و پاسخ‌های دارای بیش از ۱۰۲۴ "
        "token از دادۀ نهایی حذف شدند. سپس، دادۀ آموزش در قالب یادگیری تقابلی آماده شد؛ "
        "به‌طوری‌که هر نمونه شامل یک پرسش، یک پاسخ مثبت و هفت پاسخ منفی سخت بود. نامزدهای "
        "منفی ابتدا با مدل پایه بازیابی، سپس بازرتبه‌بندی و برای کنترل احتمال منفی کاذب "
        "بررسی شدند.",
        "مدل BGE-M3 با دادۀ آموزش ریزتنظیم شد. ارزیابی اولیه با مجموعۀ آزمون ثابت شامل "
        "۱٬۶۰۱ پرسش و ۱٬۶۰۱ پاسخ یکتا انجام شد. عملکرد مدل ریزتنظیم‌شده با مدل پایۀ "
        "BGE-M3 و دو مدل عمومی jinaai/jina-embeddings-v3 و "
        "Snowflake/snowflake-arctic-embed-l-v2.0 بر پایۀ معیارهای Recall@1، Recall@5 و "
        "MRR@10 مقایسه شد. نتایج اولیه نشان داد مدل ریزتنظیم‌شده به‌ترتیب مقادیر ۸۱٫۷۰، "
        "۹۴٫۹۴ و ۸۷٫۶۷ درصد را در این سه معیار به‌دست آورده است؛ در حالی که مدل پایه "
        "به‌ترتیب ۷۴٫۳۳، ۹۱٫۹۴ و ۸۱٫۹۷ درصد کسب کرده است.",
        "تحلیل کیفی نشان داد مدل ریزتنظیم‌شده در پرسش‌های محاوره‌ای، چندقیدی و اصطلاحی، "
        "پاسخ دقیق را بهتر از متن‌های صرفاً مشابه بازیابی می‌کند. بااین‌حال، بررسی نمونه‌ها "
        "ضرورت کنترل هم‌پوشانی معنایی میان داده‌های آموزش و آزمون را آشکار کرد. بنابراین، "
        "اجرای audit کامل نشتی داده و ارزیابی دوباره بر مجموعۀ آزمون پالایش‌شده، برای تثبیت "
        "نهایی نتایج پیشنهاد می‌شود.",
    )
    for paragraph in paragraphs:
        add_text(document, paragraph)
    add_text(
        document,
        "واژگان کلیدی: بازیابی معنایی، embedding فارسی، BGE-M3، ریزتنظیم، نمونۀ منفی سخت، "
        "پرسمان، RAG، بازیابی متراکم.",
        bold=True,
    )


def add_table_of_contents(document):
    paragraph = document.add_paragraph()
    set_rtl_paragraph(paragraph)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
    paragraph._p.append(field)
    return paragraph


def add_outline(document):
    document.add_page_break()
    add_text(document, "فهرست مطالب", size=18, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_table_of_contents(document)

    document.add_page_break()
    add_text(document, "فصل اول: کلیات پژوهش", size=18, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(document, "۱-۱. مقدمه", size=15, bold=True)
    add_text(
        document,
        "گسترش منابع متنی دیجیتال و افزایش نیاز کاربران به دسترسی سریع به اطلاعات، "
        "اهمیت سامانه‌های بازیابی اطلاعات را افزایش داده است. در بسیاری از سامانه‌های "
        "پرسش‌وپاسخ، کاربر پرسش خود را با زبان طبیعی مطرح می‌کند و انتظار دارد پاسخ "
        "مرتبط، دقیق و قابل‌فهم را در رتبه‌های ابتدایی نتایج دریافت کند. بااین‌حال، "
        "تطابق واژگان موجود در پرسش و پاسخ همواره برای تشخیص ارتباط معنایی کافی نیست؛ "
        "زیرا یک مفهوم می‌تواند با عبارت‌ها و واژه‌های متفاوت بیان شود.",
    )
    add_text(
        document,
        "در زبان فارسی، تنوع شکل نوشتار، وجود نیم‌فاصله، تفاوت نویسه‌های فارسی و عربی، "
        "صورت‌های گوناگون واژه‌ها و تنوع بیان، بازیابی پاسخ مرتبط را دشوارتر می‌کند. "
        "روش‌های بازیابی واژگانی معمولاً بر اشتراک واژه‌ها تکیه دارند؛ ازاین‌رو ممکن است "
        "پاسخی که از نظر مفهومی مرتبط است اما واژگان متفاوتی دارد، در رتبه‌های پایین‌تر "
        "نمایش داده شود. بازیابی معنایی با تبدیل متن‌ها به بازنمایی برداری، می‌کوشد این "
        "محدودیت را با سنجش نزدیکی مفهومی پرسش و پاسخ کاهش دهد.",
    )
    add_text(
        document,
        "بازیابی معنایی یکی از اجزای اصلی سامانه‌های تولید مبتنی بر بازیابی یا RAG است. "
        "در معماری RAG، پیش از تولید پاسخ توسط مدل زبانی بزرگ، یک retriever اسناد یا "
        "پاسخ‌های مرتبط را از corpus بازیابی می‌کند و آن‌ها را به‌عنوان زمینه در اختیار "
        "مدل مولد قرار می‌دهد. بنابراین، کیفیت مرحله بازیابی مستقیماً بر کیفیت زمینه "
        "ورودی مدل مولد اثر می‌گذارد. اگر پاسخ یا سند مرتبط در رتبه‌های ابتدایی بازیابی "
        "نشود، مدل مولد نیز به اطلاعات مناسب دسترسی نخواهد داشت و ممکن است پاسخ ناقص، "
        "نامرتبط یا غیرمستند تولید کند. ازاین‌رو، بهبود مدل embedding و افزایش رتبه "
        "پاسخ صحیح، علاوه بر کاربرد مستقیم در جست‌وجوی معنایی، می‌تواند زیرساخت یک "
        "سامانه RAG فارسی قابل‌اعتمادتر را فراهم کند.",
    )
    add_text(
        document,
        "مدل‌های embedding چندزبانه، امکان نمایش پرسش و پاسخ فارسی در یک فضای برداری "
        "مشترک را فراهم می‌کنند. با وجود توانایی مدل‌های عمومی، ویژگی‌های زبانی و محتوایی "
        "یک مجموعه تخصصی پرسش‌وپاسخ می‌تواند با داده‌های آموزشی عمومی متفاوت باشد. در چنین "
        "شرایطی، ریزتنظیم مدل بر داده‌های همان حوزه و استفاده از نمونه‌های منفی مناسب، "
        "می‌تواند مرز میان پاسخ صحیح و پاسخ‌های ظاهراً مشابه اما نامرتبط را دقیق‌تر کند.",
    )
    add_text(
        document,
        "پژوهش حاضر به بهبود بازیابی پاسخ در مجموعه پرسش‌وپاسخ پرسمان می‌پردازد. در این "
        "کار، ابتدا داده‌ها پالایش و تکرارهای پرسش و پاسخ حذف می‌شوند. سپس برای هر پرسش، "
        "نمونه‌های منفی سخت با ترکیب بازیابی اولیه، بازرتبه‌بندی و کنترل کیفیت تولید شده و "
        "مدل BGE-M3 بر این داده‌ها ریزتنظیم می‌شود. در پایان، عملکرد مدل پایه و مدل "
        "ریزشده با استفاده از معیارهای رتبه‌بندی روی یک مجموعه آزمون ثابت مقایسه خواهد شد.",
    )
    add_text(document, "۱-۲. بیان مسئله", size=15, bold=True)
    for paragraph_text in (
        "سامانه‌های پرسش‌وپاسخ فارسی در بسیاری از حوزه‌ها با حجم زیادی از پرسش‌های "
        "کاربران و پاسخ‌های متنی مواجه‌اند. در چنین سامانه‌هایی، کاربر معمولاً انتظار "
        "ندارد صرفاً فهرستی از متن‌های دارای واژه‌های مشترک با پرسش خود را دریافت کند؛ "
        "بلکه انتظار دارد پاسخ مرتبط با مقصود او، در رتبه‌های ابتدایی نمایش داده شود. "
        "برای نمونه، ممکن است یک کاربر پرسش خود را با زبان محاوره‌ای، شکل کوتاه‌شده یا "
        "واژه‌هایی متفاوت از متن پاسخ مطرح کند. در این وضعیت، اتکا به تطابق واژگانی "
        "نمی‌تواند ارتباط واقعی میان پرسش و پاسخ را به‌طور کامل تشخیص دهد.",
        "مجموعه پرسش‌وپاسخ پرسمان شامل تعداد زیادی پرسش و پاسخ فارسی است که می‌تواند "
        "برای ساخت سامانه جست‌وجوی معنایی، بازیابی پاسخ و همچنین به‌عنوان منبع دانش در "
        "معماری‌های RAG استفاده شود. بااین‌حال، استفاده مستقیم از داده خام برای آموزش "
        "مدل بازیابی با چند مسئله همراه است. بخشی از رکوردها ممکن است پرسش معتبر نداشته "
        "باشند، متن پاسخ به‌اشتباه در ستون پرسش قرار گرفته باشد، طول پرسش یا پاسخ برای "
        "آموزش مناسب نباشد، یا چند رکورد دارای پرسش یا پاسخ یکسان باشند. وجود این موارد "
        "نه‌تنها کیفیت داده آموزشی را کاهش می‌دهد، بلکه می‌تواند اعتبار نتایج ارزیابی را "
        "نیز تحت تأثیر قرار دهد.",
        "برای مثال، اگر یک پاسخ یکسان در داده آموزش و آزمون وجود داشته باشد، مدل ممکن "
        "است به‌جای یادگیری رابطه معنایی میان پرسش و پاسخ، از تکرار متن بهره ببرد. "
        "همچنین، اگر چند پرسش تکراری یا بسیار مشابه در مجموعه آزمون حضور داشته باشند، "
        "برخی نمونه‌ها وزن بیشتری در metric نهایی می‌گیرند. ازاین‌رو، پیش از آموزش مدل "
        "باید داده‌ها پالایش شوند، رکوردهای نامناسب کنار گذاشته شوند و تکرارهای پرسش و "
        "پاسخ کنترل شوند. این فرایند باید به‌صورت شفاف و بازتولیدپذیر انجام شود تا "
        "مشخص باشد هر رکورد به چه علت در داده باقی مانده یا حذف شده است.",
        "پس از آماده‌سازی داده، مسئله اصلی به انتخاب و آموزش مدل بازیابی مربوط می‌شود. "
        "مدل‌های embedding می‌توانند پرسش و پاسخ را به بردارهایی در یک فضای مشترک "
        "تبدیل کنند. در زمان جست‌وجو، بردار پرسش با بردار همه پاسخ‌های corpus مقایسه "
        "می‌شود و پاسخ‌ها بر اساس شباهت رتبه‌بندی می‌شوند. با وجود توانایی مدل‌های "
        "عمومی چندزبانه، این مدل‌ها لزوماً با نوع پرسش‌ها، شیوه نگارش کاربران و محتوای "
        "پاسخ‌های مجموعه پرسمان سازگار نشده‌اند. بنابراین، ممکن است پاسخ درست را در "
        "میان نتایج اولیه پیدا کنند، اما نتوانند آن را به رتبه اول یا رتبه‌های بسیار "
        "بالا منتقل کنند.",
        "مدل BGE-M3 به‌عنوان یک مدل embedding چندزبانه انتخاب شده است، زیرا امکان "
        "بازنمایی مستقیم متن‌های فارسی و استفاده از آن در بازیابی متراکم را فراهم "
        "می‌کند. بااین‌حال، استفاده از مدل پایه به‌تنهایی پاسخ مسئله نیست. مدل باید با "
        "نمونه‌های آموزشی مناسب ریزتنظیم شود تا بیاموزد یک پرسش مشخص به کدام پاسخ مرتبط "
        "است و چه پاسخ‌هایی، با وجود شباهت ظاهری یا مفهومی، پاسخ صحیح آن پرسش محسوب "
        "نمی‌شوند. کیفیت این نمونه‌های آموزشی، به‌ویژه نمونه‌های منفی، نقش "
        "تعیین‌کننده‌ای در کیفیت مدل نهایی دارد.",
        "نمونه منفی تصادفی معمولاً پاسخ نامرتبطی است که مدل به‌سادگی می‌تواند آن را از "
        "پاسخ مثبت تفکیک کند. چنین نمونه‌هایی برای شروع آموزش مفید هستند، اما الزاماً "
        "مدل را برای تمایز میان پاسخ صحیح و پاسخ‌های نزدیک آماده نمی‌کنند. در مقابل، "
        "نمونه منفی سخت پاسخی است که از نظر موضوع، واژگان یا شباهت معنایی به پرسش نزدیک "
        "است، اما پاسخ دقیق آن نیست. استفاده از این نمونه‌ها می‌تواند مدل را وادار کند "
        "تفاوت‌های ظریف‌تر میان پاسخ درست و پاسخ‌های نزدیک اما نامناسب را یاد بگیرد.",
        "بااین‌حال، تولید hard negative نیز مسئله‌ای مستقل است. پاسخ‌هایی که در "
        "رتبه‌های بالای retrieval قرار می‌گیرند، همیشه نادرست نیستند؛ برخی از آن‌ها "
        "ممکن است پاسخی هم‌ارز، مکمل یا قابل‌قبول برای همان پرسش باشند. وارد کردن چنین "
        "پاسخ‌هایی به‌عنوان negative باعث تولید false negative می‌شود. false negative "
        "به مدلی آموزش می‌دهد که یک پاسخ درست یا نزدیک به درست را از پرسش دور کند و در "
        "نتیجه می‌تواند کیفیت بازیابی را کاهش دهد. بنابراین، انتخاب hard negative باید "
        "با سازوکاری انجام شود که علاوه بر شباهت، درستی یا نادرستی پاسخ نسبت به پرسش را "
        "نیز بررسی کند.",
        "در این پژوهش، برای حل این مسئله، یک فرایند چندمرحله‌ای طراحی شده است. ابتدا "
        "مدل BGE-M3 پاسخ‌های مشابه هر پرسش را بازیابی می‌کند. سپس یک reranker "
        "کاندیدهای نزدیک را با مشاهده هم‌زمان پرسش و پاسخ دقیق‌تر رتبه‌بندی می‌کند. در "
        "مرحله بعد، یک مدل زبانی بزرگ کاندیدهای برتر را با توجه به پرسش و پاسخ مثبت "
        "بررسی می‌کند و مواردی را که هم‌ارز، مبهم یا نامناسب هستند از مجموعه negative "
        "کنار می‌گذارد. در نهایت، برای هر پرسش، یک پاسخ مثبت و هفت negative معتبر برای "
        "آموزش انتخاب می‌شود.",
        "اهمیت این مسئله به کاربرد آن در سامانه‌های RAG نیز مربوط است. در معماری RAG، "
        "مدل زبانی بزرگ بر اساس زمینه بازیابی‌شده پاسخ تولید می‌کند. اگر مرحله "
        "retrieval پاسخ‌ها یا اسناد نامرتبط را در رتبه‌های ابتدایی قرار دهد، مدل مولد "
        "با زمینه نامناسب مواجه می‌شود. در نتیجه، حتی یک مدل مولد توانمند نیز ممکن است "
        "پاسخ نادرست، ناقص یا بدون پشتوانه تولید کند. بنابراین، بهبود کیفیت retrieval "
        "یک مسئله مقدماتی و ضروری برای ساخت RAG فارسی قابل‌اعتماد است، نه صرفاً یک "
        "بهینه‌سازی جانبی.",
        "بر این اساس، مسئله اصلی پژوهش حاضر چنین تعریف می‌شود: آیا پالایش داده‌های "
        "پرسمان، تولید hard negativeهای بازبینی‌شده و ریزتنظیم مدل BGE-M3 می‌تواند "
        "کیفیت بازیابی پاسخ فارسی را نسبت به مدل پایه بهبود دهد؟ برای پاسخ به این "
        "پرسش، مدل پایه و مدل ریزتنظیم‌شده روی یک مجموعه آزمون ثابت و با corpus یکسان "
        "ارزیابی می‌شوند. رتبه پاسخ درست برای هر پرسش ثبت می‌شود و معیارهای Recall@1، "
        "Recall@5 و MRR@10 برای مقایسه عملکرد دو مدل محاسبه خواهند شد.",
    ):
        add_text(document, paragraph_text)

    add_text(document, "۱-۳. اهمیت و ضرورت پژوهش", size=15, bold=True)
    for paragraph_text in (
        "رشد سریع منابع متنی فارسی و گسترش استفاده از سامانه‌های دیجیتال، نیاز به "
        "ابزارهای دقیق برای دسترسی به اطلاعات را افزایش داده است. در بسیاری از "
        "سامانه‌های پرسش‌وپاسخ، حجم پاسخ‌های موجود به اندازه‌ای است که یافتن پاسخ "
        "مناسب به‌صورت دستی یا با جست‌وجوی ساده دشوار می‌شود. کاربر انتظار دارد "
        "سامانه، علاوه بر یافتن متن‌های دارای واژه مشترک، مفهوم موردنظر او را نیز درک "
        "کند و پاسخ مناسب را در رتبه‌های ابتدایی نتایج نمایش دهد. بنابراین، بهبود "
        "بازیابی معنایی از جنبه کاربردی و تجربه کاربر اهمیت دارد.",
        "اهمیت این مسئله در داده‌های فارسی دوچندان است. زبان فارسی دارای ویژگی‌های "
        "نوشتاری و زبانی متنوعی است که می‌تواند بازیابی مبتنی بر واژه را با چالش "
        "روبه‌رو کند. تفاوت میان نویسه‌های فارسی و عربی، استفاده‌های متفاوت از "
        "نیم‌فاصله، وجود شکل‌های گوناگون نوشتاری یک واژه، تنوع ساختار جمله و بیان "
        "محاوره‌ای پرسش‌ها، باعث می‌شود دو متن دارای معنای نزدیک، لزوماً واژه‌های "
        "مشترک زیادی نداشته باشند. در نتیجه، سامانه‌ای که تنها بر تطابق واژگانی تکیه "
        "کند، ممکن است پاسخ مرتبط را نادیده بگیرد یا در رتبه‌های پایین نمایش دهد.",
        "مجموعه پرسش‌وپاسخ پرسمان از نظر محتوایی یک منبع ارزشمند برای بررسی این مسئله "
        "است. این مجموعه شامل پرسش‌های واقعی کاربران و پاسخ‌های مرتبط است و می‌تواند "
        "بازتاب‌دهنده الگوهای متنوع پرسش‌گری در زبان فارسی باشد. چنین داده‌ای، علاوه "
        "بر کاربرد در جست‌وجوی پاسخ، برای ساخت سامانه‌های راهنمای هوشمند، دستیارهای "
        "گفت‌وگومحور و سامانه‌های مبتنی بر بازیابی نیز قابل استفاده است. بااین‌حال، "
        "استفاده مؤثر از این منبع نیازمند آن است که داده‌ها از نظر کیفیت، تکرار و طول "
        "متن کنترل شوند.",
        "پالایش داده در این پژوهش اهمیت اساسی دارد، زیرا کیفیت داده ورودی مستقیماً بر "
        "رفتار مدل اثر می‌گذارد. وجود پرسش‌های ناقص، متن‌های نامرتبط، پاسخ‌های تکراری "
        "یا رکوردهایی که ساختار صحیح پرسش‌وپاسخ ندارند، می‌تواند مدل را به سمت "
        "یادگیری الگوهای نادرست سوق دهد. همچنین، اگر رکوردهای تکراری در splitهای "
        "متفاوت آموزش و آزمون قرار گیرند، نتایج ارزیابی ممکن است واقع‌بینانه نباشند. "
        "بنابراین، طراحی فیلترهای روشن برای حذف موارد نامناسب و ثبت علت حذف هر رکورد، "
        "هم برای کیفیت آموزش و هم برای اعتبار علمی پژوهش ضروری است.",
        "اهمیت دیگر این پژوهش به نقش نمونه‌های منفی در یادگیری بازیابی متراکم مربوط "
        "می‌شود. مدل بازیابی تنها با مشاهده پاسخ مثبت نمی‌آموزد که کدام پاسخ‌ها باید "
        "در رتبه پایین‌تر قرار گیرند. برای آموزش مؤثر، لازم است مدل پاسخ صحیح را از "
        "پاسخ‌های نامرتبط یا نزدیک اما نادرست تفکیک کند. اگر نمونه‌های منفی بسیار ساده "
        "و کاملاً نامرتبط باشند، مدل می‌تواند بدون یادگیری تفاوت‌های معنایی ظریف، "
        "آن‌ها را از پاسخ مثبت جدا کند. در مقابل، نمونه‌های منفی سخت مدل را با "
        "مواردی روبه‌رو می‌کنند که از نظر موضوع یا واژگان به پرسش نزدیک‌اند، اما پاسخ "
        "صحیح آن نیستند.",
        "باوجوداین، hard negative mining بدون کنترل کیفیت می‌تواند خطرناک باشد. "
        "پاسخ‌های نزدیک ممکن است در برخی موارد واقعاً قابل‌قبول، هم‌ارز یا مکمل پاسخ "
        "مثبت باشند. استفاده از این پاسخ‌ها به‌عنوان negative، false negative ایجاد "
        "می‌کند. در چنین شرایطی مدل به‌اشتباه یاد می‌گیرد پاسخ‌های درست یا تا حدی درست "
        "را از پرسش دور کند. ضرورت استفاده از reranker و مدل زبانی بزرگ در این پژوهش "
        "از همین مسئله ناشی می‌شود. reranker می‌تواند کاندیدهای نزدیک را دقیق‌تر مرتب "
        "کند و LLM می‌تواند با در نظر گرفتن پرسش و پاسخ صحیح، کاندیدهای هم‌ارز یا "
        "مبهم را از نمونه‌های منفی واقعی جدا کند.",
        "این پژوهش از منظر سامانه‌های RAG نیز اهمیت دارد. در معماری RAG، مدل زبانی "
        "بزرگ به‌تنهایی منبع نهایی دانش محسوب نمی‌شود، بلکه برای تولید پاسخ از "
        "متن‌های بازیابی‌شده استفاده می‌کند. به همین دلیل، مرحله retrieval نقش "
        "تعیین‌کننده‌ای در کیفیت پاسخ نهایی دارد. اگر پاسخ صحیح یا سند مرتبط در میان "
        "نتایج اولیه بازیابی حضور نداشته باشد، مدل مولد زمینه مناسبی برای تولید پاسخ "
        "نخواهد داشت. در نتیجه، کیفیت پایین retrieval می‌تواند به پاسخ ناقص، نامرتبط "
        "یا بدون استناد منجر شود. بهبود embedding و رتبه‌بندی پاسخ‌ها می‌تواند کیفیت "
        "زمینه ورودی به مدل مولد را افزایش دهد و مسیر ساخت یک سامانه RAG فارسی "
        "دقیق‌تر را فراهم کند.",
        "اهمیت علمی این کار در ارائه یک فرایند بازتولیدپذیر نیز قرار دارد. در این "
        "پژوهش، مراحل پالایش داده، تخصیص شناسه پایدار، تفکیک داده‌ها، تولید negative، "
        "آموزش مدل و ارزیابی به‌صورت مشخص ثبت می‌شوند. این شفافیت امکان می‌دهد که "
        "نتایج مدل پایه و مدل ریزتنظیم‌شده با test ثابت و تنظیمات ارزیابی یکسان "
        "مقایسه شوند. علاوه بر این، ذخیره رتبه پاسخ صحیح برای هر پرسش باعث می‌شود "
        "تحلیل خطا صرفاً به یک مقدار کلی metric محدود نشود و بتوان پرسش‌هایی را که "
        "مدل در آن‌ها ضعف دارد بررسی کرد.",
        "از نظر کاربردی، مدل حاصل می‌تواند به‌عنوان مؤلفه بازیابی در سامانه‌های "
        "جست‌وجوی معنایی فارسی به‌کار رود. چنین مدلی می‌تواند پاسخ‌های مرتبط را برای "
        "کاربر رتبه‌بندی کند، زمینه مناسب را برای یک مدل مولد فراهم آورد، یا در کنار "
        "روش‌های واژگانی در یک سامانه hybrid استفاده شود. همچنین، خط لوله پالایش و "
        "تولید hard negative را می‌توان برای مجموعه‌های پرسش‌وپاسخ فارسی دیگر نیز با "
        "تغییرات محدود به‌کار برد. بنابراین، دستاورد پژوهش فقط یک مدل آموزش‌دیده نیست، "
        "بلکه مجموعه‌ای از روش‌های داده‌محور و قابل تکرار برای بهبود retrieval فارسی "
        "است.",
        "در مجموع، ضرورت این پژوهش از سه جنبه قابل بیان است: نخست، نیاز عملی به "
        "بازیابی دقیق پاسخ‌های فارسی؛ دوم، نیاز علمی به بررسی اثر پالایش داده و "
        "نمونه‌های منفی سخت بر مدل‌های embedding؛ و سوم، نیاز زیرساختی سامانه‌های RAG "
        "به retrieval قابل‌اعتماد. پژوهش حاضر می‌کوشد با ترکیب این سه جنبه، مدلی "
        "ارائه کند که پاسخ درست را با احتمال بیشتری در رتبه‌های ابتدایی بازیابی قرار "
        "دهد و زمینه لازم را برای توسعه سامانه‌های هوشمند فارسی فراهم کند.",
    ):
        add_text(document, paragraph_text)
    add_text(document, "۱-۴. اهداف پژوهش", size=15, bold=True)
    add_text(document, "هدف این پژوهش، بهبود کیفیت بازیابی معنایی پاسخ‌های فارسی در مجموعه پرسش‌وپاسخ پرسمان است. منظور از بهبود کیفیت بازیابی، افزایش احتمال قرار گرفتن پاسخ صحیح در رتبه‌های ابتدایی نتایج برای هر پرسش است. برای دستیابی به این هدف، مدل BGE-M3 بر داده‌های پالایش‌شده و دارای نمونه‌های منفی سخت ریزتنظیم می‌شود و عملکرد آن با مدل پایه مقایسه خواهد شد.")
    add_text(document, "۱-۴-۱. هدف اصلی", size=14, bold=True)
    add_text(document, "هدف اصلی پژوهش، طراحی و ارزیابی یک خط لوله بازتولیدپذیر برای ریزتنظیم مدل BGE-M3 به‌منظور بازیابی معنایی پاسخ‌های فارسی در مجموعه پرسمان است؛ به‌گونه‌ای که مدل ریزتنظیم‌شده بتواند پاسخ صحیح را نسبت به مدل پایه در رتبه‌های بالاتری بازیابی کند.")
    add_text(document, "۱-۴-۲. اهداف فرعی", size=14, bold=True)
    objectives = (
        ("۱. پالایش و استانداردسازی داده‌های پرسمان", "شناسایی و حذف رکوردهای نامعتبر، پرسش‌های ناقص، متن‌های نامناسب، پرسش‌ها و پاسخ‌های تکراری و پاسخ‌های نامتناسب از نظر طول، به‌منظور ساخت داده‌ای قابل‌اعتماد برای آموزش و ارزیابی."),
        ("۲. جلوگیری از نشت داده", "ساخت splitهای آموزش، اعتبارسنجی و آزمون با test ثابت، به‌گونه‌ای که داده آزمون در تولید negative، تنظیم پارامترها و آموزش مدل استفاده نشود."),
        ("۳. تولید نمونه‌های منفی سخت", "بازیابی پاسخ‌های مشابه برای هر پرسش با BGE-M3 و انتخاب کاندیدهایی که از نظر معنایی نزدیک‌اند، اما پاسخ صحیح محسوب نمی‌شوند."),
        ("۴. افزایش دقت انتخاب negative با reranker", "بازرتبه‌بندی کاندیدهای بازیابی‌شده با مشاهده هم‌زمان پرسش و پاسخ، برای انتخاب نمونه‌های دشوارتر و مرتبط‌تر برای بررسی."),
        ("۵. کاهش false negative با LLM", "کنترل کیفیت کاندیدها با توجه به پرسش و پاسخ مثبت، تا پاسخ‌های هم‌ارز یا مبهم به‌اشتباه به‌عنوان negative وارد داده آموزش نشوند."),
        ("۶. ساخت داده آموزش ساخت‌یافته", "ساخت یک گروه آموزشی برای هر پرسش شامل یک پاسخ مثبت و هفت پاسخ منفی معتبر، به‌منظور آموزش مقایسه‌ای مدل."),
        ("۷. ریزتنظیم BGE-M3 در حالت dense", "آموزش مؤلفه dense مدل به‌گونه‌ای که شباهت پرسش با پاسخ مثبت افزایش و شباهت آن با پاسخ‌های منفی کاهش یابد."),
        ("۸. مقایسه منصفانه مدل‌ها", "مقایسه مدل پایه و مدل ریزتنظیم‌شده با test، corpus، pooling از نوع cls و شیوه رتبه‌بندی یکسان."),
        ("۹. ارزیابی کمی بازیابی", "محاسبه Recall@1، Recall@5 و MRR@10 و ثبت rank پاسخ صحیح برای هر پرسش، به‌منظور سنجش دقیق تغییر عملکرد مدل."),
        ("۱۰. مستندسازی خطا و قابلیت استفاده در RAG", "ذخیره رتبه پاسخ صحیح برای فراهم‌کردن تحلیل خطا در ادامه پژوهش و تبیین استفاده از مدل ریزتنظیم‌شده به‌عنوان retriever در RAG؛ پیاده‌سازی و ارزیابی کامل RAG خارج از دامنه آزمایش‌های فعلی است."),
    )
    for title, description in objectives:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)
    add_text(document, "۱-۵. سؤال‌های پژوهش", size=15, bold=True)
    add_text(document, "با توجه به مسئله پژوهش، هدف اصلی این مطالعه بررسی امکان بهبود بازیابی معنایی پاسخ‌های فارسی در مجموعه پرسش‌وپاسخ پرسمان است. در این پژوهش، مدل پایه BGE-M3 و مدل ریزتنظیم‌شده با داده‌های پالایش‌شده و hard negativeهای بازبینی‌شده مقایسه می‌شوند. پرسش‌های پژوهش به‌گونه‌ای تنظیم شده‌اند که با داده‌ها، فرایند اجرا و معیارهای ارزیابی موجود در پروژه قابل پاسخ‌گویی باشند.")
    add_text(document, "۱-۵-۱. سؤال اصلی پژوهش", size=14, bold=True)
    add_text(document, "آیا ریزتنظیم مدل BGE-M3 بر داده‌های پالایش‌شده پرسمان و نمونه‌های منفی سخت بازبینی‌شده، کیفیت بازیابی معنایی پاسخ‌های فارسی را نسبت به مدل پایه بهبود می‌دهد؟")
    add_text(document, "برای پاسخ به این سؤال، مدل پایه BAAI/bge-m3 و مدل ریزتنظیم‌شده با مجموعه آزمون، corpus پاسخ‌ها و تنظیمات ارزیابی یکسان مقایسه می‌شوند. معیارهای Recall@1، Recall@5 و MRR@10 مبنای اصلی سنجش هستند. افزایش مقدار این معیارها در مدل ریزتنظیم‌شده نسبت به مدل پایه، نشانه بهبود عملکرد بازیابی خواهد بود.")
    add_text(document, "۱-۵-۲. سؤال‌های فرعی پژوهش", size=14, bold=True)
    questions = (
        ("سؤال فرعی اول", "آیا پالایش داده‌های پرسمان و حذف رکوردهای نامعتبر و تکراری، مجموعه‌ای مناسب‌تر برای آموزش و ارزیابی مدل بازیابی فراهم می‌کند؟ پاسخ این سؤال با گزارش تعداد داده‌های باقی‌مانده پس از هر فیلتر، کنترل یکتایی پرسش و پاسخ و تثبیت splitهای آموزش، اعتبارسنجی و آزمون ارائه خواهد شد."),
        ("سؤال فرعی دوم", "آیا بازیابی اولیه با BGE-M3 و بازرتبه‌بندی با reranker می‌تواند کاندیدهای مناسب‌تری برای hard negative mining فراهم کند؟ پاسخ این سؤال با ثبت رتبه و امتیاز retrieval و reranker برای پاسخ‌های کاندید بررسی می‌شود."),
        ("سؤال فرعی سوم", "آیا کنترل کیفیت hard negativeها با مدل زبانی بزرگ، احتمال ورود پاسخ‌های هم‌ارز یا مبهم را به داده آموزش کاهش می‌دهد؟ فقط کاندیدهای دارای برچسب negative وارد داده آموزش می‌شوند و نسبت برچسب‌های negative، equivalent و uncertain گزارش خواهد شد."),
        ("سؤال فرعی چهارم", "آیا استفاده از یک پاسخ مثبت و هفت hard negative معتبر برای هر پرسش، مدل را در تفکیک پاسخ صحیح از پاسخ‌های نزدیک اما نادرست تقویت می‌کند؟ پاسخ این سؤال از طریق مقایسه عملکرد مدل پایه و مدل ریزتنظیم‌شده روی مجموعه آزمون ثابت بررسی می‌شود."),
        ("سؤال فرعی پنجم", "ریزش‌تنظیم مدل BGE-M3 چه اثری بر رتبه پاسخ صحیح در مجموعه آزمون پرسمان دارد؟ رتبه پاسخ صحیح برای هر پرسش ثبت و با Recall@1، Recall@5 و MRR@10 تحلیل می‌شود."),
        ("سؤال فرعی ششم", "آیا بهبود retrieval پس از ریزتنظیم، مدل را به یک retriever مناسب‌تر برای استفاده در معماری RAG فارسی تبدیل می‌کند؟ پاسخ این سؤال در سطح قابلیت استفاده و بر پایه کیفیت retrieval است؛ پیاده‌سازی و ارزیابی کامل سامانه RAG در محدوده آزمایش‌های فعلی قرار ندارد."),
    )
    for title, question in questions:
        add_text(document, title, size=14, bold=True)
        add_text(document, question)
    add_text(document, "۱-۵-۳. شیوه پاسخ‌گویی به سؤال‌های پژوهش", size=14, bold=True)
    add_text(document, "پاسخ به سؤال‌های پژوهش بر ترکیبی از شواهد کمی و کیفی استوار است. شواهد کمی شامل تعداد رکوردهای باقی‌مانده پس از پالایش، تعداد negativeهای معتبر، تعداد پرسش‌های قابل استفاده برای آموزش و metricهای بازیابی مدل‌ها است. شواهد کیفی شامل نمونه‌هایی از رکوردهای حذف‌شده، نمونه‌های hard negative، برچسب‌های LLM و بررسی پرسش‌هایی است که پاسخ صحیح آن‌ها در رتبه پایین‌تری قرار گرفته است.")
    add_text(document, "مقایسه اصلی تنها میان مدل پایه و مدل ریزتنظیم‌شده انجام می‌شود؛ بنابراین، نتیجه پژوهش نشان‌دهنده اثر کل خط لوله آماده‌سازی داده و ریزتنظیم است. در صورت اجرای آزمایش‌های ablation در آینده، مانند مقایسه با negative تصادفی یا حذف مرحله LLM، می‌توان سهم مستقل هر جزء از خط لوله را نیز با دقت بیشتری اندازه‌گیری کرد.")
    add_text(document, "۱-۶. نوآوری پژوهش", size=15, bold=True)
    add_text(document, "نوآوری این پژوهش در معرفی یک مدل پایه کاملاً جدید نیست، بلکه در طراحی و ارزیابی یک خط لوله یکپارچه برای بهبود بازیابی معنایی پاسخ‌های فارسی در مجموعه پرسمان قرار دارد. مدل BGE-M3 یک مدل embedding چندزبانه است، اما استفاده مؤثر از مدل عمومی در داده تخصصی فارسی نیازمند آماده‌سازی داده، طراحی نمونه آموزشی مناسب و ارزیابی بازتولیدپذیر است. در این پژوهش، این اجزا در یک فرایند مشخص و قابل تکرار به یکدیگر متصل شده‌اند.")
    innovations = (
        ("۱-۶-۱. نوآوری در پالایش داده‌های پرسش‌وپاسخ فارسی", "فرایند پالایش داده تنها به حذف داده خالی محدود نیست، بلکه اعتبار ساختاری پرسش، شباهت پرسش به پاسخ، طول غیرعادی پرسش، نشانه‌های متن مقاله‌گونه، تکرار پرسش و تکرار پاسخ را بررسی می‌کند. هر رکورد حذف‌شده با دلیل مشخص ثبت می‌شود؛ بنابراین، فرایند پالایش قابل بازبینی و بازتولید است."),
        ("۱-۶-۲. نوآوری در استفاده از شناسه پایدار رکوردها", "برای هر رکورد باقی‌مانده، یک شناسه پایدار بر اساس شماره سطر فایل اولیه ساخته می‌شود. این شناسه در داده پالایش‌شده، splitها، کاندیدهای hard negative، برچسب‌های LLM و فایل‌های ارزیابی حفظ می‌شود و امکان پیگیری هر رکورد و تحلیل خطا را فراهم می‌کند."),
        ("۱-۶-۳. نوآوری در تولید hard negative چندمرحله‌ای", "hard negativeها تصادفی انتخاب نمی‌شوند. BGE-M3 ابتدا پاسخ‌های نزدیک را بازیابی می‌کند و reranker کاندیدها را با دقت بیشتر مرتب می‌سازد. این ساختار چندمرحله‌ای میان سرعت retrieval و دقت reranking تعادل برقرار می‌کند و negativeهای نزدیک‌تر و دشوارتری فراهم می‌سازد."),
        ("۱-۶-۴. نوآوری در کنترل false negative با مدل زبانی بزرگ", "LLM نقش تولید پاسخ جدید ندارد، بلکه نقش کنترل کیفیت داده را بر عهده می‌گیرد. هر کاندید با مشاهده پرسش و پاسخ مثبت به‌عنوان negative، equivalent یا uncertain برچسب می‌خورد و تنها پاسخ‌های negative وارد داده آموزش می‌شوند؛ بنابراین، ریسک ورود پاسخ‌های درست یا مبهم به مجموعه منفی کاهش می‌یابد."),
        ("۱-۶-۵. نوآوری در ساخت داده آموزش ساخت‌یافته", "برای هر پرسش، یک پاسخ مثبت و هفت hard negative معتبر انتخاب می‌شود. این ساختار ثابت، آموزش مقایسه‌ای مدل را ممکن می‌سازد و متن پاسخ‌ها، شناسه negativeها، رتبه retrieval و امتیاز reranker را برای کنترل کیفیت و تحلیل نگه می‌دارد."),
        ("۱-۶-۶. نوآوری در ارزیابی بازتولیدپذیر", "مدل پایه و مدل ریزتنظیم‌شده با test ثابت، corpus یکسان، pooling از نوع cls و شباهت کسینوسی یکسان ارزیابی می‌شوند. رتبه پاسخ صحیح برای هر پرسش و metricهای Recall@1، Recall@5 و MRR@10 ذخیره می‌شوند. گزارش HTML نیز مقایسه مدل‌ها و تنظیمات اجرا را مستند می‌کند."),
        ("۱-۶-۷. تبیین جایگاه مدل در معماری RAG فارسی", "هرچند مدل مولد و ارزیابی کامل RAG در محدوده فعلی پژوهش پیاده‌سازی نشده است، مدل ریزتنظیم‌شده می‌تواند به‌عنوان retriever پاسخ‌ها یا اسناد مرتبط را برای مدل مولد بازیابی کند. بهبود رتبه پاسخ صحیح، ظرفیت آن را برای فراهم‌کردن زمینه مناسب‌تر در RAG فارسی نشان می‌دهد."),
    )
    for title, description in innovations:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)
    add_text(document, "۱-۷. ساختار پایان‌نامه", size=15, bold=True)
    for paragraph_text in (
        "این پایان‌نامه در پنج فصل تنظیم شده است. فصل نخست به کلیات پژوهش اختصاص دارد و شامل مقدمه، بیان مسئله، اهمیت و ضرورت پژوهش، اهداف، سؤال‌های پژوهش، نوآوری‌ها و ساختار کلی پایان‌نامه است. هدف این فصل، معرفی مسئله بازیابی معنایی پاسخ‌های فارسی در مجموعه پرسمان و تبیین مسیر کلی پژوهش است.",
        "فصل دوم به مبانی نظری و پیشینه پژوهش می‌پردازد. در این فصل، مفاهیم بازیابی اطلاعات، بازیابی واژگانی و بازیابی متراکم، embedding متن، شباهت کسینوسی، یادگیری تقابلی، hard negative mining، reranking و معماری RAG بررسی می‌شوند. همچنین، مدل BGE-M3، مدل‌های embedding فارسی و چندزبانه، از جمله خانواده Tooka-SBERT، و پژوهش‌های مرتبط با بازیابی معنایی فارسی مرور خواهند شد.",
        "فصل سوم روش انجام پژوهش را شرح می‌دهد. ابتدا مجموعه داده پرسمان، ساختار اولیه داده و مراحل پالایش آن معرفی می‌شود. سپس فیلترهای اعمال‌شده، روش تفکیک داده‌ها، فرایند تولید hard negative با retrieval، reranker و LLM، ساخت داده آموزش و تنظیمات ریزتنظیم مدل BGE-M3 ارائه می‌شود.",
        "فصل چهارم به آزمایش‌ها، نتایج و تحلیل آن‌ها اختصاص دارد. محیط اجرا، تنظیمات آموزش و ارزیابی، و نتایج مدل پایه و مدل ریزتنظیم‌شده با معیارهای Recall@1، Recall@5 و MRR@10 گزارش می‌شوند. نمونه‌هایی از خطاهای بازیابی، پرسش‌های دشوار و محدودیت‌های نتایج نیز در این فصل تحلیل می‌شوند.",
        "فصل پنجم شامل جمع‌بندی پژوهش، پاسخ نهایی به سؤال‌های پژوهش، محدودیت‌ها و پیشنهادهایی برای کارهای آینده است. در پایان پایان‌نامه، فهرست منابع و پیوست‌ها شامل جدول فیلترهای داده، تنظیمات آموزش، نمونه داده‌ها، prompt برچسب‌گذاری LLM، تحلیل خطا و گزارش‌های ارزیابی ارائه می‌شوند.",
    ):
        add_text(document, paragraph_text)

    document.add_page_break()
    add_text(document, "فصل دوم: مبانی نظری و پیشینه پژوهش", size=18, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(document, "۲-۱. مقدمه", size=15, bold=True)
    for paragraph_text in (
        "بازیابی اطلاعات یکی از حوزه‌های بنیادی پردازش زبان طبیعی و سامانه‌های هوشمند است. هدف اصلی در این حوزه، یافتن و رتبه‌بندی متن‌هایی است که با نیاز اطلاعاتی کاربر ارتباط بیشتری دارند. در ساده‌ترین حالت، کاربر یک پرسش یا عبارت جست‌وجو وارد می‌کند و سامانه باید از میان مجموعه بزرگی از اسناد، پاسخ‌ها یا passageها، موارد مرتبط را انتخاب کند. کیفیت این انتخاب بر تجربه کاربر و همچنین بر عملکرد سامانه‌های بالادستی، مانند سامانه‌های پرسش‌وپاسخ، جست‌وجوی سازمانی و RAG، اثر مستقیم دارد.",
        "روش‌های سنتی بازیابی اطلاعات معمولاً بر تطابق واژه‌ها میان پرسش و سند تکیه دارند. در این روش‌ها، واژه‌های موجود در پرسش با واژه‌های اسناد مقایسه می‌شوند و اسنادی که واژه‌های مشترک بیشتری دارند، امتیاز بالاتری می‌گیرند. این رویکردها از نظر سرعت، سادگی و قابلیت تفسیر مزیت دارند، اما در تشخیص ارتباط معنایی محدود هستند. برای مثال، اگر پرسش از واژه «علت» استفاده کند و پاسخ با عبارت «دلیل» یا «سبب» مفهوم یکسانی را بیان کند، روش واژگانی ممکن است ارتباط میان آن‌ها را به‌طور کامل تشخیص ندهد.",
        "با پیشرفت شبکه‌های عصبی و مدل‌های زبانی مبتنی بر Transformer، بازنمایی متنی از سطح واژه به سطح جمله و سند توسعه یافت. مدل‌های embedding متن، هر پرسش یا پاسخ را به یک بردار عددی با طول ثابت تبدیل می‌کنند. در این فضا، متن‌هایی که از نظر معنایی به یکدیگر نزدیک‌تر هستند، باید بردارهای نزدیک‌تری داشته باشند. بنابراین، به‌جای تطابق مستقیم واژه‌ها، می‌توان ارتباط پرسش و پاسخ را با شباهت بردارهای آن‌ها اندازه‌گیری کرد. Sentence-BERT از نمونه‌های مهم این رویکرد است که امکان تولید embeddingهای مستقل و مقایسه سریع آن‌ها با شباهت کسینوسی را فراهم می‌کند.",
        "بازیابی متراکم یا Dense Retrieval بر همین ایده استوار است. در این روش، یک encoder پرسش و یک encoder پاسخ یا سند را به بردار تبدیل می‌کنند. سپس پاسخ‌ها بر اساس شباهت بردارشان با بردار پرسش رتبه‌بندی می‌شوند. مزیت بازیابی متراکم این است که می‌تواند متن‌های مرتبط را حتی در صورت نبود تطابق واژگانی مستقیم بازیابی کند. پژوهش‌های مربوط به Dense Passage Retrieval نشان داده‌اند که این رویکرد می‌تواند برای بازیابی passage در سامانه‌های پرسش‌وپاسخ مؤثر باشد.",
        "با وجود مزیت مدل‌های embedding عمومی و چندزبانه، عملکرد آن‌ها در یک حوزه تخصصی به کیفیت داده و شباهت میان داده آموزشی عمومی و داده هدف وابسته است. داده‌های پرسش‌وپاسخ پرسمان دارای ویژگی‌های زبانی، موضوعی و ساختاری خاص خود هستند. پرسش‌های کاربران ممکن است محاوره‌ای، کوتاه، دارای خطای نوشتاری یا از نظر ساختار با پاسخ‌ها متفاوت باشند. پاسخ‌ها نیز ممکن است طولانی، رسمی و شامل توضیحات چندبخشی باشند. بنابراین، استفاده از مدل عمومی بدون سازگارسازی با این داده‌ها ممکن است پاسخ درست را بازیابی کند، اما رتبه آن را به اندازه کافی بالا نبرد.",
        "یکی از عوامل مهم در سازگارسازی مدل، انتخاب داده آموزشی مناسب است. در یادگیری تقابلی، مدل باید بیاموزد که پرسش به پاسخ مثبت نزدیک و از پاسخ‌های منفی دور باشد. کیفیت نمونه‌های منفی در این فرایند نقش مهمی دارد. negativeهای تصادفی معمولاً بسیار آسان‌اند، زیرا ارتباط کمی با پرسش دارند. در مقابل، hard negativeها به پرسش نزدیک هستند، اما پاسخ صحیح نیستند و مدل را برای تشخیص تفاوت‌های ظریف‌تر میان پاسخ درست و پاسخ‌های مشابه آماده می‌کنند. بااین‌حال، انتخاب نادرست hard negative می‌تواند false negative ایجاد کند و به مدل آسیب بزند.",
        "مرحله بازیابی در معماری RAG نیز اهمیت ویژه‌ای دارد. در RAG، مدل مولد پاسخ خود را بر اساس زمینه بازیابی‌شده تولید می‌کند. اگر retriever نتواند اسناد یا پاسخ‌های مرتبط را در رتبه‌های نخست قرار دهد، مدل مولد با زمینه ناکافی یا نامرتبط مواجه می‌شود. ازاین‌رو، بهبود retriever را می‌توان یکی از پیش‌نیازهای ساخت سامانه RAG قابل‌اعتماد دانست. مدل embedding ریزتنظیم‌شده در این پژوهش می‌تواند در آینده به‌عنوان مؤلفه بازیابی در چنین معماری‌ای استفاده شود.",
        "در این فصل، ابتدا مفاهیم بازیابی اطلاعات، بازنمایی برداری متن، شباهت کسینوسی و بازیابی متراکم بررسی می‌شوند. سپس یادگیری تقابلی، نمونه‌های منفی سخت، reranking و نقش مدل‌های زبانی بزرگ در کنترل کیفیت داده معرفی خواهند شد. در ادامه، مدل BGE-M3، مدل‌های embedding فارسی و چندزبانه، از جمله خانواده Tooka-SBERT، و پژوهش‌های مرتبط مرور می‌شوند. هدف فصل دوم، ایجاد مبنای نظری لازم برای تحلیل روش پیشنهادی و مشخص‌کردن جایگاه پژوهش حاضر در میان کارهای پیشین است.",
    ):
        add_text(document, paragraph_text)

    add_text(document, "بازیابی اطلاعات به فرایند یافتن، انتخاب و رتبه‌بندی اسناد، پاسخ‌ها یا متن‌هایی گفته می‌شود که با نیاز اطلاعاتی کاربر ارتباط دارند. در پژوهش حاضر، پرسش کاربر نقش query و پاسخ‌های مجموعه پرسمان نقش corpus را دارند. هدف این است که پاسخ صحیح هر پرسش، نسبت به سایر پاسخ‌ها امتیاز بالاتری دریافت کند و در رتبه‌های ابتدایی قرار گیرد. کیفیت سامانه تنها به یافتن پاسخ صحیح محدود نیست؛ رتبه پاسخ صحیح نیز اهمیت دارد. ازاین‌رو، معیارهای Recall@1، Recall@5 و MRR@10 در این پژوهش به‌کار می‌روند.")
    add_text(document, "۲-۲. بازیابی اطلاعات: از روش‌های واژگانی تا بازیابی متراکم", size=15, bold=True)
    retrieval_sections = (
        ("۲-۲-۱. بازیابی واژگانی", "در روش‌های واژگانی یا sparse retrieval، پرسش و اسناد بر اساس واژه‌های موجود در آن‌ها نمایش داده می‌شوند. مدل‌هایی مانند TF-IDF و BM25 با توجه به حضور، فراوانی و اهمیت واژه‌ها امتیاز ارتباط را محاسبه می‌کنند. این روش‌ها سریع، ساده و قابل تفسیر هستند و برای نام‌های خاص، عبارت‌های دقیق و کلیدواژه‌های کمیاب مزیت دارند. بااین‌حال، اگر پرسش و پاسخ با واژه‌های متفاوت یک مفهوم مشترک را بیان کنند، روش واژگانی ممکن است ارتباط واقعی میان آن‌ها را به‌طور کامل تشخیص ندهد."),
        ("۲-۲-۲. چالش روش‌های واژگانی در زبان فارسی", "تفاوت میان نویسه‌های فارسی و عربی، شکل‌های متفاوت نوشتن کلمات، نیم‌فاصله، پسوندها، پیشوندها، گونه‌های محاوره‌ای و تفاوت‌های املایی می‌تواند بازیابی مبتنی بر واژه را دشوار کند. برای نمونه، شکل‌های «می‌رود»، «میرود» و «می رود» از نظر ظاهری متفاوت‌اند، اما در بسیاری از موقعیت‌ها یک مفهوم را منتقل می‌کنند. همچنین، پرسش‌های کاربران ممکن است کوتاه و محاوره‌ای باشند، در حالی که پاسخ‌ها رسمی و چندبخشی‌اند."),
        ("۲-۲-۳. بازنمایی برداری متن", "embedding متن، پرسش یا پاسخ را به برداری عددی با طول ثابت تبدیل می‌کند؛ متن‌های دارای معنای مشابه باید در این فضا بردارهای نزدیک‌تری داشته باشند. Sentence-BERT با ساختار siamese، تولید embeddingهای مستقل و مقایسه سریع آن‌ها با شباهت کسینوسی را امکان‌پذیر می‌کند. تعریف استاندارد شباهت کسینوسی برای بردار پرسش q و بردار پاسخ d در رابطه (۲-۱) ارائه می‌شود. در این پژوهش، بردارها L2-normalize می‌شوند؛ در نتیجه، ضرب داخلی آن‌ها معادل شباهت کسینوسی است."),
        ("۲-۲-۴. بازیابی متراکم", "در Dense Retrieval، embedding همه پاسخ‌های corpus از پیش تولید می‌شود. هنگام دریافت پرسش جدید، embedding پرسش ساخته و با embedding پاسخ‌ها مقایسه می‌شود. پاسخ‌هایی که شباهت بیشتری دارند، رتبه بالاتری می‌گیرند. این روش می‌تواند متن‌های مرتبط را حتی در صورت نبود تطابق واژگانی مستقیم بازیابی کند. پژوهش Dense Passage Retrieval نشان داده است که بازنمایی متراکم پرسش و passage برای بازیابی متن در سامانه‌های پرسش‌وپاسخ مؤثر است."),
        ("۲-۲-۵. مقایسه روش‌های واژگانی و متراکم", "روش‌های واژگانی و متراکم لزوماً جایگزین کامل یکدیگر نیستند. روش واژگانی برای تطابق دقیق و بازیابی متراکم برای ارتباط مفهومی، مترادف‌ها و تفاوت بیان مناسب‌تر است. ترکیب این دو روش hybrid retrieval نام دارد. بااین‌حال، تمرکز پژوهش حاضر بر بازیابی متراکم است، زیرا هدف آن سنجش مستقیم اثر ریزتنظیم BGE-M3 و hard negativeهای بازبینی‌شده بر embeddingهای پرسش و پاسخ فارسی است."),
        ("۲-۲-۶. جایگاه بازیابی متراکم در پژوهش حاضر", "در این پژوهش، پاسخ‌های پرسمان corpus را تشکیل می‌دهند و هر پرسش باید پاسخ صحیح خود را از میان آن‌ها بازیابی کند. مدل پایه و مدل ریزتنظیم‌شده پرسش‌ها و پاسخ‌ها را به embedding تبدیل، نرمال‌سازی و با شباهت کسینوسی رتبه‌بندی می‌کنند. مدل ریزتنظیم‌شده با پرسش‌های پرسمان، پاسخ‌های صحیح و hard negativeهای معتبر آموزش می‌بیند؛ بنابراین، انتظار می‌رود پاسخ درست را نسبت به پاسخ‌های ظاهراً مشابه اما نادرست دقیق‌تر رتبه‌بندی کند."),
    )
    for title, description in retrieval_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۲-۳. بازنمایی برداری متن، Transformer و شباهت کسینوسی", size=15, bold=True)
    add_text(document, "بازنمایی برداری متن یا text embedding فرایندی است که در آن یک متن به برداری عددی با طول ثابت تبدیل می‌شود. هدف این است که بردارها اطلاعات معنایی متن را حفظ کنند؛ به‌گونه‌ای که پرسش و پاسخ‌های مرتبط در فضای برداری به یکدیگر نزدیک و متن‌های نامرتبط از یکدیگر دور باشند. این بازنمایی پایه بازیابی متراکم، خوشه‌بندی متن، تشخیص شباهت معنایی و سامانه‌های RAG است.")
    embedding_sections = (
        ("۲-۳-۱. Transformer و بازنمایی وابسته به زمینه", "در بازنمایی‌های قدیمی، هر واژه معمولاً یک بردار ثابت داشت و معنای آن بدون توجه کافی به جمله نمایش داده می‌شد. مدل‌های Transformer با سازوکار self-attention، برای هر token بازنمایی وابسته به زمینه تولید می‌کنند. مدل برای هر token بررسی می‌کند کدام tokenهای دیگر برای درک آن مهم‌تر هستند؛ در نتیجه وابستگی میان واژه‌های دور از هم نیز قابل مدل‌سازی است. استفاده از multi-head attention به مدل اجازه می‌دهد انواع مختلفی از ارتباط‌های زبانی و معنایی را به‌صورت هم‌زمان بررسی کند."),
        ("۲-۳-۲. BERT، pooling و embedding متن", "BERT یکی از مدل‌های مهم مبتنی بر encoder Transformer است که بازنمایی دوسویه tokenها را با توجه به زمینه سمت راست و چپ تولید می‌کند. برای retrieval، خروجی tokenها باید به یک بردار واحد برای کل پرسش یا پاسخ تبدیل شود؛ این مرحله pooling نام دارد. در روش cls pooling، بازنمایی token ویژه CLS نماینده کل متن در نظر گرفته می‌شود. در پژوهش حاضر، BGE-M3 در حالت dense و با pooling از نوع cls برای ساخت embedding پرسش‌ها و پاسخ‌ها استفاده می‌شود."),
        ("۲-۳-۳. شباهت کسینوسی", "اگر بردار پرسش با q و بردار پاسخ با d نمایش داده شود، شباهت کسینوسی مطابق رابطه (۲-۱) محاسبه می‌شود. هرچه مقدار آن بزرگ‌تر باشد، پرسش و پاسخ از نظر برداری نزدیک‌تر هستند."),
    )
    for title, description in embedding_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)
        if title.endswith("شباهت کسینوسی"):
            add_cosine_equation(document, "۲-۱")
            add_text(document, "در این پژوهش، embeddingها با L2-normalization نرمال می‌شوند. پس از نرمال‌سازی، طول هر بردار برابر یک است؛ بنابراین، محاسبه شباهت کسینوسی مطابق رابطه (۲-۲) به ضرب داخلی ساده تبدیل می‌شود.")
            add_cosine_equation(document, "۲-۲", normalized=True)
    add_text(document, "۲-۳-۴. ارتباط با ریزتنظیم مدل", size=14, bold=True)
    add_text(document, "در بازیابی متراکم، پاسخ‌های corpus از پیش encode می‌شوند و هنگام دریافت پرسش جدید، embedding آن با embedding پاسخ‌ها مقایسه می‌شود. در ریزتنظیم، مدل می‌آموزد شباهت پرسش با پاسخ مثبت را افزایش و شباهت آن با hard negativeها را کاهش دهد. بنابراین، ریزتنظیم تلاشی برای بازآرایی فضای برداری مدل بر اساس روابط معنایی داده‌های پرسمان است؛ در صورت موفقیت، پاسخ صحیح باید در رتبه بالاتری نسبت به پاسخ‌های ظاهراً مشابه اما نادرست قرار گیرد.")

    add_text(document, "۲-۴. یادگیری تقابلی، نمونه‌های منفی سخت و بازرتبه‌بندی", size=15, bold=True)
    for paragraph_text in (
        "بازیابی متراکم تنها به معماری مدل وابسته نیست، بلکه کیفیت داده‌ای که مدل با آن آموزش می‌بیند نیز نقشی تعیین‌کننده دارد. هدف آموزش در بازیابی متراکم آن است که مدل یاد بگیرد پرسش و پاسخ مرتبط را در فضای برداری به یکدیگر نزدیک و پرسش و پاسخ نامرتبط را از هم دور کند. برای این منظور، معمولاً از یادگیری تقابلی استفاده می‌شود. در این رویکرد، هر نمونه آموزشی شامل یک پرسش، یک پاسخ مثبت و یک یا چند پاسخ منفی است. پاسخ مثبت، پاسخ درست یا مرتبط با پرسش است و پاسخ‌های منفی، متن‌هایی هستند که نباید برای آن پرسش در رتبه‌های بالا قرار گیرند.",
        "در پژوهش حاضر، هر رکورد آموزشی بر مبنای یک پرسش از مجموعه پرسمان ساخته می‌شود. پاسخ ثبت‌شده برای همان پرسش، نمونه مثبت را تشکیل می‌دهد. سپس پاسخ‌های منفی با فرایندی چندمرحله‌ای انتخاب می‌شوند تا صرفاً نامرتبط نباشند، بلکه از نظر معنایی نیز به پرسش نزدیک باشند. چنین نمونه‌هایی باعث می‌شوند مدل تنها تفاوت میان موضوعات کاملاً بی‌ارتباط را یاد نگیرد، بلکه بتواند میان پاسخ درست و پاسخ‌هایی که ظاهراً مناسب به نظر می‌رسند نیز تمایز ایجاد کند.",
    ):
        add_text(document, paragraph_text)

    contrastive_sections = (
        ("۲-۴-۱. مفهوم یادگیری تقابلی در بازیابی متراکم", "در یادگیری تقابلی، مدل از مقایسه هم‌زمان نمونه‌های مثبت و منفی یاد می‌گیرد. اگر بردار پرسش با q، بردار پاسخ مثبت با d+ و بردار پاسخ منفی با d− نمایش داده شود، هدف کلی آموزش آن است که شباهت پرسش با پاسخ مثبت از شباهت همان پرسش با پاسخ‌های منفی بیشتر شود. در عمل، مدل باید نه فقط یک پاسخ منفی، بلکه مجموعه‌ای از پاسخ‌های منفی را از پاسخ مثبت متمایز کند. هرچه پاسخ‌های منفی به پاسخ مثبت یا پرسش نزدیک‌تر باشند، مسئله آموزشی دشوارتر و در عین حال آموزنده‌تر می‌شود."),
        ("۲-۴-۲. نمونه مثبت و نقش آن در آموزش", "نمونه مثبت باید واقعاً پاسخ مناسب پرسش باشد. در داده‌های پرسش‌وپاسخ، وجود یک پیوند اولیه میان پرسش و پاسخ به‌تنهایی تضمین نمی‌کند که نمونه برای آموزش مناسب است. پاسخ ممکن است خالی، بسیار کوتاه، تکراری، ناقص، خارج از موضوع یا دارای طول غیرمتعارف باشد. به همین علت، پیش از ساخت داده آموزشی، داده‌های پرسمان پالایش شدند و رکوردهای نامناسب کنار گذاشته شدند. در این پژوهش، برای هر پرسش باقی‌مانده پس از فیلترها، پاسخ متناظر آن به‌عنوان پاسخ مثبت انتخاب می‌شود. همچنین شناسه یکتایی برای رکوردها در نظر گرفته شده است تا ارتباط پرسش، پاسخ مثبت، پاسخ‌های منفی، داده آموزش و داده آزمون در تمام مراحل قابل ردیابی باشد."),
        ("۲-۴-۳. نمونه منفی تصادفی و محدودیت آن", "ساده‌ترین شیوه ساخت نمونه منفی، انتخاب تصادفی پاسخ‌هایی از سایر رکوردهای مجموعه است. این شیوه هزینه محاسباتی پایینی دارد و در مراحل ابتدایی آموزش می‌تواند مفید باشد. بااین‌حال، بخش زیادی از پاسخ‌های تصادفی کاملاً نامرتبط با پرسش‌اند. مدل معمولاً تمایز میان پاسخ مثبت و چنین پاسخ‌های کاملاً نامرتبطی را به‌سرعت یاد می‌گیرد. بنابراین، ادامه آموزش با تعداد زیادی منفی تصادفی ممکن است اثر محدودی بر کیفیت رتبه‌بندی پاسخ‌های نزدیک و دشوار داشته باشد. در مسئله واقعی بازیابی نیز پاسخ‌های اشتباه مسئله‌ساز غالباً پاسخ‌های کاملاً بی‌ربط نیستند، بلکه پاسخ‌هایی هستند که واژه‌ها، موضوع یا ساختار معنایی مشابهی با پرسش دارند، اما پاسخ دقیق آن نیستند."),
        ("۲-۴-۴. نمونه منفی سخت", "نمونه منفی سخت یا Hard Negative پاسخی است که از نظر مدل یا از نظر محتوای زبانی به پرسش نزدیک است، اما پاسخ صحیح آن محسوب نمی‌شود. چنین پاسخ‌هایی ممکن است در بازیابی اولیه رتبه بالایی کسب کنند، زیرا با پرسش هم‌پوشانی واژگانی یا معنایی دارند. استفاده از نمونه منفی سخت باعث می‌شود مدل ناچار شود مرزهای دقیق‌تری میان مفاهیم نزدیک یاد بگیرد. در نتیجه، مدل تنها به کلیدواژه‌ها یا شباهت سطحی وابسته نمی‌ماند و برای تشخیص پاسخ درست باید ارتباط دقیق‌تر میان پرسش و پاسخ را بازنمایی کند. این ویژگی برای مجموعه پرسمان اهمیت ویژه‌ای دارد، زیرا بسیاری از پرسش‌ها در حوزه‌های نزدیک به هم قرار دارند و پاسخ‌ها نیز ممکن است از واژگان مشترک دینی، اخلاقی، اجتماعی یا فقهی استفاده کنند."),
        ("۲-۴-۵. مسئله پاسخ منفی کاذب", "پاسخ منفی کاذب زمانی رخ می‌دهد که متنی به‌عنوان منفی وارد داده آموزشی شود، اما از نظر محتوایی پاسخ درست، پاسخ مکمل یا پاسخی قابل قبول برای پرسش باشد. در چنین شرایطی، مدل در فرایند آموزش تشویق می‌شود پرسش را از پاسخ مرتبط دور کند؛ در حالی که این فاصله‌دادن با هدف اصلی بازیابی در تضاد است. رتبه بالا در بازیابی اولیه برای منفی‌بودن یک پاسخ کافی نیست. بازیابی اولیه فقط مجموعه‌ای از نامزدهای نزدیک را تولید می‌کند؛ سپس باید با یک مرحله دقیق‌تر بررسی شود که آیا هر نامزد واقعاً منفی مناسب است یا خیر."),
        ("۲-۴-۶. بازیابی اولیه با BGE-M3", "در روش این پژوهش، ابتدا مدل پایه BGE-M3 برای هر پرسش آموزشی، پاسخ‌های مشابه را از میان پاسخ‌های corpus بازیابی می‌کند. تعداد نامزدهای اولیه بیشتر از تعداد منفی‌های نهایی است تا فضای مناسبی برای انتخاب وجود داشته باشد. برای هر پرسش حدود ۱۰۰ پاسخ نزدیک بازیابی می‌شود. این مرحله با استفاده از embedding پرسش و embedding پاسخ‌ها انجام می‌شود و پاسخ‌ها بر اساس شباهت برداری با پرسش رتبه‌بندی می‌شوند. پاسخ صحیح همان پرسش از مجموعه نامزدها کنار گذاشته می‌شود، زیرا باید به‌عنوان نمونه مثبت باقی بماند."),
        ("۲-۴-۷. نقش بازرتبه‌بندی در انتخاب نامزدهای منفی", "بازرتبه‌بندی یا Reranking مرحله‌ای است که در آن نامزدهای بازیابی‌شده با مدلی دقیق‌تر ارزیابی و مرتب می‌شوند. در بازیابی متراکم، برای افزایش سرعت، پرسش و پاسخ معمولاً جداگانه به embedding تبدیل می‌شوند و شباهت میان بردارها محاسبه می‌شود. این روش برای جست‌وجو در corpus بزرگ مناسب است، اما ممکن است جزئیات رابطه میان یک پرسش مشخص و یک پاسخ مشخص را به‌طور کامل بررسی نکند. مدل reranker پرسش و پاسخ را به‌صورت یک زوج ورودی دریافت می‌کند و ارتباط آن‌ها را با دقت بیشتری می‌سنجد. در این پژوهش، پس از بازیابی حدود ۱۰۰ نامزد، از reranker برای رتبه‌بندی دقیق‌تر آن‌ها استفاده می‌شود و تعدادی از پاسخ‌های برتر، به‌جز پاسخ صحیح، برای بررسی بیشتر انتخاب می‌شوند."),
        ("۲-۴-۸. کنترل کیفیت با مدل زبانی بزرگ", "در مرحله بعد، پاسخ‌های نامزد با توجه به پرسش و پاسخ مثبت ارزیابی می‌شوند. نقش مدل زبانی بزرگ در این مرحله، تولید پاسخ یا جایگزینی پاسخ مثبت نیست؛ بلکه کنترل کیفیت داده آموزشی است. مدل بررسی می‌کند که آیا پاسخ نامزد واقعاً نامرتبط یا نادرست است، یا آن‌که پاسخ صحیح، هم‌ارز، مکمل یا بسیار نزدیک به پاسخ مثبت محسوب می‌شود. اگر نامزد پاسخ درست نباشد، اما از نظر معنایی نزدیک و گمراه‌کننده باشد، می‌تواند hard negative مناسبی باشد. استفاده از مدل زبانی بزرگ در این نقش نظارتی، ریسک ورود false negative را کاهش می‌دهد و کیفیت مجموعه آموزشی را افزایش می‌دهد."),
        ("۲-۴-۹. انتخاب هفت نمونه منفی برای هر پرسش", "پس از بازیابی اولیه، بازرتبه‌بندی و کنترل کیفیت، برای هر پرسش یک پاسخ مثبت و هفت پاسخ منفی انتخاب می‌شود. انتخاب هفت منفی برای هر پرسش، به مدل اجازه می‌دهد پاسخ مثبت را نه فقط با یک گزینه اشتباه، بلکه با مجموعه‌ای از پاسخ‌های نزدیک مقایسه کند. این ساختار با تنظیم داده آموزش مدل در پروژه هم‌راستا است و هر گروه آموزشی را به شکل یک پرسش، یک پاسخ مثبت و هفت پاسخ منفی سازمان می‌دهد. این تعداد یک انتخاب طراحی است و نه یک قاعده ثابت برای همه مسائل؛ توازنی عملی میان دشواری نمونه‌ها، اندازه داده و ظرفیت محاسباتی فراهم می‌کند."),
        ("۲-۴-۱۰. ارتباط این فرایند با مسئله پژوهش", "فرایند ساخت داده آموزشی در این پژوهش، صرفاً آماده‌سازی فنی داده نیست، بلکه بخش اصلی روش پیشنهادی را تشکیل می‌دهد. مدل پایه BGE-M3 برای داده‌های چندزبانه و وظایف گوناگون طراحی شده است؛ اما برای بازیابی پاسخ در مجموعه پرسمان، لازم است مدل تفاوت‌های موضوعی، زبانی و ساختاری پرسش‌ها و پاسخ‌های فارسی این مجموعه را دقیق‌تر یاد بگیرد. پاسخ مثبت، جهت مطلوب فضای برداری را مشخص می‌کند و hard negativeهای بازبینی‌شده مرزهایی را تعیین می‌کنند که مدل باید از آن‌ها عبور نکند. اثر این فرایند در فصل نتایج با معیارهای Recall@1، Recall@5 و MRR@10 و با مقایسه مدل پایه و مدل ریزتنظیم‌شده ارزیابی خواهد شد."),
    )
    for title, description in contrastive_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)
        if title.startswith("۲-۴-۱."):
            add_text(document, "رابطه (۲-۳) هدف رتبه‌سازی در یادگیری تقابلی را نشان می‌دهد: شباهت پرسش با پاسخ صحیح باید از شباهت همان پرسش با هر پاسخ منفی بیشتر باشد. در این رابطه، زیرنویس i به رکورد آموزشی و علامت‌های مثبت و منفی به‌ترتیب به پاسخ صحیح و پاسخ منفی اشاره دارند.")
            add_ranking_objective_equation(document, "۲-۳")

    add_text(document, "۲-۵. بازیابی تقویت‌شده با تولید و جایگاه بازیابی معنایی در RAG", size=15, bold=True)
    for paragraph_text in (
        "بازیابی تقویت‌شده با تولید یا Retrieval-Augmented Generation که به اختصار RAG نامیده می‌شود، معماری‌ای برای پاسخ‌گویی مبتنی بر مدل‌های زبانی بزرگ است که در آن مدل مولد پیش از تولید پاسخ، اطلاعات مرتبط را از یک منبع بیرونی بازیابی می‌کند. ایده اصلی RAG این است که مدل زبانی نباید تنها بر دانش نهفته در پارامترهای خود تکیه کند؛ بلکه باید بتواند برای هر پرسش، اسناد، پاسخ‌ها یا قطعه‌های متنی مرتبط را پیدا کند و پاسخ نهایی را بر پایه آن‌ها تولید یا تنظیم کند.",
        "رشد استفاده از مدل‌های زبانی بزرگ، نیاز به سامانه‌هایی را افزایش داده است که پاسخ آن‌ها قابل استناد، مرتبط با داده سازمان یا منبع مشخص و قابل به‌روزرسانی باشد. در بسیاری از کاربردها، اطلاعات موردنیاز کاربر در داده‌های داخلی سازمان، پایگاه دانش، اسناد تخصصی، پرسش‌وپاسخ‌های پیشین یا منابعی وجود دارد که بخشی از داده آموزش مدل زبانی نبوده‌اند. RAG با افزودن مرحله بازیابی، امکان استفاده از این منابع را بدون آموزش مجدد کامل مدل مولد فراهم می‌کند.",
    ):
        add_text(document, paragraph_text)

    rag_sections = (
        ("۲-۵-۱. اجزای اصلی معماری RAG", "یک سامانه RAG معمولاً از سه جزء اصلی تشکیل می‌شود: corpus یا منبع دانش، بازیاب و مدل مولد. corpus مجموعه‌ای از اسناد، پاسخ‌ها، قطعه‌های متن یا رکوردهای اطلاعاتی است که سامانه باید از میان آن‌ها اطلاعات مرتبط را پیدا کند. بازیاب وظیفه دارد با دریافت پرسش کاربر، موارد مرتبط را از corpus انتخاب و رتبه‌بندی کند. مدل مولد نیز پرسش و اطلاعات بازیابی‌شده را دریافت می‌کند و پاسخ نهایی را تولید می‌نماید. در برخی معماری‌ها، اجزای دیگری مانند بازرتبه‌بند، ماژول بازنویسی پرسش، کنترل‌کننده کیفیت، فیلتر ایمنی، مدیریت حافظه مکالمه و سامانه استناددهی نیز وجود دارد."),
        ("۲-۵-۲. جریان پردازش در RAG", "فرایند RAG با دریافت پرسش کاربر آغاز می‌شود. ابتدا ممکن است پرسش نرمال‌سازی، اصلاح یا بازنویسی شود تا برای بازیابی مناسب‌تر باشد. سپس بازیاب، پرسش را به یک نمایش مناسب، مانند embedding، تبدیل می‌کند و نزدیک‌ترین رکوردها را از corpus می‌یابد. پاسخ‌های برتر ممکن است با reranker بازرتبه‌بندی شوند تا نامزدهای دقیق‌تر در اختیار مدل مولد قرار گیرند. در ادامه، متن‌های منتخب همراه با پرسش کاربر در قالب یک prompt به مدل زبانی بزرگ داده می‌شوند. مدل مولد باید با استفاده از شواهد بازیابی‌شده، پاسخ نهایی را تولید کند و در صورت ناکافی‌بودن شواهد، محدودیت پاسخ را اعلام نماید."),
        ("۲-۵-۳. بازیاب به‌عنوان گلوگاه کیفیت RAG", "مدل مولد تنها به اطلاعاتی دسترسی دارد که در prompt آن قرار می‌گیرد. اگر بازیاب پاسخ یا سند مرتبط را در میان نتایج ابتدایی قرار ندهد، مدل مولد نیز به شواهد لازم دسترسی نخواهد داشت. بنابراین، حتی یک مدل زبانی توانمند ممکن است در صورت بازیابی ضعیف، پاسخی نادرست، ناقص یا نامرتبط تولید کند. از این دیدگاه، بازیاب گلوگاه کیفیت در RAG است. معیارهایی مانند Recall@k مستقیماً به این مسئله مربوط‌اند؛ برای مثال، Recall@5 نشان می‌دهد چه نسبتی از پرسش‌ها پاسخ صحیح خود را در میان پنج نتیجه اول دارند."),
        ("۲-۵-۴. نقش embedding در بازیابی RAG", "در RAG مبتنی بر بازیابی متراکم، پرسش کاربر و متن‌های corpus به بردارهای عددی تبدیل می‌شوند. embedding پرسش با embedding همه پاسخ‌ها یا passageهای موجود مقایسه می‌شود و مواردی که شباهت بیشتری دارند، بازیابی می‌شوند. این روش به مدل اجازه می‌دهد میان عبارت‌های دارای معنای نزدیک، حتی در نبود تطابق دقیق واژگانی، ارتباط برقرار کند. برای زبان فارسی، این ویژگی اهمیت زیادی دارد؛ زیرا کاربر ممکن است پرسش را به‌صورت محاوره‌ای، کوتاه یا با املای متفاوت مطرح کند، در حالی که پاسخ موجود در corpus رسمی، طولانی یا دارای ساختار نوشتاری متفاوت است."),
        ("۲-۵-۵. آماده‌سازی corpus برای RAG", "کیفیت RAG فقط به مدل بازیاب وابسته نیست؛ کیفیت corpus نیز اهمیت اساسی دارد. اسناد تکراری، متن‌های ناقص، پاسخ‌های خالی، رکوردهای نامرتبط و متن‌های بیش‌ازحد بلند می‌توانند بازیابی را دشوار و فضای ذخیره‌سازی را ناکارآمد کنند. به همین علت، پالایش داده‌ها باید پیش از تولید embedding و ساخت شاخص بازیابی انجام شود. در corpusهای مبتنی بر اسناد بلند، معمولاً متن‌ها به قطعه‌های کوچک‌تر تقسیم می‌شوند. این فرایند chunking نام دارد. در مجموعه پرسمان، واحد اصلی پژوهش پاسخ پرسش است؛ بنابراین، پاسخ پالایش‌شده نقش passage یا واحد بازیابی را دارد."),
        ("۲-۵-۶. تعداد نتایج بازیابی‌شده و پنجره زمینه", "بازیاب معمولاً تنها یک نتیجه را به مدل مولد نمی‌دهد، بلکه تعداد مشخصی از نتایج برتر را انتخاب می‌کند. انتخاب مقدار k یک تصمیم طراحی است. اگر تعداد نتایج بسیار کم باشد، ممکن است شواهد کافی وارد prompt نشود. اگر تعداد نتایج بسیار زیاد باشد، prompt طولانی می‌شود و علاوه بر افزایش هزینه، احتمال ورود متن‌های نامرتبط نیز بالا می‌رود. در سامانه‌ای که تنها پنج پاسخ یا passage برتر را به مدل مولد می‌دهد، قرارگرفتن پاسخ صحیح در پنج رتبه اول اهمیت مستقیم دارد. این موضوع دلیل استفاده از Recall@5 در ارزیابی پژوهش حاضر است."),
        ("۲-۵-۷. بازرتبه‌بندی و کنترل ارتباط", "بازیابی اولیه معمولاً با سرعت بالا و در فضای بزرگی از پاسخ‌ها انجام می‌شود. پس از آن، reranker می‌تواند تعداد محدودی از نتایج برتر را با دقت بیشتری بررسی کند. در معماری RAG، این مرحله موجب می‌شود متن‌هایی که ارتباط دقیق‌تری با پرسش دارند در ابتدای زمینه ورودی مدل مولد قرار گیرند. بازرتبه‌بندی به‌ویژه زمانی مفید است که پاسخ‌های موجود از نظر موضوعی نزدیک‌اند، اما فقط یکی از آن‌ها پاسخ دقیق پرسش است."),
        ("۲-۵-۸. مسئله توهم و اتکاپذیری پاسخ", "مدل‌های زبانی بزرگ ممکن است متنی روان و ظاهراً معتبر تولید کنند، اما بخشی از آن با شواهد موجود سازگار نباشد. این پدیده معمولاً توهم یا Hallucination نامیده می‌شود. RAG می‌تواند احتمال چنین وضعیتی را کاهش دهد، زیرا مدل مولد را به استفاده از زمینه بازیابی‌شده هدایت می‌کند؛ اما RAG به‌تنهایی تضمین‌کننده صحت پاسخ نیست. برای افزایش اتکاپذیری، لازم است در طراحی prompt از مدل خواسته شود فقط بر مبنای شواهد ارائه‌شده پاسخ دهد، موارد بدون شواهد را قطعی بیان نکند و در صورت ناکافی‌بودن اطلاعات، این موضوع را اعلام کند."),
        ("۲-۵-۹. جایگاه پژوهش حاضر در معماری RAG", "پژوهش حاضر یک سامانه RAG کامل، شامل مدل مولد، طراحی prompt، کنترل ادعاها و رابط کاربری، پیاده‌سازی نمی‌کند. تمرکز آن بر بخش بازیابی است: ساخت corpus پالایش‌شده از پاسخ‌های پرسمان، تولید embedding برای پرسش و پاسخ، بازیابی پاسخ صحیح و ریزتنظیم BGE-M3 با نمونه‌های منفی سخت بازبینی‌شده. با وجود این، خروجی پژوهش برای توسعه RAG فارسی کاربرد مستقیم دارد. مدل ریزتنظیم‌شده می‌تواند به‌عنوان retriever در مرحله نخست یک سامانه RAG استفاده شود و پاسخ‌های برتر بازیابی‌شده می‌توانند به مدل مولد داده شوند تا پاسخ نهایی بر اساس آن‌ها تهیه شود."),
        ("۲-۵-۱۰. جمع‌بندی", "RAG معماری‌ای است که کیفیت پاسخ نهایی آن وابستگی زیادی به کیفیت شواهد بازیابی‌شده دارد. بازیابی متراکم، reranking و پالایش corpus سه عامل مهم برای رساندن شواهد مرتبط به مدل مولد هستند. پژوهش حاضر با تمرکز بر ریزتنظیم بازیاب BGE-M3 برای داده‌های پرسش‌وپاسخ فارسی پرسمان، در پی آن است که پاسخ صحیح را در رتبه‌های بالاتر قرار دهد و زمینه مناسب‌تری برای کاربردهای آتی RAG فارسی فراهم سازد. افزایش Recall@1، Recall@5 و MRR@10 فقط بهبود عددی در یک آزمایش نیست، بلکه نشان می‌دهد بازیاب با احتمال بیشتری شواهد مناسب را در اختیار مراحل بعدی یک سامانه پاسخ‌گو قرار می‌دهد."),
    )
    for title, description in rag_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۲-۶. مدل BGE-M3 و کاربرد آن در بازیابی معنایی فارسی", size=15, bold=True)
    for paragraph_text in (
        "مدل BGE-M3 مدل پایه پژوهش حاضر است. این مدل از خانواده BGE تولیدشده توسط Beijing Academy of Artificial Intelligence است و برای ساخت embedding متن و بازیابی اطلاعات طراحی شده است. هدف استفاده از BGE-M3 در این پژوهش، تبدیل پرسش‌ها و پاسخ‌های فارسی پرسمان به بازنمایی‌های برداری و رتبه‌بندی پاسخ‌ها بر اساس نزدیکی معنایی آن‌ها با پرسش است.",
        "انتخاب مدل پایه در بازیابی معنایی اهمیت زیادی دارد. مدل باید بتواند متن‌های فارسی را با کیفیت قابل قبول بازنمایی کند، از متون کوتاه و بلند پشتیبانی کند، امکان ریزتنظیم داشته باشد و در زیرساخت‌های موجود قابل اجرا باشد. BGE-M3 به دلیل چندزبانه‌بودن، طول ورودی نسبتاً زیاد، پشتیبانی از شیوه‌های مختلف بازیابی و سازگاری با کتابخانه FlagEmbedding، برای مسئله پژوهش حاضر انتخاب شده است.",
    ):
        add_text(document, paragraph_text)

    bge_sections = (
        ("۲-۶-۱. خانواده BGE", "BGE مخفف Beijing Academy of Artificial Intelligence General Embedding است و به مجموعه‌ای از مدل‌های embedding اشاره دارد که برای وظایف بازیابی، شباهت معنایی، خوشه‌بندی و رتبه‌بندی متن توسعه یافته‌اند. ایده محوری این خانواده آن است که متن‌ها به بردارهای عددی تبدیل شوند؛ به‌گونه‌ای که متن‌های مرتبط در فضای برداری به یکدیگر نزدیک‌تر باشند. در کاربردهای بازیابی، embedding اسناد یا پاسخ‌های corpus معمولاً از پیش تولید و ذخیره می‌شود و در زمان دریافت پرسش، فقط embedding پرسش ساخته و با بردارهای ذخیره‌شده مقایسه می‌شود."),
        ("۲-۶-۲. مفهوم چندزبانه، چندکارکردی و چندمقیاسی", "حرف M در نام مدل به multilingual، multi-functionality و multi-granularity اشاره دارد. ویژگی multilingual به توانایی مدل در پردازش زبان‌های گوناگون و ساخت فضای معنایی مشترک برای آن‌ها اشاره دارد. ویژگی multi-functionality به پشتیبانی هم‌زمان مدل از بازیابی متراکم، بازیابی تنک و بازیابی چندبرداری مربوط است. ویژگی multi-granularity نیز به توانایی پردازش متن‌هایی با طول‌های متفاوت، از پرسش‌های کوتاه و passageها تا اسناد بلند، اشاره دارد. مدل BAAI/bge-m3 دارای embeddingهای ۱۰۲۴بعدی و طول ورودی حداکثر ۸۱۹۲ token است."),
        ("۲-۶-۳. معماری پایه مدل", "BGE-M3 بر پایه encoder مدل XLM-RoBERTa توسعه یافته است. معماری encoder برای ساخت embedding در بازیابی مناسب است، زیرا می‌تواند یک متن را بدون نیاز به تولید متن جدید به بازنمایی برداری تبدیل کند. در مدل‌های encoder، هر token ورودی با توجه به tokenهای دیگر همان متن بازنمایی می‌شود. این ویژگی برای فارسی اهمیت دارد؛ زیرا شکل‌های نوشتاری، ساختار جمله، تفاوت میان زبان رسمی و محاوره‌ای و وابستگی‌های معنایی می‌توانند در تشخیص ارتباط میان پرسش و پاسخ مؤثر باشند."),
        ("۲-۶-۴. بازیابی متراکم در BGE-M3", "در بازیابی متراکم، BGE-M3 برای هر پرسش و پاسخ یک embedding واحد تولید می‌کند و پاسخ‌ها بر پایه شباهت برداری با پرسش رتبه‌بندی می‌شوند. در مقاله BGE-M3، بازنمایی token ویژه CLS برای بازیابی متراکم استفاده می‌شود. در پروژه حاضر نیز مدل در حالت dense و با cls pooling برای ساخت embedding پرسش‌ها و پاسخ‌ها به‌کار رفته است. این انتخاب با معیارهای Recall@1، Recall@5 و MRR@10، که بر اساس رتبه‌بندی پاسخ‌ها محاسبه می‌شوند، هم‌راستا است."),
        ("۲-۶-۵. بازیابی تنک و بازیابی چندبرداری", "BGE-M3 علاوه بر بازیابی متراکم، از بازیابی تنک و بازیابی چندبرداری پشتیبانی می‌کند. بازیابی تنک برای تطبیق واژگانی وزن‌دار و بازیابی چندبرداری برای تطبیق دقیق‌تر میان اجزای پرسش و پاسخ کاربرد دارد. مدل امکان ترکیب این شیوه‌ها را نیز فراهم می‌کند. بااین‌حال، پژوهش حاضر فقط بازیابی متراکم را ارزیابی می‌کند تا اثر پالایش داده، hard negativeهای بازبینی‌شده و ریزتنظیم مدل بر یک تنظیم آزمایشی مشخص و قابل مقایسه سنجیده شود. بنابراین، نتایج پژوهش حاضر ارزیابی همه قابلیت‌های BGE-M3 نیستند."),
        ("۲-۶-۶. آموزش مدل و خودتقطیری دانشی", "مقاله BGE-M3 روش خودتقطیری دانشی را برای آموزش یک مدل واحد با چند قابلیت بازیابی معرفی می‌کند. در این رویکرد، امتیازهای ارتباط حاصل از قابلیت‌های مختلف بازیابی به‌عنوان سیگنال آموزشی ترکیب می‌شوند. آموزش مدل از یک فرایند چندمرحله‌ای بهره می‌برد که شامل پیش‌آموزش بر داده‌های بدون برچسب و ریزتنظیم بر داده‌های برچسب‌دار و مصنوعی است. نمونه‌های منفی سخت نیز در مرحله ریزتنظیم وارد فرایند آموزش می‌شوند."),
        ("۲-۶-۷. دلیل انتخاب BGE-M3 برای داده‌های پرسمان", "داده‌های پرسمان شامل پرسش‌های واقعی کاربران و پاسخ‌های فارسی نسبتاً مفصل است. پرسش می‌تواند کوتاه، محاوره‌ای یا دارای املای غیررسمی باشد، در حالی که پاسخ رسمی، طولانی و چندبخشی است. بنابراین، تطابق واژگانی ساده برای بازیابی پاسخ مناسب کافی نیست. مدل چندزبانه BGE-M3 نقطه شروع مناسبی برای این مسئله است؛ بااین‌حال، چندزبانه‌بودن به‌تنهایی تضمین نمی‌کند که مدل در دامنه خاص پرسمان بهترین عملکرد را داشته باشد. به همین علت، مدل پایه با داده‌های پالایش‌شده همان دامنه و نمونه‌های منفی سخت ریزتنظیم شده است."),
        ("۲-۶-۸. طول ورودی مدل و محدودیت عملی پژوهش", "BGE-M3 از نظر معماری امکان پردازش تا ۸۱۹۲ token را فراهم می‌کند. بااین‌حال، استفاده از حداکثر طول ورودی در همه مراحل آموزش لزوماً تصمیم مناسبی نیست، زیرا مصرف حافظه GPU، زمان پردازش و هزینه آموزش را افزایش می‌دهد. در این پژوهش، طول پاسخ‌ها با tokenizer مدل BGE-M3 بررسی شده و پاسخ‌های بلندتر از ۱۰۲۴ token از داده نهایی کنار گذاشته شده‌اند. این محدودیت یک تصمیم مربوط به طراحی پژوهش است، نه محدودیت ذاتی BGE-M3."),
        ("۲-۶-۹. ریزتنظیم BGE-M3 در پژوهش حاضر", "در ریزتنظیم، پارامترهای مدل پایه با استفاده از داده‌های ساخت‌یافته پرسمان به‌روزرسانی می‌شوند. هر نمونه آموزشی شامل یک پرسش، یک پاسخ مثبت و هفت پاسخ منفی است. پاسخ‌های منفی از طریق بازیابی اولیه، بازرتبه‌بندی و کنترل کیفیت انتخاب می‌شوند تا از نظر معنایی به پرسش نزدیک، اما از نظر محتوایی پاسخ صحیح نباشند. نتیجه مورد انتظار آن است که مدل ریزتنظیم‌شده پاسخ صحیح را در رتبه‌های بالاتری نسبت به مدل پایه قرار دهد. این ادعا با ارزیابی یکسان دو مدل بر مجموعه آزمون ثابت بررسی می‌شود."),
        ("۲-۶-۱۰. محدودیت‌ها و ملاحظات", "BGE-M3 یک مدل عمومی و چندزبانه است و عملکرد آن در هر زبان و هر دامنه باید به‌صورت تجربی ارزیابی شود. مدل‌کارت BGE-M3 نیز تصریح می‌کند که عملکرد آن در بازیابی تک‌زبانه ممکن است از مدل‌هایی که مشخصاً برای همان زبان طراحی شده‌اند بهتر نباشد. محدودیت دیگر آن است که پژوهش حاضر تنها حالت dense را بررسی می‌کند. استفاده از بازیابی ترکیبی، sparse retrieval، multi-vector retrieval، شاخص‌گذاری تقریبی و reranking نهایی در سامانه عملی می‌تواند بر کیفیت و سرعت اثر بگذارد و باید در مطالعات بعدی سنجیده شود."),
        ("۲-۶-۱۱. جایگاه BGE-M3 در ارزیابی‌های چندزبانه و فارسی", "مدل BGE-M3 از مدل‌های شناخته‌شده و پرکاربرد embedding چندزبانه است که به‌دلیل پشتیبانی هم‌زمان از بازیابی متراکم، بازیابی تنک و بازیابی چندبرداری، در پژوهش‌ها و سامانه‌های بازیابی اطلاعات کاربرد دارد. مقاله اصلی BGE-M3 عملکرد مدل را در بازیابی چندزبانه، بازیابی میان‌زبانی و بازیابی اسناد بلند بررسی کرده است. در این مقاله، بازنمایی CLS برای بازیابی متراکم استفاده می‌شود و در مرحله ریزتنظیم، برای هر پرسش هفت نمونه منفی به‌کار گرفته شده است. این نکته با ساختار داده آموزشی پژوهش حاضر، شامل یک پاسخ مثبت و هفت پاسخ منفی برای هر پرسش، هم‌راستا است."),
        ("۲-۶-۱۲. شواهد عملکرد در بنچ‌مارک فارسی چراغ", "چندزبانه‌بودن BGE-M3 به این معنا نیست که مدل در همه زبان‌ها و همه دامنه‌ها عملکرد یکسان یا بهترین عملکرد را خواهد داشت. ازاین‌رو، استفاده از BGE-M3 در این پژوهش به‌عنوان یک مدل پایه معتبر انجام شده است و اثر ریزتنظیم آن باید به‌صورت مستقل بر مجموعه آزمون پرسمان سنجیده شود. در بنچ‌مارک فارسی چراغ، cheRAGh-Embedding، مدل BAAI/bge-m3 در میان مدل‌های ارزیابی‌شده رتبه پنجم را کسب کرده است. میانگین MRR این مدل ۰٫۸۶۴۹ و میانگین Recall@5 آن ۰٫۹۲۲۹ گزارش شده است. این بنچ‌مارک مدل‌ها را در پنج حوزه آموزشی، عمومی، حقوقی، دینی و علمی ارزیابی می‌کند. بنابراین، انتخاب BGE-M3 به‌عنوان مدل پایه برای داده‌های فارسی پرسمان، علاوه بر ویژگی‌های معماری مدل، با نتایج قابل توجه آن در یک بنچ‌مارک فارسی نیز پشتیبانی می‌شود. بااین‌حال، این رتبه به‌معنای برتری قطعی مدل در همه مجموعه‌داده‌های فارسی نیست و ارزیابی اختصاصی بر داده پرسمان همچنان ضروری است."),
    )
    for title, description in bge_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۲-۷. ارزیابی بازیابی و معیارهای رتبه‌بندی", size=15, bold=True)
    for paragraph_text in (
        "ارزیابی در بازیابی اطلاعات با ارزیابی تولید متن تفاوت دارد. در بازیابی، هدف اصلی این نیست که مدل پاسخ جدیدی تولید کند؛ بلکه باید پاسخ صحیح یا سند مرتبط را از میان مجموعه‌ای از پاسخ‌ها پیدا و در رتبه‌های ابتدایی قرار دهد. بنابراین، خروجی مدل برای هر پرسش یک فهرست مرتب‌شده از پاسخ‌های corpus است و کیفیت مدل با بررسی جایگاه پاسخ صحیح در این فهرست سنجیده می‌شود.",
        "در پژوهش حاضر، هر پرسش آزمون دقیقاً یک پاسخ صحیح متناظر دارد. همه پاسخ‌های باقی‌مانده در corpus به‌عنوان نامزدهای بازیابی در نظر گرفته می‌شوند. مدل برای پرسش embedding می‌سازد، embedding آن را با embedding پاسخ‌های corpus مقایسه می‌کند و پاسخ‌ها را بر اساس امتیاز شباهت مرتب می‌سازد. سپس رتبه پاسخ صحیح ثبت می‌شود. این فرایند برای همه پرسش‌های آزمون تکرار می‌شود.",
    ):
        add_text(document, paragraph_text)

    evaluation_sections = (
        ("۲-۷-۱. رتبه پاسخ صحیح", "رتبه پاسخ صحیح برای یک پرسش، جایگاه آن پاسخ در فهرست نتایج مرتب‌شده است. اگر پاسخ صحیح نخستین نتیجه باشد، رتبه آن برابر ۱ است و اگر پس از سه پاسخ دیگر قرار گیرد، رتبه آن ۴ خواهد بود. هرچه رتبه کوچک‌تر باشد، عملکرد مدل بهتر است. ثبت رتبه پاسخ صحیح برای هر پرسش اهمیت زیادی دارد، زیرا معیارهای کلی تنها میانگینی از عملکرد مدل را نشان می‌دهند؛ اما رتبه هر پرسش امکان تحلیل خطا را فراهم می‌کند."),
        ("۲-۷-۲. معیار Recall@k", "Recall@k نشان می‌دهد چه نسبتی از پرسش‌ها پاسخ صحیح خود را در میان k نتیجه اول بازیابی کرده‌اند. در این رابطه، Q مجموعه پرسش‌های آزمون، |Q| تعداد پرسش‌ها و rank(q) رتبه پاسخ صحیح پرسش q است. تابع شاخص I در صورتی که پاسخ صحیح در میان k نتیجه اول باشد مقدار ۱ و در غیر این صورت مقدار ۰ می‌گیرد. برای مثال، اگر از میان ۱۰۰ پرسش آزمون، پاسخ صحیح ۹۲ پرسش در میان پنج نتیجه اول باشد، مقدار Recall@5 برابر ۰٫۹۲ یا ۹۲ درصد است."),
        ("۲-۷-۳. Recall@1 و Recall@5", "Recall@1 سخت‌گیرانه‌ترین معیار در میان معیارهای مورد استفاده این پژوهش است و فقط زمانی موفقیت را ثبت می‌کند که پاسخ صحیح دقیقاً در رتبه اول باشد. Recall@5 انعطاف بیشتری دارد. در بسیاری از کاربردهای واقعی، به‌ویژه در RAG، بازیاب چند نتیجه برتر را در اختیار مرحله بعدی قرار می‌دهد. اگر پاسخ صحیح در میان پنج نتیجه اول باشد، reranker یا مدل مولد همچنان می‌تواند از آن استفاده کند. استفاده هم‌زمان از Recall@1 و Recall@5 تفاوت میان کیفیت پاسخ اول و کیفیت مجموعه نتایج ابتدایی را روشن می‌کند."),
        ("۲-۷-۴. معیار MRR@10", "MRR مخفف Mean Reciprocal Rank است و علاوه بر وجود پاسخ صحیح در نتایج ابتدایی، به رتبه دقیق آن نیز حساس است. برای هر پرسش، معکوس رتبه پاسخ صحیح محاسبه می‌شود. اگر پاسخ صحیح در رتبه اول باشد مقدار آن ۱، در رتبه دوم مقدار آن ۱/۲ و در رتبه پنجم مقدار آن ۱/۵ خواهد بود. در MRR@10، فقط رتبه‌های ۱ تا ۱۰ در محاسبه دخالت دارند و اگر پاسخ صحیح پایین‌تر از رتبه دهم باشد، سهم آن پرسش صفر در نظر گرفته می‌شود."),
        ("۲-۷-۵. دلیل انتخاب معیارهای پژوهش", "انتخاب Recall@1، Recall@5 و MRR@10 سه جنبه مکمل از عملکرد بازیاب را پوشش می‌دهد. Recall@1 کیفیت نتیجه اول را می‌سنجد. Recall@5 بررسی می‌کند که آیا پاسخ صحیح در مجموعه کوچکی از نتایج قابل ارائه به کاربر یا مدل مولد حضور دارد یا خیر. MRR@10 نیز کیفیت رتبه‌بندی پاسخ صحیح را در ده نتیجه نخست نشان می‌دهد. افزایش هم‌زمان این سه معیار نشانه آن است که مدل هم در بازیابی پاسخ صحیح و هم در رتبه‌بندی آن بهتر عمل کرده است."),
        ("۲-۷-۶. مجموعه آزمون ثابت و مقایسه منصفانه", "برای مقایسه معتبر مدل‌ها، مجموعه آزمون باید از داده آموزش جدا باشد و در طول آزمایش‌ها ثابت بماند. اگر مدل پایه و مدل ریزتنظیم‌شده بر پرسش‌ها یا corpusهای متفاوتی ارزیابی شوند، تفاوت معیارها الزاماً ناشی از کیفیت مدل نخواهد بود. به همین علت، در پژوهش حاضر هر دو مدل بر یک مجموعه آزمون مشترک و با corpus پاسخ یکسان ارزیابی می‌شوند. پالایش داده و حذف تکرارها پیش از تفکیک آموزش و آزمون نیز از نشت داده جلوگیری می‌کند."),
        ("۲-۷-۷. تنظیمات مؤثر بر ارزیابی", "مقدار معیارهای بازیابی تنها به وزن‌های مدل وابسته نیست. tokenizer، طول مجاز ورودی، روش pooling، نرمال‌سازی embeddingها، نوع شباهت، اندازه batch، روش ساخت corpus، وجود یا نبود reranker و تعداد نتایج بازیابی‌شده نیز می‌توانند نتیجه را تغییر دهند. بنابراین، برای مقایسه منصفانه باید این تنظیمات برای مدل پایه و مدل ریزتنظیم‌شده یکسان نگه داشته شوند. در پژوهش حاضر، رتبه پاسخ صحیح هر پرسش و مقدار معیارهای نهایی نیز ذخیره می‌شود تا ارزیابی مدل‌های بعدی با همان تنظیمات تکرارپذیر باشد."),
        ("۲-۷-۸. تحلیل خطا بر اساس رتبه‌ها", "ذخیره رتبه هر پرسش، امکان تحلیل عمیق‌تر از یک عدد میانگین را فراهم می‌کند. پرسش‌هایی که پاسخ صحیح آن‌ها در رتبه ۱ قرار گرفته‌اند نمونه‌های موفق مدل هستند. پرسش‌هایی که پاسخ صحیح آن‌ها در رتبه‌های ۲ تا ۵ قرار گرفته‌اند نشان می‌دهند مدل پاسخ درست را بازیابی کرده، اما در تفکیک آن از پاسخ‌های نزدیک نیاز به بهبود دارد. پرسش‌هایی که پاسخ صحیح آن‌ها خارج از ده رتبه اول قرار گرفته‌اند، نمونه‌های شکست جدی‌تر به شمار می‌روند."),
        ("۲-۷-۹. محدودیت معیارهای رتبه‌بندی", "Recall و MRR کیفیت بازیابی را می‌سنجند، اما همه جنبه‌های کیفیت یک سامانه پاسخ‌گو را پوشش نمی‌دهند. این معیارها بررسی نمی‌کنند که پاسخ بازیابی‌شده از نظر نگارشی روان است یا خیر، آیا برای کاربر قابل فهم است، آیا کامل است، یا در یک سامانه RAG مدل مولد چگونه از آن استفاده می‌کند. همچنین، در این پژوهش برای هر پرسش یک پاسخ صحیح اصلی در نظر گرفته شده است؛ در برخی موارد ممکن است پاسخ دیگری نیز از نظر محتوایی قابل قبول باشد، اما در داده مرجع به‌عنوان پاسخ صحیح همان پرسش ثبت نشده باشد."),
        ("۲-۷-۱۰. جمع‌بندی", "معیارهای Recall@1، Recall@5 و MRR@10 ابزارهای اصلی سنجش کیفیت بازیابی در پژوهش حاضر هستند. این معیارها نشان می‌دهند مدل با چه دقتی پاسخ صحیح را بازیابی می‌کند و آن را در چه جایگاهی قرار می‌دهد. استفاده از مجموعه آزمون ثابت، corpus یکسان، تنظیمات ارزیابی مشترک و ذخیره رتبه هر پرسش، مقایسه مدل پایه و مدل ریزتنظیم‌شده را قابل اعتماد و بازتولیدپذیر می‌سازد."),
    )
    for title, description in evaluation_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)
        if title.startswith("۲-۷-۲."):
            add_recall_equation(document, "۲-۴")
        if title.startswith("۲-۷-۴."):
            add_mrr_equation(document, "۲-۵")

    add_text(document, "۲-۸. پیشینه پژوهش و شکاف تحقیق", size=15, bold=True)
    related_work_sections = (
        ("۲-۸-۱. Sentence-BERT و شکل‌گیری embeddingهای جمله", "Sentence-BERT یا SBERT از پژوهش‌های مهم در تبدیل مدل‌های BERT به مدل‌هایی مناسب برای ساخت embedding جمله است. در BERT معمولی، برای سنجش شباهت دو جمله باید هر زوج جمله به‌صورت هم‌زمان به مدل داده شود؛ این روش برای مقایسه یک پرسش با corpus بزرگ هزینه محاسباتی زیادی دارد. SBERT با ساختار siamese و triplet، embedding مستقل برای هر جمله تولید می‌کند تا شباهت آن‌ها با معیارهایی مانند کسینوسی محاسبه شود. اهمیت SBERT در این است که مدل‌های embedding را برای جست‌وجوی معنایی، خوشه‌بندی و بازیابی سریع عملی کرد."),
        ("۲-۸-۲. بازیابی متراکم پرسش‌وپاسخ", "پژوهش Dense Passage Retrieval یا DPR نشان داد که می‌توان با یک ساختار دو-رمزگذار، پرسش و passage را جداگانه به بردار تبدیل کرد و پاسخ مرتبط را از میان corpus بازیابی نمود. در این رویکرد، مدل می‌آموزد بردار پرسش را به بردار passage صحیح نزدیک کند. DPR در بازیابی پرسش‌وپاسخ حوزه‌باز اثرگذار بود، زیرا نشان داد بازیابی متراکم، حتی با تعداد محدودی نمونه نظارت‌شده، می‌تواند در برخی مجموعه‌داده‌ها از BM25 پیشی بگیرد. بااین‌حال، کارایی این روش به کیفیت زوج‌های پرسش و پاسخ، انتخاب پاسخ‌های منفی و شباهت میان داده آموزش و دامنه هدف وابسته است."),
        ("۲-۸-۳. نقش نمونه‌های منفی سخت", "نمونه‌های منفی سخت از مسائل مهم در آموزش بازیاب‌های متراکم هستند. پاسخ‌های منفی تصادفی معمولاً به‌راحتی از پاسخ صحیح متمایز می‌شوند و اطلاعات آموزشی محدودی برای مدل فراهم می‌کنند. در مقابل، پاسخ‌های نزدیک اما نادرست، مدل را مجبور می‌کنند تفاوت‌های دقیق‌تر معنایی را یاد بگیرد. روش ANCE از نمونه‌های منفی استخراج‌شده با بازیابی تقریبی نزدیک‌ترین همسایه‌ها استفاده می‌کند. این ایده مبنای مرحله بازیابی اولیه در پژوهش حاضر است؛ با این تفاوت که نامزدهای نزدیک پس از بازیابی، بازرتبه‌بندی و کنترل کیفیت نیز می‌شوند تا خطر ورود پاسخ منفی کاذب کاهش یابد."),
        ("۲-۸-۴. مسئله پاسخ منفی کاذب", "هر پاسخ نزدیک الزاماً یک negative مناسب نیست. در مجموعه‌های پرسش‌وپاسخ، ممکن است یک پاسخ برای پرسش دیگری ثبت شده باشد، اما از نظر محتوایی پاسخ قابل قبول یا مکمل پرسش فعلی نیز محسوب شود. واردکردن چنین پاسخی به‌عنوان negative، مدل را به دورکردن دو متن مرتبط تشویق می‌کند و می‌تواند کیفیت بازیابی را کاهش دهد. بخش مهم روش پیشنهادی این پژوهش، افزودن کنترل کیفیت پس از بازیابی و reranking است. در این مرحله، پاسخ‌های نامزد با توجه به پرسش و پاسخ مثبت بررسی می‌شوند و موارد هم‌ارز، مکمل یا مبهم کنار گذاشته می‌شوند."),
        ("۲-۸-۵. مدل‌های embedding چندزبانه", "رشد مدل‌های چندزبانه باعث شده است که بازیابی معنایی برای زبان‌هایی با داده آموزشی محدودتر نیز امکان‌پذیر شود. مدل‌هایی مانند mDPR، mContriever، multilingual-e5 و BGE-M3 برای ساخت فضای معنایی مشترک میان زبان‌ها طراحی شده‌اند. مزیت این مدل‌ها، بهره‌گیری از داده‌های متنوع و توانایی بالقوه بازیابی تک‌زبانه و میان‌زببانی است. بااین‌حال، عملکرد مدل چندزبانه در یک دامنه خاص تضمین‌شده نیست و پرسش‌های پرسمان از نظر موضوع، لحن، طول و سبک نگارش با داده‌های عمومی متفاوت‌اند."),
        ("۲-۸-۶. BGE-M3 و مدل‌های چندکارکردی بازیابی", "BGE-M3 بر بازیابی متراکم، بازیابی تنک و بازیابی چندبرداری تمرکز دارد و برای بیش از ۱۰۰ زبان و متن‌هایی تا ۸۱۹۲ token طراحی شده است. مقاله اصلی این مدل از آموزش چندمرحله‌ای، داده‌های چندزبانه و خودتقطیری دانشی برای تقویت هم‌زمان قابلیت‌های بازیابی استفاده می‌کند. همچنین در مرحله ریزتنظیم مقاله، برای هر پرسش هفت نمونه منفی استفاده شده است. در پژوهش حاضر، BGE-M3 فقط در حالت dense ارزیابی می‌شود تا اثر ریزتنظیم، داده پالایش‌شده و نمونه‌های منفی سخت در یک پیکربندی روشن بررسی شود."),
        ("۲-۸-۷. شکاف تحقیق", "مرور پژوهش‌ها نشان می‌دهد که مدل‌های عمومی و چندزبانه قدرتمندی برای بازیابی متراکم وجود دارند، اما سازگارکردن آن‌ها با corpus واقعی پرسش‌وپاسخ فارسی همچنان نیازمند طراحی داده و ارزیابی اختصاصی است. در داده‌های پرسمان، پاسخ‌های تکراری، پاسخ‌های بلند، پرسش‌های محاوره‌ای و پاسخ‌های نزدیک از نظر موضوعی می‌توانند آموزش و ارزیابی را دچار خطا کنند. شکاف مورد توجه پژوهش حاضر، نبود یک خط لوله بازتولیدپذیر برای تبدیل داده خام پرسمان به داده مناسب ریزتنظیم بازیابی است."),
        ("۲-۸-۸. جایگاه پژوهش حاضر", "پژوهش حاضر مدل پایه جدیدی از ابتدا آموزش نمی‌دهد و یک سامانه RAG کامل نیز پیاده‌سازی نمی‌کند. سهم آن، طراحی و ارزیابی یک فرایند داده‌محور برای ریزتنظیم BGE-M3 در بازیابی پاسخ‌های فارسی پرسمان است. این خط لوله شامل پالایش مرحله‌ای، حذف تکرار پرسش و پاسخ، کنترل طول پاسخ با tokenizer مدل، تفکیک ثابت داده آزمون، تولید نامزدهای نزدیک، بازرتبه‌بندی، کنترل کیفیت پاسخ‌های منفی و ساخت رکوردهای آموزش با یک پاسخ مثبت و هفت پاسخ منفی است. مدل پایه و مدل ریزتنظیم‌شده با مجموعه آزمون، corpus و تنظیمات یکسان مقایسه می‌شوند."),
    )
    for title, description in related_work_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۲-۹. جمع‌بندی فصل", size=15, bold=True)
    for paragraph_text in (
        "در این فصل، مبانی نظری و پیشینه لازم برای پژوهش حاضر بررسی شد. ابتدا تحول بازیابی اطلاعات از روش‌های واژگانی به بازیابی متراکم توضیح داده شد. روش‌های واژگانی مانند TF-IDF و BM25 بر تطابق واژه‌ها متکی هستند و برای عبارت‌های دقیق، نام‌های خاص و کلیدواژه‌های کمیاب مفیدند؛ اما در تشخیص ارتباط میان پرسش و پاسخ‌هایی که با واژه‌های متفاوت یک مفهوم مشترک را بیان می‌کنند، محدودیت دارند. بازیابی متراکم با تبدیل متن‌ها به embedding و سنجش شباهت آن‌ها در فضای برداری، راهی برای کاهش این محدودیت فراهم می‌کند.",
        "سپس مفاهیم بازنمایی برداری متن، Transformer، BERT، pooling و شباهت کسینوسی مطرح شد. در بازیابی متراکم، پرسش و پاسخ به بردار تبدیل می‌شوند و پاسخ‌ها بر اساس میزان شباهت با پرسش رتبه می‌گیرند. در پژوهش حاضر، مدل BGE-M3 در حالت dense و با cls pooling به‌کار گرفته شده است. نرمال‌سازی بردارها نیز باعث می‌شود ضرب داخلی آن‌ها معادل شباهت کسینوسی باشد و رتبه‌بندی پاسخ‌ها به‌صورت کارآمد انجام شود.",
        "بخش دیگری از فصل به یادگیری تقابلی و اهمیت نمونه‌های منفی سخت اختصاص یافت. مدل embedding فقط با مشاهده پاسخ‌های کاملاً نامرتبط، توانایی کافی برای تشخیص پاسخ درست از پاسخ‌های نزدیک اما نادرست به دست نمی‌آورد. به همین علت، hard negativeها در ساخت داده آموزش اهمیت دارند. با وجود این، پاسخ نزدیک ممکن است در واقع پاسخی قابل قبول، مکمل یا هم‌ارز باشد. این وضعیت false negative نام دارد و می‌تواند به مدل آسیب بزند. روش پژوهش حاضر با بازیابی اولیه، reranking و کنترل کیفیت، در پی انتخاب پاسخ‌های منفی نزدیک اما معتبر است.",
        "همچنین جایگاه بازیاب در معماری RAG بررسی شد. در یک سامانه RAG، مدل مولد تنها به شواهدی دسترسی دارد که بازیاب در اختیار آن قرار می‌دهد. بنابراین، قرارگرفتن پاسخ صحیح در رتبه‌های ابتدایی، بر کیفیت پاسخ نهایی اثر مستقیم دارد. پژوهش حاضر یک سامانه RAG کامل پیاده‌سازی نمی‌کند؛ اما مدل ریزتنظیم‌شده و corpus پالایش‌شده آن می‌توانند به‌عنوان بخش retriever در یک سامانه RAG فارسی استفاده شوند.",
        "مدل BGE-M3 به‌عنوان مدل پایه معرفی شد. این مدل از قابلیت‌های چندزبانه، چندکارکردی و چندمقیاسی پشتیبانی می‌کند و در پژوهش حاضر تنها قابلیت بازیابی dense آن ارزیابی می‌شود. ویژگی‌های مدل، از جمله embeddingهای ۱۰۲۴بعدی، ظرفیت پردازش متن بلند و امکان ریزتنظیم، آن را برای مسئله بازیابی پاسخ در پرسمان مناسب می‌سازد. بااین‌حال، هیچ مدل عمومی یا چندزبانه‌ای بدون ارزیابی در دامنه هدف برتر تلقی نمی‌شود؛ به همین علت، مدل پایه و مدل ریزتنظیم‌شده بر داده آزمون ثابت پرسمان مقایسه خواهند شد.",
        "در ادامه، معیارهای ارزیابی Recall@1، Recall@5 و MRR@10 بررسی شدند. این معیارها سه جنبه مکمل از کیفیت بازیابی را نشان می‌دهند: کیفیت پاسخ اول، حضور پاسخ صحیح در چند نتیجه ابتدایی و دقت رتبه‌بندی پاسخ صحیح. ثبت رتبه هر پرسش نیز امکان تحلیل خطا و بررسی نمونه‌های موفق و ناموفق را فراهم می‌کند. ثابت نگه‌داشتن مجموعه آزمون، corpus و تنظیمات ارزیابی، شرط لازم برای مقایسه منصفانه مدل‌ها است.",
        "در بخش پیشینه پژوهش، نقش Sentence-BERT در فراگیرشدن embeddingهای جمله، نقش DPR در بازیابی متراکم پرسش‌وپاسخ و نقش ANCE در انتخاب نمونه‌های منفی سخت مرور شد. بررسی این پژوهش‌ها نشان می‌دهد که کیفیت بازیاب به تعامل میان مدل پایه، داده آموزشی، نوع negativeها و روش ارزیابی وابسته است. شکاف مورد توجه پژوهش حاضر، طراحی یک فرایند بازتولیدپذیر برای تبدیل داده خام پرسش‌وپاسخ فارسی پرسمان به داده‌ای مناسب برای ریزتنظیم بازیاب متراکم است.",
        "بر این اساس، فصل سوم به روش اجرای پژوهش اختصاص دارد. در آن فصل، منبع داده، مراحل پالایش، ایجاد شناسه یکتا، حذف تکرارها، کنترل طول پاسخ‌ها، تفکیک داده آموزش و آزمون، تولید پاسخ‌های منفی سخت، ساخت داده آموزش، تنظیمات ریزتنظیم BGE-M3 و روش ارزیابی مدل‌ها به‌صورت دقیق تشریح خواهد شد.",
    ):
        add_text(document, paragraph_text)

    document.add_page_break()
    add_text(document, "فصل سوم: روش اجرای پژوهش", size=18, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(document, "۳-۱. مقدمه", size=15, bold=True)
    for paragraph_text in (
        "این فصل روش اجرای پژوهش حاضر را تشریح می‌کند. هدف پژوهش، سنجش اثر پالایش داده‌های پرسش‌وپاسخ فارسی پرسمان، تولید نمونه‌های منفی سخت و ریزتنظیم مدل BAAI/bge-m3 بر کیفیت بازیابی پاسخ صحیح است. برای دستیابی به این هدف، یک فرایند مرحله‌ای طراحی شده است که از داده خام آغاز می‌شود و به مقایسه مدل پایه و مدل ریزتنظیم‌شده بر مجموعه آزمون ثابت پایان می‌یابد.",
        "طرح پژوهش از نوع تجربی و کمّی است. در این طرح، مدل پایه به‌عنوان خط مبنا در نظر گرفته می‌شود. سپس همان مدل با داده‌های آموزشی ساخته‌شده از corpus پرسمان ریزتنظیم می‌شود. در پایان، هر دو مدل با corpus پاسخ یکسان، پرسش‌های آزمون یکسان و تنظیمات بازیابی یکسان ارزیابی می‌شوند. بنابراین، اختلاف معیارهای ارزیابی را می‌توان به اثر فرایند ساخت داده و ریزتنظیم مدل نسبت داد، نه به تفاوت در مجموعه آزمون یا corpus بازیابی.",
        "روش پیشنهادی از چند مرحله وابسته به هم تشکیل شده است. ابتدا داده خام پرسمان خوانده و فیلترهای کیفیت بر پرسش‌ها و پاسخ‌ها اعمال می‌شود. در ادامه، شناسه پایدار برای رکوردهای باقی‌مانده ساخته می‌شود و پرسش‌ها و پاسخ‌های تکراری حذف می‌شوند. سپس داده پالایش‌شده به دو بخش آموزش و آزمون تفکیک می‌شود؛ به‌گونه‌ای که داده آزمون در مراحل آموزش و انتخاب نمونه‌های آموزشی استفاده نشود.",
        "در بخش ساخت داده آموزش، برای هر پرسش یک پاسخ مثبت و چند پاسخ منفی انتخاب می‌شود. پاسخ مثبت، پاسخ متناظر همان پرسش است. برای انتخاب پاسخ‌های منفی، ابتدا مدل پایه پاسخ‌های نزدیک را بازیابی می‌کند. سپس نامزدها با reranker بازرتبه‌بندی می‌شوند. پاسخ صحیح کنار گذاشته می‌شود و پاسخ‌های نزدیک با کمک کنترل کیفیت بررسی می‌شوند تا پاسخ هم‌ارز، مکمل یا پاسخ منفی کاذب وارد داده آموزش نشود. در نهایت، برای هر پرسش یک پاسخ مثبت و هفت پاسخ منفی معتبر انتخاب می‌شود.",
        "مدل BGE-M3 با داده آموزشی حاصل ریزتنظیم می‌شود. در این پژوهش، تمرکز بر قابلیت بازیابی dense مدل است. پرسش و پاسخ به embedding تبدیل می‌شوند و مدل می‌آموزد شباهت پرسش با پاسخ مثبت را نسبت به پاسخ‌های منفی افزایش دهد. تنظیمات آموزش، از جمله طول ورودی، تعداد نمونه‌های منفی، تعداد دوره‌های آموزش، نرخ یادگیری و شیوه pooling، ثبت می‌شوند تا آزمایش قابل بازتولید باشد.",
        "پیش از ریزتنظیم، عملکرد مدل پایه بر مجموعه آزمون محاسبه می‌شود. پس از آموزش نیز مدل ریزتنظیم‌شده دقیقاً بر همان پرسش‌ها و همان corpus ارزیابی می‌شود. در هر ارزیابی، رتبه پاسخ صحیح برای هر پرسش ذخیره می‌شود و معیارهای Recall@1، Recall@5 و MRR@10 محاسبه می‌شوند. نگهداری رتبه‌ها، علاوه بر تولید معیارهای کلی، امکان تحلیل خطا و بررسی پرسش‌هایی را فراهم می‌کند که مدل در بازیابی آن‌ها موفق یا ناموفق بوده است.",
        "برای افزایش قابلیت بازتولید، مراحل کلیدی پژوهش در notebookهای مستقل پیاده‌سازی شده‌اند. یک notebook برای پالایش داده پرسمان، یک notebook برای ساخت داده آموزش، یک notebook برای کنترل کیفیت نمونه‌های منفی سخت و یک notebook برای ارزیابی مدل‌ها ایجاد شده است. خروجی هر مرحله در مسیر مشخص ذخیره می‌شود و شناسه رکوردها میان مراحل حفظ می‌شود. به این ترتیب، می‌توان مسیر هر نمونه را از داده خام تا داده آموزش یا آزمون دنبال کرد.",
        "در ادامه فصل، ابتدا منبع داده و ساختار اولیه آن معرفی می‌شود. سپس مراحل پالایش داده، ساخت شناسه، حذف تکرارها و کنترل طول متن‌ها تشریح خواهد شد. پس از آن، روش تفکیک آموزش و آزمون، تولید hard negative، ساخت داده آموزش، تنظیمات ریزتنظیم BGE-M3 و روش ارزیابی نهایی ارائه می‌شود.",
    ):
        add_text(document, paragraph_text)

    add_text(document, "۳-۲. منبع داده و چالش‌های کیفیت مجموعه پرسمان", size=15, bold=True)
    for paragraph_text in (
        "داده مورد استفاده در این پژوهش از مجموعه پرسش‌وپاسخ فارسی پرسمان تهیه شده است. پرسمان خود را پاسخ‌گوی سؤالات دانشگاهیان معرفی می‌کند و محتوای آن حوزه‌هایی مانند قرآن و حدیث، کلام و دین‌پژوهی، اخلاق، احکام، مشاوره، تاریخ و سیره، سیاست، فرهنگ و مهدویت را پوشش می‌دهد. این تنوع موضوعی، مجموعه را برای سنجش بازیابی معنایی فارسی ارزشمند می‌کند؛ اما هم‌زمان مسئله بازیابی را دشوارتر می‌سازد، زیرا پرسش‌ها و پاسخ‌ها از نظر واژگان، سبک نگارش، طول و سطح تخصص یکدست نیستند.",
        "داده خام پژوهش در قالب CSV شامل ۴۷٬۳۶۳ رکورد، بدون احتساب سطر سرستون، دریافت شد. هر رکورد دست‌کم شامل متن پرسش و متن پاسخ بود. وجود پیوند اولیه میان پرسش و پاسخ یک مزیت مهم برای ساخت داده بازیابی محسوب می‌شود، زیرا پاسخ متناظر هر پرسش می‌تواند نمونه مثبت اولیه باشد. بااین‌حال، چنین پیوندی به‌تنهایی برای استفاده مستقیم در آموزش مدل embedding کافی نیست. داده پرسش‌وپاسخ وبی معمولاً برای نمایش به کاربر طراحی شده است، نه برای آموزش کنترل‌شده مدل بازیابی.",
    ):
        add_text(document, paragraph_text)

    data_quality_sections = (
        ("۳-۲-۱. ناهمگونی موضوعی و سبکی", "پرسش‌های پرسمان از حوزه‌های محتوایی گوناگون آمده‌اند. برخی پرسش‌ها کوتاه و مستقیم‌اند، مانند یک پرسش فقهی یا اعتقادی مشخص؛ برخی دیگر شرح مفصل تجربه فردی، مسئله خانوادگی یا زمینه تاریخی هستند. پاسخ‌ها نیز از یک جمله کوتاه تا متن‌های تحلیلی چندپاراگرافی، نقل‌قول، ارجاع و فهرست را دربرمی‌گیرند. این ناهمگونی از یک سو باعث می‌شود corpus به داده واقعی کاربران نزدیک باشد و از سوی دیگر بازیابی را دشوار می‌کند. مدل باید بتواند میان پرسش کوتاه و پاسخ بلند، زبان محاوره‌ای و زبان رسمی، یا واژه‌های متفاوت با معنای مشابه ارتباط برقرار کند."),
        ("۳-۲-۲. نبود ساختار پرسشی در بخشی از ستون پرسش", "یکی از مهم‌ترین مشکلات داده خام، این بود که همه مقادیر ستون پرسش واقعاً پرسش نبودند. بخشی از رکوردها فاقد علامت سؤال بودند یا متن موجود در ستون پرسش از نظر ساختار به مقاله، پاسخ، توضیح طولانی، فهرست منابع یا بخشی از یک مطلب شباهت داشت. اگر یک متن مقاله‌مانند به‌عنوان query وارد آموزش شود، مدل به‌جای یادگیری رابطه طبیعی میان پرسش کاربر و پاسخ، با رابطه‌ای غیرواقعی میان دو متن بلند روبه‌رو می‌شود. برای کنترل این مسئله، وجود علامت سؤال فارسی یا لاتین در انتهای پرسش به‌عنوان یک شرط اولیه در نظر گرفته شد."),
        ("۳-۲-۳. پرسش‌های بسیار بلند و مقاله‌مانند", "در بخشی از داده، متن ستون پرسش از حد یک پرسش طبیعی بسیار طولانی‌تر بود. پرسش‌هایی با طول نرمال‌شده ۱۲۰۰ کاراکتر یا بیشتر کنار گذاشته شدند. همچنین متن‌هایی که ترکیبی از چند نشانه مقاله‌مانند داشتند، حذف شدند. این نشانه‌ها شامل طول زیاد، عبارت‌های رایج در مقاله یا نوشتار، تعداد زیاد ارجاع، فهرست شماره‌دار و چند پاراگراف مجزا بود. این تصمیم به‌معنای نامعتبر دانستن پرسش‌های مفصل نیست؛ برای کاهش حذف نادرست پرسش‌های طولانی واقعی، وجود یک نشانه به‌تنهایی دلیل حذف نبود و ترکیب چند نشانه لازم بود."),
        ("۳-۲-۴. جابه‌جایی یا تکرار متن میان ستون پرسش و پاسخ", "برخی رکوردها در ستون پرسش متنی داشتند که با یکی از پاسخ‌های مجموعه یکسان بود. این حالت می‌تواند ناشی از خطای ورود داده، جابه‌جایی ستون‌ها، درج پاسخ به‌جای پرسش یا تکرار ناخواسته محتوا باشد. چنین رکوردی برای آموزش بازیابی مناسب نیست، زیرا رابطه پرسش و پاسخ را به یک رابطه غیرطبیعی متن با متن تبدیل می‌کند. برای شناسایی این موارد، متن‌ها پیش از مقایسه نرمال‌سازی شدند تا تفاوت‌های ظاهری مانع تشخیص تکرار دقیق نشوند."),
        ("۳-۲-۵. تفاوت‌های نگارشی زبان فارسی", "زبان فارسی در داده‌های وب با شکل‌های نگارشی متنوعی ظاهر می‌شود. تفاوت میان ی فارسی و عربی، ک فارسی و عربی، نیم‌فاصله، فاصله عادی، فاصله زائد، نویسه‌های نامرئی، علائم نگارشی و شیوه نوشتن اعداد می‌تواند دو متن معنایی یکسان را از نظر رشته‌ای متفاوت نشان دهد. نرمال‌سازی در این پژوهش شامل یکسان‌سازی نویسه‌های فارسی و عربی، حذف یا همسان‌سازی نویسه‌های نامرئی، تبدیل فاصله‌های متنوع به فاصله معمول و کوچک‌سازی متن برای مقایسه بود. متن اصلی رکوردها برای استفاده مدل و نمایش نمونه‌ها حفظ شد تا داده واقعی corpus تغییر محتوایی پیدا نکند."),
        ("۳-۲-۶. پرسش‌ها و پاسخ‌های تکراری", "تکرار در داده‌های پرسش‌وپاسخ دو شکل مهم دارد: پرسش‌های تکراری یا بسیار نزدیک و پاسخ‌های تکراری که برای چند پرسش متفاوت استفاده شده‌اند. وجود پرسش تکراری در داده آموزش باعث می‌شود برخی موضوع‌ها وزن بیشتری از سایر موضوع‌ها پیدا کنند. مهم‌تر از آن، وجود یک پرسش یا پاسخ تکراری در دو بخش آموزش و آزمون می‌تواند نشت داده ایجاد کند. پاسخ تکراری برای ارزیابی بازیابی نیز مسئله‌ساز است؛ مدل ممکن است پاسخ درست از نظر محتوا را بازیابی کند، اما به‌دلیل متفاوت‌بودن شناسه سند، رتبه آن به اشتباه ناموفق ثبت شود. به همین علت، در ساخت corpus آزمون، فقط نخستین رخداد پاسخ تکراری حفظ شد تا هر پرسش یک پاسخ مرتبط یکتا داشته باشد."),
        ("۳-۲-۷. پاسخ‌های خالی، ناقص یا نامتعارف", "پاسخ خالی یا پاسخ بسیار ناقص برای آموزش و ارزیابی مفید نیست. پاسخ خالی امکان ساخت نمونه مثبت را از بین می‌برد و پاسخ بسیار کوتاه ممکن است فاقد زمینه لازم برای پاسخ‌گویی باشد. از سوی دیگر، پاسخ‌های بسیار بلند نیز با وجود ظرفیت طولانی مدل BGE-M3، هزینه پردازش و آموزش را افزایش می‌دهند و ممکن است چندین موضوع یا نقل‌قول طولانی را در خود داشته باشند. طول پاسخ‌ها با tokenizer مدل BAAI/bge-m3 اندازه‌گیری شد، نه با شمارش کاراکتر. پاسخ‌های دارای بیش از ۱۰۲۴ token از مجموعه نهایی کنار گذاشته شدند. این مقدار یک تصمیم طراحی برای کنترل هزینه و همگن‌سازی داده است، نه محدودیت ذاتی BGE-M3."),
        ("۳-۲-۸. تعادل میان کیفیت داده و حفظ تنوع", "پالایش بیش‌ازحد داده ممکن است تنوع زبانی و موضوعی corpus را کاهش دهد. برای نمونه، حذف همه پرسش‌های طولانی می‌تواند بخشی از مسائل واقعی کاربران را کنار بگذارد. در مقابل، پالایش ضعیف باعث ورود داده‌های خراب، تکراری و غیرپرسشی به آموزش می‌شود. بنابراین، فیلترهای این پژوهش به‌گونه‌ای طراحی شدند که بر نشانه‌های نسبتاً روشن خطای ساختاری تمرکز کنند، نه بر حذف صرف متن‌های بلند یا غیررسمی. پس از فیلتر اولیه کیفیت پرسش‌ها، ۱۷٬۳۵۱ رکورد باقی ماند. مراحل بعدی شامل حذف تکرارها، کنترل طول پاسخ با tokenizer، ایجاد شناسه پایدار و تفکیک آموزش و آزمون، کیفیت داده نهایی را بیشتر افزایش دادند."),
        ("۳-۲-۹. ملاحظات کاربردی و محدودیت منبع", "پرسمان یک corpus واقعی پرسش‌وپاسخ است، نه مجموعه‌داده‌ای که از ابتدا برای benchmark بازیابی طراحی شده باشد. این ویژگی مزیت مهمی دارد، زیرا مدل با زبان واقعی کاربران و پاسخ‌های واقعی مواجه می‌شود. اما محدودیت‌هایی نیز ایجاد می‌کند: همه پاسخ‌ها لزوماً از نظر سبک یا سطح جزئیات همسان نیستند، ممکن است برای یک پرسش چند پاسخ قابل قبول وجود داشته باشد و پاسخ ثبت‌شده همیشه تنها پاسخ ممکن نباشد. بنابراین، معیارهای خودکار پژوهش باید با احتیاط تفسیر شوند. Recall@k و MRR@10 نشان می‌دهند مدل تا چه حد پاسخ مرجع ثبت‌شده را بازیابی می‌کند، اما به‌تنهایی تمام پاسخ‌های ممکن و قابل قبول را پوشش نمی‌دهند."),
        ("۳-۲-۱۰. جمع‌بندی", "مجموعه پرسمان به‌دلیل فارسی‌بودن، واقع‌گرایی پرسش‌ها، تنوع موضوعی و داشتن پیوند پرسش و پاسخ، منبع مناسبی برای پژوهش بازیابی معنایی است. بااین‌حال، داده خام دارای چالش‌هایی مانند متن غیرپرسشی در ستون پرسش، پرسش‌های مقاله‌مانند، تکرار پرسش و پاسخ، تفاوت‌های نگارشی فارسی، پاسخ‌های خالی یا بسیار بلند و خطر نشت داده بود. به همین دلیل، پالایش داده یک مرحله جانبی نیست، بلکه بخشی از روش پژوهش محسوب می‌شود. خروجی این مرحله corpus کنترل‌شده‌ای است که امکان ساخت داده آموزش، انتخاب نمونه‌های منفی سخت و ارزیابی منصفانه مدل پایه و مدل ریزتنظیم‌شده را فراهم می‌کند."),
    )
    for title, description in data_quality_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۳-۳. فرایند پالایش و آماده‌سازی داده‌ها", size=15, bold=True)
    for paragraph_text in (
        "پس از بررسی ویژگی‌ها و چالش‌های داده خام پرسمان، لازم بود مجموعه‌داده پیش از ورود به مرحله آموزش و ارزیابی پالایش شود. هدف از پالایش، حذف ساده رکوردها یا کاهش حجم داده نبود؛ بلکه ایجاد مجموعه‌ای قابل اعتماد بود که در آن هر رکورد تا حد امکان یک رابطه معتبر میان پرسش و پاسخ را نمایش دهد. در بازیابی معنایی، کیفیت این رابطه اهمیت مستقیم دارد؛ زیرا مدل از روی همین زوج‌های پرسش و پاسخ یاد می‌گیرد که چه متن‌هایی باید در فضای برداری به یکدیگر نزدیک و چه متن‌هایی باید از هم دور شوند.",
        "فرایند آماده‌سازی داده در این پژوهش به‌صورت مرحله‌ای انجام شد. در هر مرحله، نوع مشخصی از خطا یا ناهمگونی داده کنترل شد و تعداد رکوردهای باقی‌مانده پیش و پس از اعمال فیلتر ثبت گردید. این ثبت مرحله‌ای فرایند پژوهش را قابل بازتولید می‌کند و مشخص می‌سازد که هر تصمیم پالایشی چه اثری بر اندازه و کیفیت مجموعه نهایی گذاشته است. اجرای مراحل با اسکریپت و notebook انجام شد تا انتخاب رکوردها بر مبنای قواعد ثابت باشد، نه بررسی موردی و غیرقابل تکرار.",
    ):
        add_text(document, paragraph_text)

    preprocessing_sections = (
        ("۳-۳-۱. یکسان‌سازی و نرمال‌سازی متن", "پیش از اعمال فیلترهای محتوایی، متن پرسش‌ها و پاسخ‌ها نرمال‌سازی شد. نرمال‌سازی برای مقایسه و تشخیص تکرار انجام شد و متن اصلی برای نگهداری در داده نهایی حفظ گردید. در زبان فارسی، تفاوت‌های ظاهراً کوچک در نویسه‌ها می‌توانند باعث شوند دو متن هم‌معنا یا حتی کاملاً یکسان، به‌صورت دو رشته متفاوت دیده شوند. در این مرحله، شکل‌های عربی و فارسی حروف ی و ک همسان‌سازی شدند. نویسه‌های نامرئی، نشانه‌های کنترل، فاصله‌های زائد، فاصله بدون‌عرض و کاراکترهای مشابه نیز مدیریت شدند. سپس فاصله‌های تکراری به یک فاصله معمولی تبدیل و ابتدا و انتهای متن پاک‌سازی شد. برای مقایسه دقیق‌تر، متن نرمال‌شده به شکل کوچک‌شده نیز در اختیار فرایند تشخیص تکرار قرار گرفت. این اقدام به‌ویژه برای پرسش‌هایی اهمیت دارد که تنها در نیم‌فاصله، شکل حروف، فاصله‌گذاری یا علامت نگارشی تفاوت دارند. بدون نرمال‌سازی، چنین رکوردهایی ممکن است به اشتباه مستقل تلقی شوند و در بخش‌های متفاوت آموزش و آزمون قرار گیرند."),
        ("۳-۳-۲. کنترل وجود ساختار پرسش", "مجموعه مورد استفاده برای بازیابی پاسخ، باید از پرسش‌های واقعی تشکیل شود. بنابراین، نخستین فیلتر بررسی کرد که متن ستون پرسش با علامت سؤال فارسی یا لاتین پایان یافته باشد. این شرط ساده، بخشی از متن‌های توضیحی، عنوان‌ها، مقاله‌ها، قطعه‌های پاسخ و مطالب ناقص را که به اشتباه در ستون پرسش قرار گرفته بودند، حذف کرد. وجود علامت سؤال به‌تنهایی تضمین نمی‌کند که متن از هر نظر یک پرسش مناسب است؛ بااین‌حال، یک شاخص اولیه و قابل تکرار برای جداسازی بسیاری از رکوردهای غیرپرسشی است. این قاعده همچنین با هدف پروژه هم‌راستا است، زیرا در کاربرد نهایی، ورودی مدل معمولاً پرسش کاربر است و مدل باید پاسخ مناسب را از corpus بازیابی کند."),
        ("۳-۳-۳. حذف پرسش‌های نامعتبر یا بسیار بلند", "پرسش‌هایی با طول نرمال‌شده ۱۲۰۰ کاراکتر یا بیشتر حذف شدند. این آستانه برای جلوگیری از ورود متن‌هایی انتخاب شد که غالباً به یک مقاله، متن چندبخشی، نقل‌قول طولانی یا شرحی فراتر از یک پرسش معمول شباهت داشتند. چنین متن‌هایی معمولاً نیاز اطلاعاتی واحد و روشنی ندارند و مناسب ارزیابی retrieval نیستند. علاوه بر طول، نشانه‌های مقاله‌مانند نیز بررسی شدند. عبارت‌هایی مانند در این مقاله، در این نوشتار، ادامه دارد، پی‌نوشت، کلیدواژه و فهرست منابع می‌توانند نشان دهند که متن موجود در ستون پرسش بخشی از یک مطلب ساختاریافته است. همچنین تعداد زیاد ارجاع‌های داخل کروشه، فهرست‌های شماره‌دار و چند پاراگراف مستقل به‌عنوان نشانه‌های کمکی در نظر گرفته شدند. برای جلوگیری از حذف پرسش‌های واقعی و مفصل، وجود تنها یک نشانه برای حذف کافی نبود. مگر در حالت پرسش‌های بسیار بلند، رکورد زمانی مقاله‌مانند تلقی شد که چند نشانه به‌طور هم‌زمان وجود داشته باشد. این رویکرد میان حفظ تنوع طبیعی پرسش‌ها و حذف خطاهای ساختاری تعادل ایجاد می‌کند."),
        ("۳-۳-۴. کنترل برابری پرسش و پاسخ", "در برخی رکوردها، متن پرسش پس از نرمال‌سازی با متن پاسخ برابر بود. این حالت می‌تواند حاصل خطای ورود داده، درج پاسخ در ستون پرسش یا تکرار ناخواسته متن باشد. نگه‌داشتن چنین نمونه‌ای در آموزش، یک رابطه مصنوعی ایجاد می‌کند؛ زیرا مدل به‌جای یادگیری تطبیق پرسش با پاسخ، به یکسانی مستقیم دو متن پاداش می‌دهد. ازاین‌رو، پرسش هر رکورد با پاسخ متناظر و نیز پاسخ‌های موجود در مجموعه مقایسه شد. رکوردهایی که پرسش آن‌ها با یک پاسخ نرمال‌شده یکسان بود، از مجموعه کنار گذاشته شدند. این فیلتر به کاهش نمونه‌های ناسالم و جلوگیری از ساده‌شدن غیرواقعی مسئله آموزش کمک می‌کند."),
        ("۳-۳-۵. کنترل طول پاسخ با توکنایزر BGE-M3", "طول پاسخ‌ها بر اساس تعداد کاراکتر محاسبه نشد، زیرا مدل زبان متن را به token پردازش می‌کند و تعداد کاراکتر معیار دقیقی برای هزینه محاسباتی یا طول ورودی مدل نیست. بنابراین، پاسخ‌های باقی‌مانده با توکنایزر مدل BAAI/bge-m3 پردازش شدند و تعداد token هر پاسخ، شامل tokenهای ویژه، اندازه‌گیری شد. در داده نهایی، پاسخ‌هایی که بیش از ۱۰۲۴ token داشتند حذف شدند. هدف از این تصمیم، همگن‌سازی طول نمونه‌ها، جلوگیری از هزینه نامتناسب در آموزش و اطمینان از سازگاری داده با تنظیم طول انتخاب‌شده برای آزمایش‌ها بود. این آستانه به این معنا نیست که BGE-M3 ذاتاً قادر به پردازش متن بلندتر نیست؛ بلکه یک تصمیم طراحی در این پژوهش برای ساخت corpus کنترل‌شده است. بررسی آماری طول پاسخ‌ها پیش از حذف نیز اهمیت دارد. این بررسی نشان می‌دهد چه سهمی از پاسخ‌ها در آستانه انتخاب‌شده قرار می‌گیرند، طول میانه و صدک‌های بالای توزیع چگونه‌اند و آیا تصمیم انتخاب‌شده باعث حذف بخش نامتناسبی از داده می‌شود یا خیر."),
        ("۳-۳-۶. ایجاد شناسه پایدار برای رکوردها", "برای هر رکورد باقی‌مانده یک شناسه یکتا و پایدار ایجاد شد. این شناسه، پرسش، پاسخ مثبت، پاسخ‌های منفی، عضویت در داده آموزش یا آزمون و نتایج ارزیابی را در مراحل مختلف به یکدیگر متصل می‌کند. استفاده از شناسه پایدار، وابستگی فرایند پژوهش به شماره سطر فایل CSV را کاهش می‌دهد؛ زیرا شماره سطر پس از اعمال فیلترها، مرتب‌سازی یا ذخیره‌سازی مجدد ممکن است تغییر کند. شناسه‌ها در داده نهایی به‌صورت قابل خواندن تولید شدند و برای نمونه، الگوی porseman-... را دنبال می‌کنند. این تصمیم باعث می‌شود نمونه‌های مشکل‌دار، رتبه پاسخ درست و خروجی مدل‌ها بعداً قابل ردیابی باشند. به‌ویژه در مرحله ارزیابی، برای هر پرسش باید مشخص باشد که پاسخ مرجع دقیق آن کدام رکورد از corpus است."),
        ("۳-۳-۷. حذف تکرارها و جلوگیری از نشت داده", "پس از پالایش اولیه، پرسش‌های تکراری حذف شدند و تنها یک نمونه از هر پرسش نرمال‌شده نگه داشته شد. این کار از افزایش مصنوعی وزن برخی موضوع‌ها در آموزش جلوگیری می‌کند. اگر یک پرسش مشابه چندین بار در داده ظاهر شود، مدل ممکن است بیشتر از آنچه واقعاً لازم است روی همان الگوی زبانی تمرکز کند. پاسخ‌های تکراری نیز در ساخت مجموعه آزمون کنترل شدند. دلیل این موضوع آن است که در ارزیابی بازیابی، هر پرسش باید یک پاسخ مرجع یکتا در corpus داشته باشد. اگر یک پاسخ با شناسه‌های متفاوت چند بار وجود داشته باشد، ممکن است مدل پاسخ درست را بازیابی کند، اما چون شناسه بازیابی‌شده با شناسه مرجع برابر نیست، نتیجه به اشتباه ناموفق ثبت شود. همچنین رکوردهای آموزش و آزمون به‌گونه‌ای ساخته شدند که پاسخ یکسان به‌طور هم‌زمان در هر دو بخش ظاهر نشود. این اصل از نشت داده جلوگیری می‌کند. نشت داده می‌تواند معیارهای ارزیابی را به‌صورت غیرواقعی افزایش دهد، زیرا مدل ممکن است در آموزش با همان پاسخ یا نمونه‌ای بسیار نزدیک مواجه شده باشد."),
        ("۳-۳-۸. خروجی پالایش و آمادگی برای مراحل بعدی", "پس از اعمال فیلترهای اولیه کیفیت پرسش، تعداد ۱۷٬۳۵۱ رکورد باقی ماند. سپس کنترل طول پاسخ بر پایه توکنایزر، حذف تکرارها، ساخت شناسه‌های یکتا و تفکیک داده انجام شد. خروجی این فرایند، داده‌ای است که برای سه کاربرد اصلی آماده شده است: ساخت داده آموزش، ساخت مجموعه آزمون مستقل و تولید corpus پاسخ‌ها برای ارزیابی retrieval. این corpus نهایی، مبنای تمام مراحل بعدی پژوهش است. ابتدا عملکرد مدل پایه BGE-M3 روی داده آزمون اندازه‌گیری می‌شود. سپس از داده آموزش برای ساخت نمونه‌های مثبت و منفی سخت و ریزتنظیم مدل استفاده می‌شود. در نهایت، همان داده آزمون، corpus و معیارهای ارزیابی برای مقایسه منصفانه مدل پایه، مدل ریزتنظیم‌شده و مدل‌های مقایسه‌ای به کار می‌روند."),
    )
    for title, description in preprocessing_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۳-۴. تفکیک داده‌های آموزش و آزمون", size=15, bold=True)
    for paragraph_text in (
        "پس از پالایش و آماده‌سازی داده‌ها، مجموعه نهایی به دو بخش آموزش و آزمون تقسیم شد. این تفکیک برای سنجش معتبر عملکرد مدل ضروری است. اگر مدل با همان نمونه‌هایی ارزیابی شود که در آموزش دیده است، مقدار معیارها بیانگر توان مدل در به‌خاطر سپردن داده‌ها خواهد بود، نه توان آن در بازیابی پاسخ برای پرسش‌های جدید. بنابراین، داده آزمون باید تا حد امکان مستقل از داده آموزش باشد.",
        "در این پژوهش، بخش عمده داده برای آموزش و حدود ده درصد برای آزمون کنار گذاشته شد. داده آموزش برای ساخت نمونه‌های مثبت و منفی سخت و سپس ریزتنظیم مدل استفاده می‌شود. داده آزمون در هیچ‌یک از مراحل انتخاب نمونه آموزشی، محاسبه loss، به‌روزرسانی وزن‌های مدل یا انتخاب تنظیمات بر اساس نتیجه نهایی دخالت ندارد. این جداسازی باعث می‌شود که مقایسه مدل پایه و مدل ریزتنظیم‌شده بر یک مجموعه ثابت و مستقل انجام شود.",
    ):
        add_text(document, paragraph_text)

    data_split_sections = (
        ("۳-۴-۱. هدف از داده آزمون مستقل", "وظیفه مدل در این پژوهش، بازیابی پاسخ مرتبط از میان مجموعه‌ای از پاسخ‌ها برای یک پرسش ورودی است. داده آزمون باید این وضعیت را شبیه‌سازی کند. برای هر پرسش آزمون، پاسخ متناظر آن به‌عنوان پاسخ صحیح یا positive document شناخته می‌شود و تمام پاسخ‌های آزمون corpus بازیابی را تشکیل می‌دهند. در زمان ارزیابی، embedding پرسش آزمون با embedding تمام پاسخ‌های corpus مقایسه می‌شود. پاسخ‌ها بر اساس شباهت مرتب می‌شوند و رتبه پاسخ صحیح به‌دست می‌آید. سپس معیارهایی مانند Recall@5 و MRR@10 محاسبه می‌شوند. چون هیچ‌یک از پرسش‌ها و پاسخ‌های آزمون در آموزش استفاده نشده‌اند، نتیجه حاصل تصویر واقع‌بینانه‌تری از توان تعمیم مدل ارائه می‌دهد."),
        ("۳-۴-۲. نسبت تفکیک داده", "در بسیاری از پژوهش‌های یادگیری ماشین، بخش کوچکی از داده برای آزمون نهایی کنار گذاشته می‌شود. نسبت دقیق به حجم داده، تنوع موضوعی، هزینه ارزیابی و هدف پژوهش بستگی دارد. در این پروژه، حدود ده درصد از داده پالایش‌شده برای آزمون اختصاص یافت و بخش باقی‌مانده در داده آموزش قرار گرفت. این نسبت دو نیاز را هم‌زمان تأمین می‌کند. از یک سو، داده آموزش باید به‌اندازه‌ای بزرگ باشد که مدل ریزتنظیم‌شده با تنوع زبانی و موضوعی مناسب مواجه شود. از سوی دیگر، داده آزمون باید به‌اندازه‌ای باشد که معیارهای بازیابی تحت تأثیر چند نمونه خاص قرار نگیرند. با توجه به تعداد نهایی نمونه‌ها، مجموعه آزمون شامل ۱٬۶۰۱ پرسش و ۱٬۶۰۱ پاسخ یکتاست. این اندازه برای محاسبه معیارهای رتبه‌بندی و مقایسه چند مدل مناسب است."),
        ("۳-۴-۳. نقش شناسه یکتا در تفکیک", "هر رکورد داده نهایی دارای شناسه پایدار است. هنگام تقسیم داده، این شناسه تعیین می‌کند که هر رکورد به کدام بخش تعلق دارد. نگه‌داری شناسه در فایل‌های خروجی ضروری است، زیرا در مراحل بعد باید بتوان ارتباط میان پرسش، پاسخ صحیح، پاسخ‌های منفی و رتبه بازیابی‌شده را به‌صورت دقیق دنبال کرد. در داده آزمون، شناسه پرسش و شناسه پاسخ صحیح به‌طور صریح ذخیره می‌شوند. بنابراین، پس از رتبه‌بندی پاسخ‌ها، می‌توان بررسی کرد که آیا پاسخ با شناسه صحیح در میان نتایج برتر ظاهر شده است یا خیر. این روش، وابستگی ارزیابی به مقایسه متن خام را کاهش می‌دهد و خطاهای ناشی از تفاوت‌های نگارشی، فاصله‌گذاری یا نمایش متن را محدود می‌کند."),
        ("۳-۴-۴. جلوگیری از نشت پاسخ‌های تکراری", "استقلال داده آزمون تنها با تقسیم تصادفی سطرها تضمین نمی‌شود. در مجموعه‌های پرسش‌وپاسخ، ممکن است یک پاسخ برای چند پرسش ثبت شده باشد یا پرسش‌های بسیار نزدیک با تفاوت‌های نگارشی وجود داشته باشند. اگر یک پاسخ یکسان در هر دو بخش آموزش و آزمون قرار بگیرد، مدل ممکن است در زمان آموزش مستقیماً با همان متن پاسخ مواجه شده باشد. در این صورت، ارزیابی نمی‌تواند به‌طور کامل توان تعمیم مدل را نشان دهد. برای کنترل این مسئله، پاسخ‌های تکراری در مجموعه آزمون حذف شدند و از هر گروه پاسخ یکسان تنها یک نمونه نگه‌داری شد. همچنین، نمونه‌های آزمون با داده آموزش مقایسه شدند تا پاسخ یکسان به‌طور هم‌زمان در هر دو بخش ظاهر نشود. این تصمیم حجم داده آزمون را کاهش می‌دهد، اما اعتبار ارزیابی را افزایش می‌دهد. وجود پاسخ تکراری، علاوه بر نشت داده، در محاسبه رتبه نیز مشکل ایجاد می‌کند. اگر پاسخ درست با چند شناسه در corpus تکرار شده باشد، مدل ممکن است متن درست را بازیابی کند، اما چون شناسه بازیابی‌شده با شناسه مرجع برابر نیست، ارزیابی به اشتباه آن را ناموفق ثبت کند. یکتا کردن پاسخ‌ها این ابهام را از بین می‌برد."),
        ("۳-۴-۵. تثبیت نمونه آزمون", "فایل آزمون پس از تولید ذخیره و در تمام آزمایش‌ها بدون تغییر استفاده شد. این موضوع برای مقایسه علمی مدل‌ها اهمیت زیادی دارد. مدل پایه BAAI/bge-m3، مدل ریزتنظیم‌شده پژوهش، مدل jinaai/jina-embeddings-v3 و مدل Snowflake/snowflake-arctic-embed-l-v2.0 همگی با همان پرسش‌ها، همان corpus پاسخ‌ها و همان قواعد رتبه‌بندی ارزیابی شدند. اگر برای هر مدل مجموعه آزمون یا corpus متفاوتی استفاده شود، تفاوت معیارها ممکن است ناشی از تفاوت داده باشد، نه تفاوت واقعی مدل‌ها. تثبیت مجموعه آزمون سبب می‌شود هر تغییر در Recall@k یا MRR@10 به‌طور مستقیم به عملکرد مدل نسبت داده شود."),
        ("۳-۴-۶. داده آموزش و نقش آن در ریزتنظیم", "داده آموزش پس از حذف موارد تکراری و کنترل کیفیت، شامل ۱۳٬۰۰۰ رکورد است. هر رکورد آموزشی یک پرسش، یک پاسخ مثبت و مجموعه‌ای از پاسخ‌های منفی را دربرمی‌گیرد. پاسخ مثبت، پاسخ ثبت‌شده و مرتبط با همان پرسش است. پاسخ‌های منفی در مراحل بعدی از میان پاسخ‌های مشابه، اما نادرست، انتخاب می‌شوند. تفکیک داده پیش از تولید نمونه‌های منفی انجام شد. در نتیجه، هنگام بازیابی نامزدهای منفی برای یک پرسش آموزشی، فقط از corpus آموزشی استفاده می‌شود و پاسخ‌های آزمون وارد داده آموزش نمی‌شوند. این تصمیم از انتقال غیرمستقیم اطلاعات آزمون به فرایند ریزتنظیم جلوگیری می‌کند."),
        ("۳-۴-۷. محدودیت‌های تفکیک تصادفی", "تفکیک تصادفی، با وجود کاربرد گسترده، همه چالش‌های ارزیابی را حل نمی‌کند. در داده‌های موضوعی مانند پرسمان، ممکن است دو پرسش با نگارش متفاوت درباره یک مسئله واحد باشند. حتی اگر متن آن‌ها دقیقاً یکسان نباشد، حضور یکی در آموزش و دیگری در آزمون می‌تواند ارزیابی را تا حدی آسان‌تر کند. همچنین، موضوع‌های پرکاربرد ممکن است در هر دو بخش ظاهر شوند. در این پژوهش، حذف تکرار دقیق پرسش و پاسخ و کنترل اشتراک پاسخ، مهم‌ترین اقدام‌ها برای کاهش این ریسک بوده‌اند. بااین‌حال، باید پذیرفت که استقلال معنایی کامل میان آموزش و آزمون در یک corpus واقعی پرسش‌وپاسخ به‌سادگی قابل تضمین نیست. این موضوع یکی از محدودیت‌های پژوهش است و در تفسیر نتایج باید در نظر گرفته شود."),
        ("۳-۴-۸. جمع‌بندی", "تفکیک داده به آموزش و آزمون، پایه ارزیابی معتبر مدل بازیابی است. در این پژوهش، مجموعه آزمون مستقل شامل ۱٬۶۰۱ پرسش و ۱٬۶۰۱ پاسخ یکتا است و در تمام مقایسه‌ها ثابت نگه داشته شد. کنترل تکرار پاسخ‌ها، نگه‌داری شناسه‌های پایدار و جلوگیری از ورود داده آزمون به فرایند تولید داده آموزشی، باعث شد نتایج ارزیابی تا حد امکان منعکس‌کننده توان واقعی مدل‌ها در بازیابی پاسخ برای پرسش‌های دیده‌نشده باشد."),
    )
    for title, description in data_split_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۳-۵. ساخت داده آموزش با نمونه‌های منفی سخت", size=15, bold=True)
    for paragraph_text in (
        "پس از آماده‌سازی داده آموزش، باید هر زوج پرسش و پاسخ به قالبی تبدیل شود که برای ریزتنظیم مدل embedding مناسب باشد. در این قالب، هر نمونه آموزشی شامل سه جزء اصلی است: یک پرسش، یک پاسخ مثبت و چند پاسخ منفی. پاسخ مثبت، پاسخ ثبت‌شده و مرتبط با همان پرسش است. پاسخ منفی، متنی است که نباید برای آن پرسش در رتبه‌های بالا قرار گیرد.",
        "هدف از این ساختار، آموزش مدل برای نزدیک‌کردن بردار پرسش به بردار پاسخ صحیح و دورکردن آن از پاسخ‌های نادرست است. کیفیت پاسخ‌های منفی در این فرایند اهمیت زیادی دارد. اگر پاسخ‌های منفی کاملاً نامرتبط باشند، مدل به‌سرعت تفاوت آن‌ها با پاسخ مثبت را یاد می‌گیرد؛ اما ممکن است در تشخیص پاسخ‌های ظاهراً مشابه و رقیب موفق نباشد. به همین دلیل، در این پژوهش از نمونه‌های منفی سخت استفاده می‌شود.",
    ):
        add_text(document, paragraph_text)

    hard_negative_sections = (
        ("۳-۵-۱. ساختار هر رکورد آموزشی", "هر رکورد آموزشی شامل متن پرسش، پاسخ مثبت، فهرستی از پاسخ‌های منفی و شناسه یکتای رکورد است. برای هر پرسش، یک پاسخ مثبت و هفت پاسخ منفی نهایی نگه‌داری می‌شود. وجود چند پاسخ منفی در یک رکورد به مدل اجازه می‌دهد پاسخ مثبت را هم‌زمان با چند گزینه رقیب مقایسه کند. این ساختار با هدف contrastive learning سازگار است؛ زیرا مدل باید پاسخ صحیح را نسبت به تمام پاسخ‌های منفی موجود در همان نمونه ترجیح دهد. پاسخ مثبت از رابطه اصلی پرسش و پاسخ در داده پالایش‌شده گرفته می‌شود. با توجه به مراحل حذف پاسخ‌های خالی، پاسخ‌های بسیار بلند و داده‌های تکراری، پاسخ مثبت هر رکورد از نظر ساختاری برای آموزش قابل استفاده است."),
        ("۳-۵-۲. مفهوم نمونه منفی سخت", "نمونه منفی سخت، پاسخی است که از نظر واژگانی، موضوعی یا معنایی به پرسش نزدیک است، اما پاسخ صحیح آن محسوب نمی‌شود. برای مثال، پرسشی درباره یک حکم فقهی ممکن است پاسخ‌های متعددی از همان حوزه را بازیابی کند. این پاسخ‌ها ممکن است واژه‌های مشترکی با پرسش داشته باشند و حتی در نگاه اول مرتبط به نظر برسند، اما مسئله دقیق پرسش را پاسخ نمی‌دهند. نمونه‌های منفی سخت نسبت به نمونه‌های منفی تصادفی ارزش آموزشی بیشتری دارند. منفی تصادفی معمولاً از موضوعی کاملاً متفاوت انتخاب می‌شود و مدل با کمترین یادگیری می‌تواند آن را از پاسخ مثبت جدا کند. در مقابل، منفی سخت مدل را مجبور می‌کند تفاوت‌های ظریف میان موضوع‌های نزدیک، قیود پرسش، نوع نیاز اطلاعاتی و پاسخ دقیق را یاد بگیرد. بااین‌حال، سخت‌بودن یک نمونه به‌تنهایی کافی نیست. اگر یک پاسخ واقعاً مرتبط یا قابل‌قبول باشد، واردکردن آن به‌عنوان منفی به مدل سیگنال نادرست می‌دهد. این وضعیت منفی کاذب نام دارد و می‌تواند کیفیت ریزتنظیم را کاهش دهد. بنابراین، انتخاب نمونه منفی سخت باید با کنترل کیفی همراه باشد."),
        ("۳-۵-۳. بازیابی نامزدهای اولیه با BGE-M3", "برای یافتن پاسخ‌های منفی سخت، ابتدا مدل پایه BAAI/bge-m3 برای پرسش‌ها و پاسخ‌های داده آموزش embedding تولید می‌کند. سپس شباهت هر پرسش با پاسخ‌های corpus آموزشی محاسبه می‌شود. پاسخ‌ها بر اساس شباهت مرتب می‌شوند و برای هر پرسش، ۱۰۰ پاسخ نزدیک‌تر به‌عنوان نامزد اولیه انتخاب می‌گردند. استفاده از ۱۰۰ نامزد به این دلیل است که مرحله نخست بازیابی باید فضای نسبتاً بزرگی از پاسخ‌های مشابه را پوشش دهد. اگر فقط تعداد کمی پاسخ بازیابی شود، ممکن است پاسخ‌های منفی مفید از ابتدا وارد فرایند نشوند. در عین حال، استفاده از همه پاسخ‌های corpus برای بازبینی دقیق هزینه محاسباتی بالایی دارد. بنابراین، بازیابی اولیه با مدل embedding یک مرحله سریع برای محدودکردن فضای جست‌وجو است. پاسخ مثبت همان پرسش از میان نامزدها حذف می‌شود. همچنین، پاسخ‌های دارای شناسه یکسان یا متن تکراری با پاسخ مثبت نباید به‌عنوان منفی وارد مراحل بعدی شوند. این کنترل از ایجاد برچسب نادرست جلوگیری می‌کند."),
        ("۳-۵-۴. بازرتبه‌بندی نامزدها", "مدل embedding برای جست‌وجوی سریع در corpus بزرگ مناسب است، اما معمولاً پرسش و پاسخ را جداگانه نمایش می‌دهد. در نتیجه، ممکن است جزئیات دقیق رابطه میان یک پرسش مشخص و یک پاسخ مشخص را به‌طور کامل تشخیص ندهد. برای افزایش دقت انتخاب نامزدهای منفی، ۱۰۰ پاسخ بازیابی‌شده وارد مرحله بازرتبه‌بندی می‌شوند. در بازرتبه‌بندی، یک مدل reranker پرسش و پاسخ را به‌صورت یک زوج دریافت می‌کند و میزان ارتباط آن‌ها را با دقت بیشتری می‌سنجد. خروجی این مرحله، ترتیب جدیدی از نامزدها است که پاسخ‌های نزدیک‌تر و مسئله‌سازتر را در رتبه‌های بالاتر قرار می‌دهد. از میان این فهرست، ۲۰ پاسخ اول، به‌استثنای پاسخ مثبت، برای بررسی و انتخاب نهایی نگه‌داری می‌شوند. تعداد ۲۰ نامزد برتر میان دو نیاز تعادل ایجاد می‌کند: از یک سو، تعداد کافی پاسخ نزدیک برای انتخاب منفی‌های مناسب فراهم می‌شود؛ از سوی دیگر، حجم بررسی کیفی به اندازه‌ای باقی می‌ماند که بتوان آن را با هزینه قابل‌قبول انجام داد."),
        ("۳-۵-۵. انتخاب هفت پاسخ منفی نهایی", "پس از حذف پاسخ صحیح و کنترل تکرارها، هفت پاسخ نخست از میان نامزدهای بازرتبه‌بندی‌شده به‌عنوان منفی‌های پیش‌فرض انتخاب می‌شوند. انتخاب پاسخ‌های رتبه‌بالا باعث می‌شود منفی‌ها از نظر معنایی به پرسش نزدیک باشند و مسئله آموزشی برای مدل ساده نباشد. با وجود این، ترتیب reranker به‌تنهایی معیار قطعی در نظر گرفته نمی‌شود. ممکن است یک پاسخ رتبه‌بالا در واقع پاسخی مکمل، پاسخی قابل‌قبول با نگاه متفاوت یا پاسخی نزدیک به پاسخ صحیح باشد. بنابراین، نامزدهای منتخب باید پیش از ثبت نهایی از نظر احتمال منفی کاذب بررسی شوند. خروجی این مرحله، برای هر پرسش، یک پاسخ مثبت و هفت پاسخ منفی است. این داده در قالب استاندارد ذخیره می‌شود تا بتوان آن را مستقیماً در فرایند ریزتنظیم مدل استفاده کرد."),
        ("۳-۵-۶. کنترل کیفی با مدل زبانی بزرگ", "در مرحله کنترل کیفیت، یک مدل زبانی بزرگ با دریافت پرسش، پاسخ مثبت و نامزد منفی بررسی می‌کند که آیا نامزد واقعاً پاسخ نادرست یا نامرتبط برای آن پرسش است یا خیر. وظیفه مدل زبانی در اینجا تولید پاسخ جدید نیست؛ بلکه نقش داور یا برچسب‌زن کیفیت را دارد. برای کاهش ابهام، دستور ارزیابی باید روشن باشد. مدل زبانی باید تشخیص دهد که آیا پاسخ منفی، پاسخ مستقیم و صحیح پرسش است، پاسخ مکمل یا قابل‌قبول پرسش است، به موضوعی نزدیک اما متفاوت مربوط است، یا با پرسش نامرتبط است. نامزدی که پاسخ صحیح یا قابل‌قبول تلقی شود، نباید به‌عنوان منفی استفاده شود. در مقابل، پاسخی که به موضوع نزدیک است اما نیاز اطلاعاتی پرسش را برآورده نمی‌کند، منفی سخت مناسبی خواهد بود. نگه‌داری دلیل کوتاه تصمیم یا برچسب کنترل کیفی برای هر نامزد، امکان بررسی بعدی و تحلیل خطا را فراهم می‌کند."),
        ("۳-۵-۷. خطرهای روش انتخاب منفی", "مهم‌ترین خطر این روش، وجود منفی کاذب است. در داده‌های دینی، اخلاقی و مشاوره‌ای، ممکن است یک پرسش بیش از یک پاسخ قابل‌قبول داشته باشد. برای نمونه، دو پاسخ ممکن است هر دو درباره یک مفهوم صحیح باشند، اما یکی توضیح کوتاه و دیگری توضیح تفصیلی ارائه کند. اگر یکی از آن‌ها مثبت و دیگری منفی برچسب بخورد، مدل با سیگنال متناقض مواجه می‌شود. خطر دوم، سوگیری مدل پایه است. نامزدهای اولیه با BGE-M3 انتخاب می‌شوند؛ بنابراین، پاسخ‌هایی که مدل پایه آن‌ها را بازیابی نمی‌کند، وارد مرحله بازرتبه‌بندی هم نخواهند شد. این مسئله طبیعی است، اما باید در تحلیل روش ذکر شود. استفاده از تعداد نسبتاً زیاد نامزد اولیه و سپس reranking، اثر این محدودیت را کاهش می‌دهد. خطر سوم، خطای داوری مدل زبانی بزرگ است. مدل زبانی ممکن است پاسخ را ناقص تفسیر کند یا به‌دلیل شباهت سطحی، یک پاسخ مرتبط را نادرست تشخیص دهد. به همین علت، خروجی کنترل کیفی باید قابل بازبینی ذخیره شود و در صورت امکان، نمونه‌ای از آن با ارزیابی انسانی بررسی گردد."),
        ("۳-۵-۸. مزیت روش چندمرحله‌ای", "روش پیشنهادی از سه لایه تشکیل شده است: بازیابی سریع با مدل embedding، بازرتبه‌بندی دقیق‌تر با reranker و کنترل کیفیت معنایی با مدل زبانی بزرگ. هر مرحله نقشی متفاوت دارد. BGE-M3 فضای بزرگ پاسخ‌ها را با سرعت مناسب محدود می‌کند؛ reranker پاسخ‌های نزدیک‌تر را با دقت بیشتری مرتب می‌سازد؛ و مدل زبانی بزرگ احتمال ورود منفی کاذب را کاهش می‌دهد. این فرایند چندمرحله‌ای، نسبت به انتخاب صرفاً تصادفی پاسخ‌های منفی، به داده‌ای آموزنده‌تر برای ریزتنظیم منجر می‌شود. انتظار می‌رود مدلی که با این نمونه‌ها آموزش می‌بیند، نه‌تنها پاسخ‌های کاملاً نامرتبط، بلکه پاسخ‌های نزدیک اما نادرست را نیز بهتر تشخیص دهد. این ویژگی، برای کاربردهای بازیابی معنایی و سامانه‌های RAG اهمیت ویژه‌ای دارد؛ زیرا در این سامانه‌ها، پاسخ‌های ظاهراً مرتبط اما نادرست می‌توانند به مرحله تولید پاسخ وارد شوند و کیفیت پاسخ نهایی را کاهش دهند."),
        ("۳-۵-۹. جمع‌بندی", "در این پژوهش، برای هر رکورد آموزشی یک پاسخ مثبت و هفت پاسخ منفی سخت در نظر گرفته می‌شود. نامزدهای منفی ابتدا از میان ۱۰۰ پاسخ مشابه با BGE-M3 بازیابی می‌شوند، سپس با reranker بازرتبه‌بندی می‌گردند و ۲۰ نامزد برتر برای انتخاب نهایی بررسی می‌شوند. هفت پاسخ نخست پس از حذف پاسخ صحیح، کنترل تکرار و بازبینی کیفیت به‌عنوان منفی‌های نهایی ذخیره می‌شوند. این روش، داده آموزشی را از زوج‌های ساده پرسش و پاسخ به نمونه‌های contrastive غنی‌تر تبدیل می‌کند. در نتیجه، مدل ریزتنظیم‌شده می‌تواند مرز دقیق‌تری میان پاسخ صحیح و پاسخ‌های معنایی نزدیک اما نادرست یاد بگیرد."),
    )
    for title, description in hard_negative_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۳-۶. ریزتنظیم مدل BGE-M3", size=15, bold=True)
    for paragraph_text in (
        "در این پژوهش، مدل BAAI/bge-m3 به‌عنوان مدل پایه انتخاب شد و سپس با داده آموزشی ساخته‌شده از مجموعه پرسمان ریزتنظیم گردید. هدف از ریزتنظیم، تغییر معماری مدل یا آموزش آن از ابتدا نیست؛ بلکه سازگارکردن بازنمایی‌های برداری مدل با زبان فارسی، ساختار پرسش‌وپاسخ پرسمان و نیاز بازیابی پاسخ مرتبط است.",
        "مدل پایه BGE-M3 پیش از ریزتنظیم، دانش عمومی چندزبانه و توانایی مناسب در ساخت embedding دارد. بااین‌حال، واژگان، شیوه طرح پرسش، موضوع‌های دینی و مشاوره‌ای، سبک پاسخ‌ها و الگوهای نگارشی موجود در پرسمان با داده‌های عمومی آموزش اولیه مدل یکسان نیستند. ریزتنظیم باعث می‌شود مدل در فضای برداری خود، پرسش‌های پرسمان را به پاسخ‌های درست نزدیک‌تر و پاسخ‌های مشابه اما نامرتبط را دورتر کند.",
    ):
        add_text(document, paragraph_text)

    add_text(document, "۳-۶-۱. مدل پایه", size=14, bold=True)
    add_text(document, "مدل پایه این پژوهش BAAI/bge-m3 است. BGE-M3 یک مدل embedding چندزبانه مبتنی بر معماری Transformer است که برای وظایف بازیابی، شباهت معنایی و retrieval طراحی شده است. در این پژوهش، از نمایش dense مدل استفاده شد؛ یعنی هر پرسش و هر پاسخ به یک بردار با طول ثابت تبدیل می‌شود و شباهت میان آن‌ها در فضای برداری محاسبه می‌گردد. انتخاب BGE-M3 چند دلیل دارد. نخست، مدل از زبان فارسی پشتیبانی می‌کند و در وظایف بازیابی چندزبانه عملکرد مناسبی دارد. دوم، مدل پایه در ارزیابی اولیه مجموعه آزمون پرسمان نیز عملکرد قابل‌قبولی نشان داد. سوم، استفاده از مدلی که از ابتدا برای embedding و retrieval آموزش دیده است، نسبت به استفاده از مدل‌های عمومی تولید متن، با هدف این پژوهش هم‌راستاتر است.")

    add_text(document, "۳-۶-۲. قالب ورودی آموزش", size=14, bold=True)
    add_text(document, "هر نمونه آموزشی شامل یک پرسش، یک پاسخ مثبت و هفت پاسخ منفی است. پرسش به‌عنوان query وارد مدل می‌شود. پاسخ مثبت همان پاسخ صحیح و متناظر با پرسش است. هفت پاسخ منفی نیز از میان پاسخ‌های معنایی نزدیک اما نادرست انتخاب شده‌اند. در هر گروه آموزشی، مدل باید تشخیص دهد کدام پاسخ نسبت به پرسش بیشترین ارتباط را دارد. از آنجا که پاسخ‌های منفی به‌صورت تصادفی انتخاب نشده‌اند، مسئله آموزشی دشوارتر و در عین حال واقع‌بینانه‌تر است. مدل نمی‌تواند تنها با تشخیص تفاوت موضوع‌های کاملاً نامرتبط موفق شود؛ بلکه باید میان پاسخ‌های نزدیک، پاسخ دقیق را شناسایی کند. تعداد گزینه‌های هر گروه آموزشی برابر با هشت است: یک پاسخ مثبت و هفت پاسخ منفی. این مقدار با پارامتر train_group_size=8 در فرایند آموزش تنظیم شد.")

    add_text(document, "۳-۶-۳. تابع هدف آموزش", size=14, bold=True)
    add_text(document, "برای آموزش مدل از یادگیری تقابلی استفاده می‌شود. فرض شود q_i بردار پرسش، d_i مثبت بردار پاسخ مثبت و d_ij منفی بردار پاسخ‌های منفی باشد. مدل شباهت میان پرسش و هر پاسخ را محاسبه می‌کند. هدف آموزش این است که شباهت q_i با d_i مثبت از شباهت همان پرسش با تمام پاسخ‌های منفی بیشتر شود. تابع loss برای یک نمونه به‌صورت زیر است:")
    add_contrastive_loss_equation(document, "۳-۱")
    add_text(document, "در این رابطه، s تابع شباهت میان embedding پرسش و پاسخ و τ پارامتر دما است. کوچک‌ترشدن loss به این معنا است که مدل پاسخ مثبت را با اطمینان بیشتری نسبت به پاسخ‌های منفی ترجیح می‌دهد. در اجرای این پژوهش، مقدار دما برابر با ۰٫۰۲ تنظیم شد. این پارامتر بر میزان حساسیت مدل به تفاوت نمره پاسخ مثبت و پاسخ‌های منفی اثر می‌گذارد. استفاده از پاسخ‌های منفی سخت در کنار این تابع هدف، باعث می‌شود مدل به تفاوت‌های کوچک‌تر در ارتباط معنایی حساس شود.")

    fine_tuning_sections = (
        ("۳-۶-۴. طول ورودی پرسش و پاسخ", "حداکثر طول ورودی پرسش و پاسخ در آموزش برابر با ۱۰۲۴ token تنظیم شد. این تنظیم با مرحله پالایش داده هم‌راستا است؛ زیرا پاسخ‌های دارای بیش از ۱۰۲۴ token پیش از ساخت داده نهایی حذف شده‌اند. بنابراین، تعیین این مقدار در آموزش نقش محدودسازی تازه برای corpus ندارد، بلکه تضمین می‌کند ورودی‌ها با قالب ثابت و قابل‌کنترل پردازش شوند. برای پرسش‌ها نیز پیش‌تر سقف ۱۲۰۰ کاراکتر اعمال شده بود تا متن‌های مقاله‌مانند و غیرپرسشی وارد داده نشوند. در نتیجه، طول ورودی‌ها در مرحله ریزتنظیم تا حد زیادی همگن است و هزینه پردازش نمونه‌های غیرعادی کاهش می‌یابد."),
        ("۳-۶-۵. تنظیمات اصلی آموزش", "ریزتنظیم مدل با یک epoch انجام شد. انتخاب یک epoch با توجه به اندازه داده آموزشی و هدف پژوهش صورت گرفت: مدل باید با دامنه پرسمان سازگار شود، بدون آنکه با تکرار زیاد داده‌ها دچار بیش‌برازش شود. تعداد گام‌های آموزش ثبت شد و خروجی نهایی به‌همراه checkpointهای میان‌دوره‌ای ذخیره گردید. نرخ یادگیری برابر با 1e-5 در نظر گرفته شد. این مقدار برای ریزتنظیم مدل‌های Transformer معمول است، زیرا نرخ یادگیری بسیار بزرگ می‌تواند دانش عمومی مدل پایه را تخریب کند و نرخ بسیار کوچک نیز ممکن است سازگاری لازم با داده جدید را ایجاد نکند. نسبت warmup برابر با ۰٫۱ تنظیم شد. در ابتدای آموزش، نرخ یادگیری به‌تدریج افزایش می‌یابد تا به مقدار اصلی برسد. این کار، به‌ویژه در ریزتنظیم مدل‌های بزرگ، پایداری به‌روزرسانی وزن‌ها را افزایش می‌دهد و احتمال نوسان شدید loss در گام‌های نخست را کاهش می‌دهد."),
        ("۳-۶-۶. تنظیمات حافظه و پردازش", "آموزش با دقت fp16 انجام شد. استفاده از نیم‌دقت، مصرف حافظه GPU را کاهش می‌دهد و سرعت محاسبات را افزایش می‌دهد؛ در حالی که برای این نوع ریزتنظیم معمولاً کیفیت قابل‌قبولی حفظ می‌شود. اندازه batch برای هر GPU برابر با یک گروه پرسش تنظیم شد. با توجه به وجود یک پاسخ مثبت و هفت پاسخ منفی، هر گروه شامل هشت متن پاسخ علاوه بر پرسش است. برای افزایش اندازه batch مؤثر بدون افزایش مصرف حافظه، از gradient_accumulation_steps=4 استفاده شد. در این روش، گرادیان چند گام محاسبه و سپس وزن‌های مدل به‌روزرسانی می‌شوند. همچنین، gradient checkpointing فعال شد. در این روش، بخشی از مقادیر میانی به‌جای نگه‌داری دائمی در حافظه، هنگام نیاز دوباره محاسبه می‌شوند. این کار هزینه محاسباتی را اندکی افزایش می‌دهد، اما مصرف حافظه را کاهش می‌دهد و امکان آموزش مدل با ورودی‌های بلندتر یا گروه‌های نمونه بزرگ‌تر را فراهم می‌سازد."),
        ("۳-۶-۷. آموزش توزیع‌شده و بازتولیدپذیری", "آموزش با دو فرایند GPU و سازوکار Distributed Data Parallel انجام شد. در این روش، داده آموزشی میان فرایندها تقسیم می‌شود، هر فرایند بخشی از batch را پردازش می‌کند و گرادیان‌ها پیش از به‌روزرسانی وزن‌ها همگام‌سازی می‌شوند. این تنظیم زمان آموزش را کاهش می‌دهد و امکان استفاده از منابع محاسباتی موجود را فراهم می‌سازد. برای بازتولیدپذیری، seed برابر با ۴۲ تنظیم شد. هرچند اجرای دقیق مدل‌های عمیق روی GPU ممکن است به‌دلیل تفاوت‌های سطح پایین در محاسبات کاملاً بیت‌به‌بیت یکسان نباشد، ثبت seed، داده آموزش، نسخه مدل پایه، پارامترها و مسیر خروجی مدل باعث می‌شود فرایند تا حد زیادی قابل تکرار باشد. مدل نهایی در مسیر خروجی اختصاصی ذخیره شد و فایل‌های لازم شامل وزن مدل، پیکربندی، tokenizer و checkpointها نگه‌داری شدند. بنابراین، مدل می‌تواند بدون اجرای دوباره آموزش، برای embedding، ارزیابی و مقایسه با مدل پایه بارگذاری شود."),
        ("۳-۶-۸. کنترل بیش‌برازش", "در ریزتنظیم روی داده دامنه‌ای، بیش‌برازش یک خطر جدی است. مدلی که بیش از حد با داده آموزش سازگار شود، ممکن است روی پرسش‌های دیده‌شده عملکرد خوبی نشان دهد، اما روی پرسش‌های جدید توانایی خود را از دست بدهد. در این پژوهش، چند اقدام برای کاهش این خطر انجام شد: استفاده از یک epoch، نگه‌داشتن داده آزمون مستقل، انتخاب نمونه‌های منفی متنوع و ارزیابی نهایی فقط روی داده آزمون ثابت. نشانه اصلی موفقیت آموزش، کاهش loss به‌تنهایی نیست. loss صرفاً نشان می‌دهد مدل در تفکیک نمونه‌های آموزشی بهتر شده است. معیار اصلی، بهبود Recall@5 و MRR@10 روی داده آزمون دیده‌نشده است. به همین دلیل، نتایج آموزش با مدل پایه و مدل‌های مقایسه‌ای روی همان corpus آزمون تحلیل می‌شوند."),
        ("۳-۶-۹. خروجی مرحله ریزتنظیم", "خروجی این مرحله، نسخه ریزتنظیم‌شده BGE-M3 است که برای embedding پرسش‌ها و پاسخ‌های فارسی پرسمان به کار می‌رود. این مدل همان معماری مدل پایه را حفظ می‌کند، اما فضای برداری آن با زوج‌های پرسش و پاسخ و نمونه‌های منفی سخت داده آموزشی سازگار شده است. در فصل بعد، عملکرد این مدل با مدل پایه BGE-M3، jinaai/jina-embeddings-v3 و Snowflake/snowflake-arctic-embed-l-v2.0 مقایسه می‌شود. این مقایسه نشان می‌دهد که ریزتنظیم دامنه‌ای تا چه حد می‌تواند بازیابی پاسخ در مجموعه پرسمان را بهبود دهد."),
    )
    for title, description in fine_tuning_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۳-۷. پروتکل ارزیابی و مقایسه مدل‌ها", size=15, bold=True)
    for paragraph_text in (
        "پس از آماده‌سازی داده و ریزتنظیم مدل، لازم است عملکرد مدل ریزتنظیم‌شده در شرایطی کنترل‌شده و قابل بازتولید ارزیابی شود. هدف ارزیابی، سنجش توان مدل در بازیابی پاسخ درست برای پرسش‌هایی است که در آموزش مشاهده نشده‌اند. همچنین، برای آن‌که بهبود حاصل از ریزتنظیم به‌درستی تفسیر شود، مدل ریزتنظیم‌شده باید با مدل پایه و مدل‌های مقایسه‌ای روی دقیقاً همان داده ارزیابی شود.",
        "در این پژوهش، تمام مدل‌ها با یک مجموعه آزمون ثابت شامل ۱٬۶۰۱ پرسش و ۱٬۶۰۱ پاسخ یکتا ارزیابی شدند. هر پرسش آزمون یک پاسخ مرجع دارد و تمام پاسخ‌های آزمون corpus بازیابی را تشکیل می‌دهند. بنابراین، برای هر پرسش، مدل باید پاسخ درست را از میان ۱٬۶۰۱ پاسخ ممکن پیدا و رتبه‌بندی کند.",
    ):
        add_text(document, paragraph_text)

    evaluation_sections_before_equations = (
        ("۳-۷-۱. مدل‌های ارزیابی‌شده", "چهار مدل در ارزیابی نهایی حضور دارند: مدل پایه BAAI/bge-m3، مدل BGE-M3 ریزتنظیم‌شده با داده پرسمان، مدل jinaai/jina-embeddings-v3 و مدل Snowflake/snowflake-arctic-embed-l-v2.0. مدل پایه BGE-M3 نشان می‌دهد که پیش از سازگارسازی با داده پرسمان، مدل چندزبانه در مسئله بازیابی پاسخ فارسی چه عملکردی دارد. مدل ریزتنظیم‌شده، اثر مستقیم داده آموزشی و روش انتخاب نمونه‌های منفی را نشان می‌دهد. دو مدل Jina و Snowflake نیز به‌عنوان مدل‌های embedding عمومی و قدرتمند وارد مقایسه شده‌اند تا نتیجه فقط محدود به مقایسه مدل با نسخه پایه خود نباشد. همه مدل‌ها بدون تغییر داده آزمون، پاسخ‌های corpus یا معیارهای رتبه‌بندی اجرا شدند. در نتیجه، هر تفاوت در معیارها به تفاوت در بازنمایی‌های برداری و قابلیت بازیابی مدل‌ها مربوط است، نه تفاوت در داده یا شیوه محاسبه."),
        ("۳-۷-۲. ساخت embedding پرسش‌ها و پاسخ‌ها", "برای هر مدل، ابتدا embedding تمام پاسخ‌های corpus آزمون تولید می‌شود. این مرحله برای هر مدل فقط یک بار انجام می‌شود و بردارهای پاسخ‌ها برای همه پرسش‌های آزمون استفاده می‌گردند. سپس embedding هر پرسش آزمون محاسبه می‌شود. در مدل‌های BGE-M3، پرسش‌ها با حالت encode_queries و پاسخ‌ها با حالت encode_corpus پردازش شدند. روش pooling در این مدل‌ها cls است؛ یعنی بردار token ویژه CLS به‌عنوان بازنمایی کل متن استفاده می‌شود. این تنظیم با شیوه معمول استفاده از BGE-M3 در بازیابی dense سازگار است. برای Jina Embeddings v3 و Snowflake Arctic Embed L v2.0، embeddingها با API و تنظیمات پیشنهادی هر مدل تولید شدند. در مدل Jina، وظیفه بازیابی برای پرسش و passage به‌صورت جداگانه مشخص شد تا مدل بتواند تفاوت نقش query و document را در نظر بگیرد. در مدل Snowflake نیز الگوی ورودی مناسب پرسش برای query و الگوی passage برای پاسخ‌ها استفاده شد."),
        ("۳-۷-۳. نرمال‌سازی embeddingها و شباهت کسینوسی", "پس از تولید embeddingها، بردار همه پرسش‌ها و پاسخ‌ها با نُرم L2 نرمال‌سازی شدند. نرمال‌سازی باعث می‌شود طول بردارها بر امتیاز شباهت اثر نگذارد و مقایسه بر مبنای جهت بردارها انجام شود. پس از نرمال‌سازی، ضرب داخلی دو بردار برابر با شباهت کسینوسی آن‌ها است. در رابطه زیر، q بردار پرسش و d بردار پاسخ است. شباهت بیشتر نشان می‌دهد که مدل، پرسش و پاسخ را از نظر معنایی نزدیک‌تر می‌داند."),
    )
    for title, description in evaluation_sections_before_equations:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)
    add_cosine_similarity_equation(document, "۳-۲")
    add_text(document, "برای هر پرسش، شباهت آن با تمام پاسخ‌های corpus محاسبه شد. سپس پاسخ‌ها بر اساس امتیاز شباهت از بیشترین به کمترین مرتب شدند. رتبه پاسخ مرجع در این فهرست، مبنای محاسبه معیارهای ارزیابی است.")

    evaluation_sections_after_cosine = (
        ("۳-۷-۴. محاسبه رتبه پاسخ صحیح", "هر پرسش آزمون دارای شناسه پاسخ صحیح است. پس از مرتب‌سازی پاسخ‌ها، موقعیت پاسخ دارای شناسه مرجع استخراج می‌شود. اگر پاسخ صحیح در جایگاه نخست قرار گرفته باشد، رتبه آن برابر یک است. اگر پاسخ صحیح ششمین نتیجه باشد، رتبه برابر شش خواهد بود. ذخیره رتبه صحیح برای هر پرسش، فقط برای محاسبه معیار کلی نیست. این اطلاعات برای تحلیل خطا نیز ضروری است. برای نمونه، می‌توان پرسش‌هایی را بررسی کرد که پاسخ صحیح آن‌ها در رتبه دوم تا پنجم قرار گرفته است، یا مواردی را که پاسخ صحیح در ده رتبه اول ظاهر نشده است. چنین تحلیلی نشان می‌دهد ضعف مدل بیشتر در کدام نوع پرسش‌ها یا پاسخ‌ها رخ می‌دهد."),
        ("۳-۷-۵. معیار Recall@1", "معیار Recall@1 بررسی می‌کند که آیا پاسخ صحیح در رتبه اول قرار گرفته است یا خیر. این معیار سخت‌گیرانه است، زیرا تنها بهترین نتیجه بازیابی را در نظر می‌گیرد. در سامانه‌ای که پاسخ رتبه اول مستقیماً به کاربر نمایش داده می‌شود، Recall@1 شاخص مهمی است. برای هر پرسش، اگر رتبه پاسخ درست برابر یک باشد، مقدار آن نمونه یک و در غیر این صورت صفر است. میانگین این مقدار برای تمام پرسش‌های آزمون، Recall@1 را تشکیل می‌دهد. افزایش این معیار نشان می‌دهد که مدل در انتخاب پاسخ درست به‌عنوان بهترین نتیجه موفق‌تر شده است."),
        ("۳-۷-۶. معیار Recall@5", "معیار Recall@5 بررسی می‌کند که آیا پاسخ صحیح در میان پنج نتیجه اول قرار گرفته است یا خیر. این معیار برای کاربردهایی اهمیت دارد که چند پاسخ نخست به کاربر، reranker یا مدل زبانی مولد داده می‌شود. در یک سامانه RAG، اگر سند مرتبط در میان چند نتیجه نخست وجود داشته باشد، مرحله بعدی می‌تواند از آن برای تولید پاسخ بهتر استفاده کند. در این پژوهش، Recall@5 یکی از معیارهای اصلی است. این معیار نشان می‌دهد که مدل تا چه حد توانسته پاسخ درست را به بخش بالای فهرست رتبه‌بندی منتقل کند؛ حتی اگر پاسخ درست در جایگاه اول نباشد."),
        ("۳-۷-۷. معیار MRR@10", "معیار MRR@10 علاوه بر حضور پاسخ صحیح در میان نتایج برتر، به رتبه دقیق آن نیز حساس است. اگر پاسخ صحیح در رتبه اول باشد، سهم آن پرسش برابر یک است. اگر در رتبه دوم باشد، سهم آن برابر یک‌دوم و اگر در رتبه دهم باشد، سهم آن برابر یک‌دهم خواهد بود. پاسخ‌هایی که رتبه‌ای بزرگ‌تر از ده دارند، سهم صفر دریافت می‌کنند. در رابطه زیر، Q مجموعه پرسش‌های آزمون و rank(q) رتبه پاسخ صحیح برای پرسش q است. MRR@10 نسبت به Recall@5 اطلاعات بیشتری ارائه می‌دهد؛ زیرا میان قرارگرفتن پاسخ درست در رتبه اول و رتبه پنجم تفاوت قائل می‌شود."),
    )
    for title, description in evaluation_sections_after_cosine:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)
        if title.startswith("۳-۷-۷"):
            add_mrr_at_10_equation(document, "۳-۳")

    evaluation_sections_final = (
        ("۳-۷-۸. تنظیمات اجرایی و مدیریت حافظه", "برای مدل‌های BGE-M3، اندازه batch تولید embedding برابر با ۱۶ بود. برای مدل‌های Jina و Snowflake، به‌دلیل مصرف بیشتر حافظه GPU، اندازه batch برابر با ۲ در نظر گرفته شد. این تفاوت فقط بر سرعت اجرا اثر دارد و تغییری در مجموعه آزمون، روش نرمال‌سازی یا شیوه محاسبه معیارها ایجاد نمی‌کند. امتیازدهی شباهت‌ها با batch جست‌وجو برابر با ۶۴ پرسش انجام شد. این روش اجازه می‌دهد شباهت هر پرسش با تمام پاسخ‌های corpus محاسبه شود، بدون آن‌که همه ماتریس شباهت به‌طور هم‌زمان در حافظه قرار گیرد. پس از پایان ارزیابی هر مدل، شیء مدل از حافظه حذف شد، garbage collection اجرا شد و حافظه cache کارت گرافیک آزاد گردید. این کار از تجمع حافظه در ارزیابی پیاپی چند مدل جلوگیری می‌کند و باعث می‌شود اجرای مقایسه بدون خطای کمبود حافظه انجام شود."),
        ("۳-۷-۹. ذخیره خروجی‌ها و بازتولیدپذیری", "برای هر مدل، رتبه پاسخ صحیح و مقادیر مربوط به هر پرسش در فایل خروجی ذخیره شدند. این خروجی شامل شناسه پرسش، متن پرسش، رتبه پاسخ درست، مقدار Recall@1، مقدار Recall@5 و reciprocal rank در MRR@10 است. علاوه بر این، یک فایل خلاصه شامل نام مدل، مسیر یا شناسه مدل، Recall@1، Recall@5 و MRR@10 تولید شد. گزارش HTML نیز برای نمایش مقایسه مدل‌ها ساخته شد. نگه‌داری این فایل‌ها باعث می‌شود نتایج بدون اجرای دوباره تمام مدل‌ها قابل مشاهده و تحلیل باشند."),
        ("۳-۷-۱۰. اصول مقایسه منصفانه", "مقایسه مدل‌ها فقط زمانی معتبر است که شرایط ارزیابی ثابت باشد. در این پژوهش، برای همه مدل‌ها موارد زیر یکسان نگه داشته شد: داده آزمون، corpus پاسخ‌ها، شناسه پاسخ صحیح، نرمال‌سازی بردارها، شباهت کسینوسی، روش رتبه‌بندی و معیارهای ارزیابی. تفاوت‌هایی مانند اندازه batch، API هر مدل یا نحوه بارگذاری مدل، صرفاً تنظیمات اجرایی هستند و بر تعریف مسئله ارزیابی اثر نمی‌گذارند. در نتیجه، بهبود مدل ریزتنظیم‌شده نسبت به مدل پایه، Jina و Snowflake را می‌توان به توان آن در ساخت embeddingهای مناسب‌تر برای داده فارسی پرسمان نسبت داد."),
    )
    for title, description in evaluation_sections_final:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۳-۸. جمع‌بندی فصل سوم", size=15, bold=True)
    for paragraph_text in (
        "در این فصل، روش اجرایی پژوهش برای بهبود بازیابی معنایی پاسخ‌های فارسی در مجموعه پرسمان تشریح شد. فرایند پژوهش از داده خام آغاز شد و پس از پالایش، ساخت داده آموزش، ریزتنظیم مدل و ارزیابی منصفانه ادامه یافت. هدف این فصل، ارائه مسیری روشن و قابل بازتولید از داده اولیه تا مدل نهایی بود.",
        "نخست، داده خام پرسمان از نظر ساختار و کیفیت بررسی شد. داده اولیه شامل پرسش‌ها و پاسخ‌های واقعی بود، اما مانند بسیاری از corpusهای وب، دارای رکوردهای غیرپرسشی، متن‌های مقاله‌مانند، تفاوت‌های نگارشی فارسی، پرسش‌ها و پاسخ‌های تکراری، پاسخ‌های خالی و پاسخ‌های بسیار بلند بود. ورود مستقیم چنین داده‌ای به آموزش می‌توانست باعث شود مدل به‌جای یادگیری رابطه طبیعی میان پرسش و پاسخ، الگوهای نادرست و نویزی را یاد بگیرد.",
        "برای کنترل این مسائل، نرمال‌سازی متن، بررسی ساختار پرسش، حذف پرسش‌های بلند یا مقاله‌مانند، کنترل برابری پرسش و پاسخ، اندازه‌گیری طول پاسخ با توکنایزر BGE-M3 و حذف پاسخ‌های بیش از ۱۰۲۴ token انجام شد. همچنین، شناسه‌های یکتا برای رکوردها ایجاد شدند و تکرارها کنترل شدند. این مراحل باعث شد داده نهایی از نظر ساختاری برای آموزش و ارزیابی بازیابی معنایی مناسب‌تر باشد.",
        "در مرحله تفکیک داده، بخش آموزش و آزمون از یکدیگر جدا شدند. مجموعه آزمون ثابت شامل ۱٬۶۰۱ پرسش و ۱٬۶۰۱ پاسخ یکتا است. پاسخ‌های تکراری در corpus آزمون کنترل شدند تا هر پرسش یک پاسخ مرجع یکتا داشته باشد. همچنین، از ورود پاسخ‌های آزمون به فرایند ساخت داده آموزش جلوگیری شد تا نشت داده رخ ندهد. این تصمیم، اعتبار معیارهای ارزیابی و مقایسه مدل‌ها را افزایش می‌دهد.",
        "داده آموزش با قالب contrastive ساخته شد. هر نمونه شامل یک پرسش، یک پاسخ مثبت و هفت پاسخ منفی است. پاسخ‌های منفی به‌صورت تصادفی انتخاب نشدند؛ بلکه ابتدا با مدل پایه BGE-M3، ۱۰۰ پاسخ نزدیک برای هر پرسش بازیابی شد. سپس نامزدها با reranker بازرتبه‌بندی شدند و از میان پاسخ‌های برتر، پس از حذف پاسخ صحیح، کنترل تکرار و بررسی احتمال منفی کاذب، هفت پاسخ منفی نهایی انتخاب شدند.",
        "این روش برای مدل اهمیت زیادی دارد، زیرا پاسخ‌های منفی سخت از نظر موضوعی و معنایی به پرسش نزدیک‌اند، اما پاسخ دقیق آن نیستند. مدل در فرایند ریزتنظیم مجبور می‌شود مرز دقیق‌تری میان پاسخ صحیح و پاسخ‌های مشابه اما نامرتبط بیاموزد. کنترل کیفی با مدل زبانی بزرگ نیز برای کاهش احتمال ورود پاسخ‌های مرتبط به‌عنوان منفی در نظر گرفته شد.",
        "مدل پایه BAAI/bge-m3 با داده آموزش ریزتنظیم شد. آموزش با یک epoch، نرخ یادگیری 1e-5، نسبت warmup برابر با ۰٫۱، دقت fp16 و داده‌های گروهی شامل یک پاسخ مثبت و هفت پاسخ منفی انجام گرفت. استفاده از gradient accumulation و gradient checkpointing امکان اجرای آموزش با مدیریت مناسب حافظه را فراهم کرد. مدل نهایی، همراه با tokenizer، پیکربندی و checkpointهای لازم ذخیره شد تا قابل بارگذاری و تکرار باشد.",
        "در بخش ارزیابی، مدل پایه BGE-M3، مدل ریزتنظیم‌شده، jinaai/jina-embeddings-v3 و Snowflake/snowflake-arctic-embed-l-v2.0 با داده آزمون یکسان مقایسه شدند. برای تمام مدل‌ها، embedding پرسش‌ها و پاسخ‌ها تولید، با نُرم L2 نرمال‌سازی و سپس بر پایه شباهت کسینوسی رتبه‌بندی شد. معیارهای Recall@1، Recall@5 و MRR@10 برای سنجش کیفیت بازیابی به کار رفتند.",
        "ذخیره رتبه پاسخ صحیح برای هر پرسش، فایل‌های معیارهای کلی و گزارش HTML باعث شد ارزیابی کاملاً قابل بررسی باشد. علاوه بر این، مدیریت حافظه GPU پس از پایان اجرای هر مدل انجام شد تا مقایسه مدل‌های مختلف بدون تجمع حافظه و خطای اجرایی انجام شود.",
        "در مجموع، روش ارائه‌شده در این فصل، یک خط لوله کامل برای سازگارسازی مدل embedding چندزبانه با داده فارسی پرسمان فراهم می‌کند. فصل چهارم به ارائه و تحلیل نتایج آزمایش‌ها اختصاص دارد. در آن فصل، عملکرد مدل ریزتنظیم‌شده با مدل پایه و دو مدل مقایسه‌ای بررسی می‌شود و میزان اثر ریزتنظیم دامنه‌ای بر بازیابی پاسخ‌های فارسی تحلیل خواهد شد.",
    ):
        add_text(document, paragraph_text)

    document.add_page_break()
    add_text(document, "فصل چهارم: نتایج آزمایش‌ها و تحلیل", size=18, bold=True)
    add_text(document, "۴-۱. مقدمه", size=15, bold=True)
    for paragraph_text in (
        "این فصل به ارائه و تحلیل نتایج آزمایش‌های انجام‌شده اختصاص دارد. هدف اصلی آزمایش‌ها، پاسخ به این پرسش است که آیا ریزتنظیم مدل BGE-M3 با داده پالایش‌شده پرسمان و نمونه‌های منفی سخت، می‌تواند بازیابی پاسخ‌های فارسی را نسبت به مدل پایه و مدل‌های embedding عمومی بهبود دهد یا خیر.",
        "تمام مدل‌ها بر روی یک مجموعه آزمون ثابت شامل ۱٬۶۰۱ پرسش و ۱٬۶۰۱ پاسخ یکتا ارزیابی شدند. برای هر پرسش، مدل باید پاسخ صحیح را از میان تمام پاسخ‌های corpus بازیابی کند. پاسخ‌ها بر اساس شباهت کسینوسی میان embedding پرسش و embedding پاسخ مرتب شدند و رتبه پاسخ مرجع برای هر پرسش ثبت گردید.",
        "استفاده از مجموعه آزمون ثابت در این فصل اهمیت اساسی دارد. اگر مدل‌ها با corpus یا پرسش‌های متفاوت ارزیابی شوند، تفاوت نتایج را نمی‌توان با اطمینان به کیفیت مدل‌ها نسبت داد. در این پژوهش، داده آزمون، پاسخ‌های corpus، نرمال‌سازی بردارها، روش رتبه‌بندی و معیارهای محاسبه‌شده برای همه مدل‌ها یکسان بوده‌اند. بنابراین، اختلاف عملکرد مشاهده‌شده، بیانگر تفاوت مدل‌ها در بازنمایی معنایی پرسش‌ها و پاسخ‌های فارسی است.",
        "چهار مدل در آزمایش‌ها ارزیابی شدند: مدل پایه BAAI/bge-m3، نسخه ریزتنظیم‌شده همان مدل با داده پرسمان، jinaai/jina-embeddings-v3 و Snowflake/snowflake-arctic-embed-l-v2.0. مقایسه با مدل پایه نشان می‌دهد که ریزتنظیم دامنه‌ای چه اثری داشته است. مقایسه با Jina و Snowflake نیز جایگاه مدل ریزتنظیم‌شده را نسبت به دو مدل embedding عمومی و قدرتمند مشخص می‌کند.",
        "سه معیار Recall@1، Recall@5 و MRR@10 برای ارزیابی استفاده شده‌اند. Recall@1 نشان می‌دهد که پاسخ صحیح در چه درصدی از پرسش‌ها در رتبه اول قرار گرفته است. Recall@5 بررسی می‌کند که پاسخ صحیح چند درصد از پرسش‌ها در میان پنج نتیجه نخست بوده است. MRR@10 نیز با درنظرگرفتن رتبه دقیق پاسخ درست در ده نتیجه اول، تصویری دقیق‌تر از کیفیت رتبه‌بندی ارائه می‌دهد.",
        "تحلیل نتایج فقط به مقایسه مقدارهای کلی معیارها محدود نمی‌شود. افزایش Recall@1 نشان می‌دهد که مدل در انتخاب بهترین پاسخ موفق‌تر است. افزایش Recall@5 بیانگر بهبود توان مدل در رساندن پاسخ صحیح به محدوده نتایج قابل استفاده در سامانه‌های بازیابی و RAG است. افزایش MRR@10 نیز نشان می‌دهد پاسخ صحیح به‌طور میانگین به رتبه‌های بالاتری منتقل شده است.",
        "در ادامه فصل، ابتدا نتایج عددی هر چهار مدل در یک جدول مقایسه‌ای ارائه می‌شود. سپس، اثر ریزتنظیم نسبت به مدل پایه و دو مدل مقایسه‌ای تحلیل می‌گردد. در بخش‌های بعدی، نتایج از منظر کاربرد در بازیابی پاسخ فارسی و سامانه‌های RAG تفسیر شده و محدودیت‌های نتایج نیز بیان می‌شود.",
    ):
        add_text(document, paragraph_text)

    add_text(document, "۴-۲. مقایسه کمی مدل‌ها", size=15, bold=True)
    add_text(document, "نتایج ارزیابی چهار مدل بر روی مجموعه آزمون ثابت پرسمان در جدول ۴-۱ ارائه شده است. در این جدول، مقدار بیشتر برای هر سه معیار نشان‌دهنده عملکرد بهتر است. مدل ریزتنظیم‌شده BGE-M3 در تمام معیارها بهترین نتیجه را کسب کرده است.")
    add_text(document, "جدول ۴-۱. مقایسه عملکرد مدل‌ها در بازیابی پاسخ‌های فارسی پرسمان", size=14, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_results_comparison_table(document)
    for paragraph_text in (
        "نتایج نشان می‌دهد مدل پایه BGE-M3، حتی بدون سازگارسازی با داده پرسمان، عملکرد قدرتمندی در بازیابی پاسخ فارسی دارد. مقدار ۷۴٫۳۳ درصد برای Recall@1 نشان می‌دهد مدل پایه در حدود سه‌چهارم پرسش‌ها، پاسخ صحیح را در رتبه اول قرار داده است. همچنین، Recall@5 برابر با ۹۱٫۹۴ درصد بیانگر آن است که پاسخ صحیح برای بخش بزرگی از پرسش‌ها در میان پنج گزینه نخست بازیابی شده است.",
        "با وجود عملکرد مناسب مدل پایه، ریزتنظیم با داده پرسمان باعث بهبود روشن هر سه معیار شده است. Recall@1 از ۷۴٫۳۳ درصد به ۸۱٫۷۰ درصد افزایش یافته است. این افزایش ۷٫۳۷ واحد درصدی مهم است، زیرا Recall@1 سخت‌گیرانه‌ترین معیار این آزمایش است و فقط زمانی موفقیت را ثبت می‌کند که پاسخ درست دقیقاً در رتبه اول باشد.",
        "Recall@5 نیز از ۹۱٫۹۴ درصد به ۹۴٫۹۴ درصد رسیده است؛ یعنی ۳٫۰۰ واحد درصد افزایش. هرچند مقدار اولیه Recall@5 بالا بوده است، افزایش آن نشان می‌دهد ریزتنظیم توانسته بخشی از پاسخ‌های صحیحی را که پیش‌تر خارج از پنج رتبه اول بوده‌اند، به محدوده قابل استفاده بازیابی منتقل کند. این نتیجه برای کاربردهای RAG اهمیت دارد؛ زیرا در این سامانه‌ها معمولاً چند سند یا پاسخ نخست به مرحله بازرتبه‌بندی یا تولید پاسخ ارسال می‌شود.",
        "مقدار MRR@10 نیز از ۸۱٫۹۷ درصد برای مدل پایه به ۸۷٫۶۷ درصد برای مدل ریزتنظیم‌شده افزایش یافته است. این اختلاف ۵٫۷۰ واحد درصدی نشان می‌دهد بهبود فقط ناشی از ورود پاسخ صحیح به ده نتیجه نخست نیست، بلکه پاسخ صحیح در بسیاری از موارد به رتبه‌های بالاتر منتقل شده است. به‌عبارت دیگر، مدل ریزتنظیم‌شده نه‌تنها پاسخ درست را بیشتر بازیابی می‌کند، بلکه آن را در جایگاه مناسب‌تری قرار می‌دهد.",
        "مدل jinaai/jina-embeddings-v3 نسبت به مدل پایه BGE-M3 عملکرد بهتری دارد. مقدار Recall@1 آن ۷۹٫۳۳ درصد، Recall@5 آن ۹۳٫۰۰ درصد و MRR@10 آن ۸۵٫۴۵ درصد است. بااین‌حال، مدل ریزتنظیم‌شده BGE-M3 در هر سه معیار از Jina بهتر عمل کرده است. اختلاف مدل ریزتنظیم‌شده نسبت به Jina در Recall@1 برابر با ۲٫۳۷ واحد درصد، در Recall@5 برابر با ۱٫۹۴ واحد درصد و در MRR@10 برابر با ۲٫۲۲ واحد درصد است.",
        "مدل Snowflake/snowflake-arctic-embed-l-v2.0 نیز عملکرد رقابتی دارد. Recall@1 آن ۷۹٫۰۱ درصد و Recall@5 آن ۹۳٫۳۲ درصد است. این مدل در Recall@5 اندکی از Jina بهتر است، اما در Recall@1 و MRR@10 مقدار پایین‌تری دارد. مدل BGE-M3 ریزتنظیم‌شده در مقایسه با Snowflake، به‌ترتیب ۲٫۶۹، ۱٫۶۲ و ۲٫۳۶ واحد درصد در Recall@1، Recall@5 و MRR@10 بهبود نشان می‌دهد.",
        "برتری مدل ریزتنظیم‌شده نسبت به هر دو مدل عمومی مقایسه‌ای، بیانگر ارزش ریزتنظیم دامنه‌ای است. مدل‌های Jina و Snowflake با وجود توانایی عمومی مناسب و داده آموزشی گسترده، با سبک پرسش‌ها، واژگان، تنوع نگارشی و الگوی پاسخ‌های فارسی پرسمان به‌طور اختصاصی آموزش ندیده‌اند. در مقابل، BGE-M3 ریزتنظیم‌شده با زوج‌های پرسش و پاسخ همین دامنه و نمونه‌های منفی سخت مواجه شده است. بنابراین، مدل توانسته تفاوت میان پاسخ دقیق و پاسخ‌های ظاهراً مشابه را بهتر یاد بگیرد.",
        "این نتایج از فرض اصلی پژوهش پشتیبانی می‌کنند: یک مدل embedding چندزبانه مناسب، با داده پاک‌سازی‌شده و ریزتنظیم هدفمند روی corpus فارسی دامنه‌ای، می‌تواند از مدل پایه و مدل‌های عمومی قدرتمند در بازیابی پاسخ‌های همان دامنه بهتر عمل کند.",
    ):
        add_text(document, paragraph_text)

    add_text(document, "۴-۳. تحلیل کیفی نمونه‌های بهبود‌یافته", size=15, bold=True)
    add_text(document, "ارقام کلی Recall@1، Recall@5 و MRR@10 نشان می‌دهند که مدل ریزتنظیم‌شده عملکرد بهتری دارد؛ اما تحلیل موردی روشن می‌کند این بهبود دقیقاً در چه نوع پرسش‌هایی رخ داده است. در این بخش، دو پرسش انتخاب شده‌اند که پاسخ صحیح آن‌ها در مدل پایه در رتبه‌های بسیار پایین قرار داشت، اما پس از ریزتنظیم به رتبه اول منتقل شد. این نمونه‌ها نشان می‌دهند ریزتنظیم، تنها شباهت واژگانی را افزایش نداده، بلکه به تشخیص قیدها و نیاز اطلاعاتی دقیق پرسش کمک کرده است.")

    add_text(document, "۴-۳-۱. درک پرسش محاوره‌ای و چندقیدی درباره پوشش و ورزش", size=14, bold=True)
    for paragraph_text in (
        "پرسش با شناسه porseman-7005 درباره این است که آیا کوتاه‌کردن موی دخترانه به شکل پسرانه، انجام ورزش سنگین یا اسب‌سواری، مصداق تشبّه زن به مرد و مشمول نهی روایی است یا خیر. پاسخ مرجع با یک حکم شرطی بیان می‌کند که با رعایت موازین شرعی، قرارنگرفتن در معرض دید نامحرم و پرهیز از آمیختگی با نامحرم، انجام این امور اشکال ندارد.",
        "مدل پایه BGE-M3 پاسخ مرجع را در رتبه ۸۹ قرار داده است. پاسخ‌های نخست مدل پایه هرچند از نظر موضوعی به برخی واژه‌های پرسش نزدیک‌اند، عمدتاً درباره موی بلند، ورزش‌های مجاز، احکام عمومی زنان یا مباحث کلی اخلاقی هستند. در نتیجه، مدل پایه توانسته موضوع‌های پراکنده پرسش را تا حدی تشخیص دهد، اما نتوانسته ارتباط میان مجموعه قیدهای پرسش و حکم مورد نیاز را به پاسخ دقیق متصل کند.",
        "پس از ریزتنظیم، پاسخ مرجع به رتبه اول منتقل شده است. این تغییر نشان می‌دهد مدل ریزتنظیم‌شده توانسته پرسش محاوره‌ای و نسبتاً طولانی را به یک نیاز اطلاعاتی مشخص تبدیل کند: حکم انجام چند فعالیت در شرایطی که حدود شرعی رعایت شود. پاسخ درست نیز دقیقاً همین ساختار شرطی را دارد و به‌جای تمرکز جداگانه بر مو، ورزش یا اسب‌سواری، معیارهای شرعی مشترک آن‌ها را بیان می‌کند.",
        "این نمونه از نظر کاربردی اهمیت دارد، زیرا پرسش‌های کاربران در سامانه‌های واقعی معمولاً کوتاه، رسمی و تک‌موضوعی نیستند. پرسش حاضر دارای نگارش محاوره‌ای، چند فعالیت متفاوت و نگرانی درباره یک حکم روایی است. بهبود رتبه پاسخ مرجع نشان می‌دهد ریزتنظیم با داده دامنه‌ای به مدل کمک کرده است تا فراتر از هم‌پوشانی واژگانی، هدف عملی پرسش و شروط پاسخ مناسب را بهتر بازنمایی کند.",
    ):
        add_text(document, paragraph_text)

    add_text(document, "۴-۳-۲. تمایز میان پرسش اصطلاحی و متن‌های کلی درباره احکام", size=14, bold=True)
    for paragraph_text in (
        "پرسش با شناسه porseman-8887 این است: فرق بین جایز نبودن و حرام چیست؟ پاسخ مرجع کوتاه و روشن است: در مقام عمل تفاوتی بین آن دو نیست.",
        "مدل پایه پاسخ مرجع را در رتبه ۲۵۳ قرار داده است و پنج پاسخ نخست آن شامل متن‌هایی درباره حکمت و علت احکام، تفاوت گناه و اشتباه، عذاب اخروی و مباحث کلی احکام است. این پاسخ‌ها واژه‌هایی مانند حرام، گناه، حکم و احکام الهی دارند و از نظر موضوعی دور نیستند؛ اما پرسش کاربر یک پرسش اصطلاحی و مقایسه‌ای است و به تعریف عملی تفاوت میان دو تعبیر نیاز دارد.",
        "مدل ریزتنظیم‌شده پاسخ مرجع را در رتبه اول قرار داده است. این مدل توانسته تشخیص دهد که کاربر به دنبال بحث فلسفی درباره علت احکام یا مفهوم گناه نیست، بلکه تفاوت کاربردی دو اصطلاح فقهی را می‌پرسد. به‌عبارت دیگر، مدل جدید میان شباهت موضوعی و انطباق با نیاز اطلاعاتی دقیق پرسش تمایز بهتری ایجاد کرده است.",
        "این نمونه از منظر کاربردی نیز مهم است. در سامانه‌های پاسخ‌گویی و RAG، پرسش‌های کوتاه اصطلاحی فراوان‌اند و معمولاً چندین متن عمومی با واژه‌های مشترک وجود دارد. اگر مدل صرفاً بر شباهت واژگانی تکیه کند، ممکن است متن‌های طولانی و کلی را پیش از پاسخ کوتاه و دقیق قرار دهد. انتقال پاسخ مرجع از رتبه ۲۵۳ به رتبه اول نشان می‌دهد ریزتنظیم مدل BGE-M3 با داده پرسمان، بازیابی پاسخ‌های فشرده و دقیق را بهبود داده است.",
    ):
        add_text(document, paragraph_text)

    add_text(document, "۴-۳-۳. جمع‌بندی تحلیل کیفی", size=14, bold=True)
    add_text(document, "هر دو نمونه نشان می‌دهند بهبود مدل ریزتنظیم‌شده صرفاً ناشی از افزایش شباهت میان پرسش و پاسخ نیست. در نمونه نخست، مدل نیاز اطلاعاتی یک پرسش محاوره‌ای و چندقیدی را به پاسخ شرطی مناسب متصل کرده است؛ در نمونه دوم، مدل نوع نیاز اطلاعاتی پرسش، یعنی تمایز اصطلاحی و عملی، را از متن‌های کلی درباره احکام جدا کرده است. بنابراین، نتایج کیفی با نتایج کمی فصل چهارم هم‌راستا هستند و نشان می‌دهند ریزتنظیم دامنه‌ای، دقت رتبه‌بندی پاسخ‌های فارسی را در پرسش‌های نزدیک و ابهام‌پذیر افزایش داده است.")

    add_text(document, "۴-۴. اعتبار نتایج و محدودیت‌های ارزیابی", size=15, bold=True)
    add_text(document, "نتایج کمی و تحلیل‌های کیفی فصل حاضر نشان می‌دهند که ریزتنظیم BGE-M3 با داده پرسمان، در بازیابی پاسخ‌های فارسی مؤثر بوده است. بااین‌حال، تفسیر علمی نتایج نیازمند توجه به محدودیت‌های داده، تعریف پاسخ صحیح و روش تفکیک آموزش و آزمون است. این محدودیت‌ها لزوماً نتیجه اصلی پژوهش را رد نمی‌کنند، اما محدوده اعتبار آن را روشن می‌سازند.")
    evaluation_limitations = (
        ("۴-۴-۱. خطر هم‌پوشانی میان آموزش و آزمون", "مهم‌ترین تهدید برای اعتبار ارزیابی، وجود هم‌پوشانی میان داده آموزش و آزمون است. این هم‌پوشانی می‌تواند به‌صورت پرسش یکسان، پاسخ یکسان، بازنویسی نزدیک یک پرسش یا زوج پرسش‌وپاسخ با محتوای تقریباً یکسان رخ دهد. در چنین شرایطی، مدل ممکن است به‌جای تعمیم به پرسش جدید، از الگوی مشاهده‌شده در آموزش استفاده کند. در نتیجه، معیارهای بازیابی بیش از توان واقعی مدل در داده‌های دیده‌نشده گزارش می‌شوند. بررسی موردی نمونه‌های کیفی نشان داد که تطابق دقیق پرسش و پاسخ با فایل واقعی آموزش، به‌تنهایی برای تشخیص همه موارد کافی نیست. ممکن است دو پرسش تنها در چند واژه یا نحوه بیان تفاوت داشته باشند، اما نیاز اطلاعاتی و پاسخ آن‌ها عملاً یکسان باشد. بنابراین، حذف تکرار دقیق یک گام ضروری است، ولی برای جلوگیری کامل از نشتی داده کفایت نمی‌کند. برای نسخه نهایی ارزیابی، باید پرسش‌های آزمون نه‌تنها از نظر تطابق دقیق، بلکه از نظر شباهت بالا با پرسش‌ها و پاسخ‌های آموزشی نیز کنترل شوند. پرسش‌هایی که دارای همزاد بسیار نزدیک در داده آموزش هستند، باید از داده آزمون حذف شوند یا به یک بخش واحد منتقل گردند. سپس تمام مدل‌ها با همان مجموعه آزمون پالایش‌شده دوباره ارزیابی شوند. این روش، مقایسه مدل پایه، مدل ریزتنظیم‌شده و مدل‌های مقایسه‌ای را معتبرتر می‌کند."),
        ("۴-۴-۲. محدودیت تک‌پاسخ مرجع", "در ارزیابی فعلی، برای هر پرسش یک پاسخ مرجع در corpus وجود دارد و معیارها بررسی می‌کنند که آیا همان پاسخ دقیق بازیابی شده است یا خیر. این روش در benchmarkهای retrieval رایج است، اما در داده واقعی پرسش‌وپاسخ محدودیت دارد. ممکن است برای یک پرسش، چند پاسخ صحیح، مکمل یا قابل‌قبول وجود داشته باشد. برای نمونه، در پرسش آیا توکل یعنی تلاش نکردن؟، مدل ریزتنظیم‌شده یک پاسخ مرتبط درباره مفهوم توکل را در رتبه اول بازیابی کرده، اما پاسخ مرجع ثبت‌شده در رتبه سوم قرار گرفته است. از دید معیار تک‌پاسخ، این مورد بهترین نتیجه محسوب نمی‌شود؛ اما از دید محتوایی، پاسخ رتبه اول نیز می‌تواند برای کاربر سودمند باشد. بنابراین، معیارهای خودکار ممکن است در برخی موارد کیفیت ادراک‌شده بازیابی را کمتر از واقع نشان دهند. یک راه‌حل در پژوهش‌های آینده، تعریف چند پاسخ مرتبط برای هر پرسش است. این پاسخ‌ها می‌توانند با داوری انسانی یا بازبینی دقیق مدل زبانی بزرگ برچسب‌گذاری شوند. در این صورت، معیارهای retrieval به‌جای وابستگی به یک شناسه پاسخ، مجموعه‌ای از پاسخ‌های قابل‌قبول را در نظر می‌گیرند."),
        ("۴-۴-۳. محدودیت پوشش موضوعی corpus", "مجموعه پرسمان شامل موضوع‌های متنوعی مانند احکام، اعتقادات، اخلاق، قرآن، تاریخ، خانواده و مشاوره است. این تنوع برای آموزش مدل مفید است، اما توزیع موضوع‌ها الزاماً یکنواخت نیست. برخی موضوع‌ها ممکن است پرسش‌ها و پاسخ‌های بیشتری داشته باشند و در نتیجه، سهم بیشتری در معیار کلی دریافت کنند. ممکن است مدل در حوزه‌های پرتکرار، مانند پرسش‌های فقهی کوتاه، عملکرد بهتری داشته باشد و در پرسش‌های بلند، مشاوره‌ای یا دارای زمینه تاریخی عملکرد ضعیف‌تری نشان دهد. معیار کلی به‌تنهایی این تفاوت را آشکار نمی‌کند. تحلیل آینده می‌تواند معیارها را به‌صورت تفکیک‌شده بر حسب موضوع، طول پرسش، طول پاسخ و نوع نیاز اطلاعاتی گزارش کند."),
        ("۴-۴-۴. محدودیت مدل‌های مقایسه‌ای", "مدل‌های Jina و Snowflake با تنظیمات رسمی و روی داده آزمون یکسان اجرا شدند، اما هر مدل معماری، داده پیش‌آموزش، دستور encoding و طراحی متفاوتی دارد. بنابراین، مقایسه آن‌ها با مدل BGE-M3 ریزتنظیم‌شده، مقایسه عملکرد در یک وظیفه واحد است، نه اثبات برتری مطلق یک معماری یا یک خانواده مدل در همه زبان‌ها و corpusها. همچنین، مدل BGE-M3 ریزتنظیم‌شده به‌طور اختصاصی با داده پرسمان سازگار شده است، در حالی که Jina و Snowflake در حالت عمومی ارزیابی شده‌اند. برتری مدل ریزتنظیم‌شده، بیشتر نشان‌دهنده ارزش سازگارسازی دامنه‌ای است تا برتری ذاتی BGE-M3 در همه کاربردها. برای مقایسه کامل‌تر، می‌توان در آینده مدل‌های مقایسه‌ای را نیز با داده آموزشی یکسان ریزتنظیم کرد."),
        ("۴-۴-۵. محدودیت ارزیابی dense retrieval", "ارزیابی حاضر بر dense retrieval و شباهت کسینوسی میان embedding پرسش و پاسخ متمرکز است. در کاربرد واقعی، یک سامانه بازیابی ممکن است از index تقریبی، فیلترهای متاداده، reranker، بازیابی hybrid یا مدل زبانی مولد استفاده کند. عملکرد مدل embedding در ارزیابی دقیق corpus کوچک، لزوماً برابر با عملکرد کل سامانه نهایی در محیط عملیاتی نیست. بااین‌حال، ارزیابی dense retrieval یک پایه مهم برای سنجش مدل است. اگر پاسخ مرتبط در رتبه‌های بالای بازیابی اولیه قرار نگیرد، مرحله reranking یا تولید پاسخ نیز معمولاً نمی‌تواند آن را جبران کند. بهبود Recall@5 و MRR@10 نشان می‌دهد مدل ریزتنظیم‌شده می‌تواند ورودی مناسب‌تری برای مراحل بعدی یک سامانه RAG فراهم کند."),
        ("۴-۴-۶. محدودیت نبود آزمون معناداری آماری", "تفاوت معیارهای مدل‌ها در این پژوهش به‌صورت درصد گزارش شده است. با وجود آن‌که اختلاف‌ها در معیارهای اصلی قابل توجه‌اند، آزمون رسمی معناداری آماری یا بازه اطمینان برای آن‌ها محاسبه نشده است. در یک ارزیابی کامل‌تر می‌توان از bootstrap روی پرسش‌های آزمون استفاده کرد تا برای اختلاف Recall و MRR بازه اطمینان به‌دست آید. این تحلیل مشخص می‌کند که آیا بهبود مشاهده‌شده تنها حاصل تعدادی محدود از پرسش‌ها است یا در نمونه‌های بازنمونه‌گیری‌شده نیز پایدار باقی می‌ماند. اجرای چنین تحلیلی، قدرت استدلال آماری پایان‌نامه را افزایش خواهد داد."),
        ("۴-۴-۷. جمع‌بندی", "نتایج فعلی شواهد مناسبی از اثر ریزتنظیم دامنه‌ای بر بازیابی پاسخ‌های فارسی پرسمان فراهم می‌کنند؛ بااین‌حال، باید با توجه به خطر هم‌پوشانی معنایی آموزش و آزمون، محدودیت تک‌پاسخ مرجع، توزیع موضوعی corpus و نبود آزمون معناداری آماری تفسیر شوند. مهم‌ترین اقدام پیش از تثبیت نتایج نهایی، اجرای audit کامل نشتی داده و ارزیابی دوباره همه مدل‌ها بر مجموعه آزمون پالایش‌شده است."),
    )
    for title, description in evaluation_limitations:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۴-۵. دلالت نتایج برای سامانه‌های بازیابی و RAG", size=15, bold=True)
    add_text(document, "نتایج این پژوهش فقط به مسئله رتبه‌بندی پاسخ در یک benchmark محدود نمی‌شوند. مدل embedding در یک سامانه واقعی بازیابی یا RAG، مرحله نخست انتخاب دانش است. اگر این مرحله نتواند پاسخ یا سند مرتبط را در میان نتایج ابتدایی قرار دهد، مرحله‌های بعدی مانند reranker یا مدل زبانی مولد نیز معمولاً به اطلاعات درست دسترسی نخواهند داشت. ازاین‌رو، بهبود معیارهای بازیابی مدل ریزتنظیم‌شده می‌تواند بر کیفیت کل سامانه اثر بگذارد.")
    rag_implication_sections = (
        ("۴-۵-۱. نقش مدل embedding در زنجیره RAG", "یک سامانه RAG معمولاً از چند مرحله تشکیل می‌شود. ابتدا پرسش کاربر دریافت و به embedding تبدیل می‌شود. سپس این بردار در index پاسخ‌ها یا اسناد جست‌وجو می‌شود و تعدادی نامزد اولیه بازیابی می‌گردد. در مرحله بعد، می‌توان نامزدها را با reranker دقیق‌تر مرتب کرد و در نهایت، پاسخ‌های منتخب به مدل زبانی بزرگ داده می‌شوند تا پاسخ نهایی تولید یا خلاصه شود. مدل embedding تعیین می‌کند چه پاسخ‌هایی اصلاً وارد مجموعه نامزدها شوند. بنابراین، این مدل مسئول انتخاب نهایی پاسخ نیست، اما مسئول ایجاد مجموعه اولیه قابل‌اعتماد است. هرچه پاسخ مرتبط در رتبه بالاتری ظاهر شود، احتمال آن‌که مرحله‌های بعدی بتوانند از آن استفاده کنند بیشتر خواهد بود."),
        ("۴-۵-۲. تفسیر Recall@5 در کاربرد عملی", "Recall@5 مدل ریزتنظیم‌شده در ارزیابی فعلی برابر با ۹۴٫۹۴ درصد است. این مقدار نشان می‌دهد که برای بخش بزرگی از پرسش‌های آزمون، پاسخ مرجع در میان پنج نتیجه نخست بازیابی شده است. در معماری RAG، پنج پاسخ نخست می‌توانند به reranker یا مدل زبانی بزرگ ارسال شوند تا پاسخ نهایی بر اساس آن‌ها تولید شود. افزایش Recall@5 نسبت به مدل پایه، از ۹۱٫۹۴ درصد به ۹۴٫۹۴ درصد، به این معنا است که مدل ریزتنظیم‌شده در تعداد بیشتری از پرسش‌ها پاسخ صحیح را وارد مرحله بعدی زنجیره می‌کند. این بهبود حتی زمانی مهم است که پاسخ مرجع دقیقاً در رتبه اول نباشد؛ زیرا reranker می‌تواند از میان پنج نامزد، پاسخ دقیق‌تر را انتخاب کند. بااین‌حال، این نتیجه نباید به‌صورت مستقیم معادل دقت پاسخ تولیدشده توسط RAG تفسیر شود. مدل زبانی مولد ممکن است با وجود دریافت پاسخ درست، اطلاعات را نادرست ترکیب کند، متن نامعتبر تولید کند یا از محتوای بازیابی‌شده استفاده نکند. بنابراین، Recall@5 کیفیت مرحله بازیابی را می‌سنجد، نه کیفیت نهایی کل سامانه RAG."),
        ("۴-۵-۳. اهمیت Recall@1 برای پاسخ مستقیم", "Recall@1 مدل ریزتنظیم‌شده برابر با ۸۱٫۷۰ درصد گزارش شد. این معیار در سناریویی اهمیت دارد که سامانه پاسخ بازیابی‌شده را بدون تولید متن جدید و به‌صورت مستقیم به کاربر نمایش می‌دهد. در چنین کاربردی، پاسخ رتبه اول بیشترین اهمیت را دارد، زیرا معمولاً کاربر فقط همان نتیجه اول را مشاهده می‌کند. افزایش Recall@1 نسبت به مدل پایه نشان می‌دهد ریزتنظیم تنها باعث افزایش پوشش پاسخ‌های مرتبط در میان نتایج ابتدایی نشده است، بلکه مدل را در انتخاب پاسخ مناسب به‌عنوان بهترین نتیجه نیز تقویت کرده است. این ویژگی برای سامانه‌های پرسش‌وپاسخ فارسی، مرکز پاسخ‌گویی، جست‌وجوی دانش و راهنمای هوشمند اهمیت عملی دارد."),
        ("۴-۵-۴. نقش MRR@10 در اولویت‌بندی پاسخ‌ها", "MRR@10 مدل ریزتنظیم‌شده برابر با ۸۷٫۶۷ درصد است. این معیار نشان می‌دهد پاسخ‌های صحیح، به‌طور میانگین، به رتبه‌های بالاتری منتقل شده‌اند. در کاربرد واقعی، تفاوت میان رتبه اول، دوم و پنجم اهمیت دارد؛ زیرا کاربران و مرحله‌های بعدی سامانه معمولاً توجه بیشتری به نتایج ابتدایی دارند. بهبود MRR@10 نشان می‌دهد مدل تنها پاسخ درست را به ده نتیجه اول وارد نکرده، بلکه توانسته آن را در بسیاری از موارد از رتبه‌های پایین‌تر به رتبه‌های ابتدایی منتقل کند. نمونه‌های کیفی فصل حاضر نیز همین الگو را نشان می‌دهند: مدل ریزتنظیم‌شده در پرسش‌های چندقیدی یا اصطلاحی، پاسخ دقیق را بالاتر از متن‌های کلی و ظاهراً مرتبط قرار داده است."),
    )
    for title, description in rag_implication_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۴-۵-۵. طراحی پیشنهادی برای کاربرد عملی", size=14, bold=True)
    add_text(document, "بر مبنای نتایج پژوهش، یک طراحی عملی برای سامانه پاسخ‌گویی فارسی می‌تواند شامل مراحل زیر باشد:")
    for item_number, item_text in enumerate((
        "پاسخ‌های معتبر پرسمان یا corpus هدف، پیش‌پردازش و به embedding تبدیل شوند.",
        "embedding پاسخ‌ها در یک index برداری، مانند FAISS یا یک پایگاه داده برداری، ذخیره شود.",
        "پرسش کاربر با مدل BGE-M3 ریزتنظیم‌شده به embedding تبدیل گردد.",
        "پنج تا ده پاسخ نزدیک‌تر از index بازیابی شوند.",
        "در صورت نیاز، پاسخ‌ها با reranker بازرتبه‌بندی شوند.",
        "پاسخ رتبه اول مستقیماً نمایش داده شود یا چند پاسخ برتر به مدل زبانی بزرگ داده شوند.",
        "مدل زبانی بزرگ پاسخ نهایی را با اتکا به منابع بازیابی‌شده تولید کند و شناسه یا منبع پاسخ را نیز نمایش دهد.",
    ), 1):
        add_text(document, f"{item_number}. {item_text}")
    add_text(document, "این طراحی باعث می‌شود مدل زبانی بزرگ به‌جای تکیه صرف بر دانش پارامتری خود، از پاسخ‌های بازیابی‌شده فارسی استفاده کند. همچنین، نمایش منبع بازیابی‌شده می‌تواند قابلیت بررسی پاسخ را برای کاربر افزایش دهد.")

    rag_final_sections = (
        ("۴-۵-۶. ملاحظات اعتماد و ایمنی", "در حوزه‌های دینی، فقهی، اخلاقی و مشاوره‌ای، بازیابی پاسخ مرتبط به‌تنهایی برای ارائه پاسخ نهایی کافی نیست. پاسخ بازیابی‌شده باید منبع، تاریخ، اعتبار و تناسب آن با شرایط پرسش کاربر را نیز در نظر بگیرد. برای نمونه، یک حکم فقهی ممکن است به مرجع تقلید، شرایط فرد یا جزئیات مسئله وابسته باشد. در یک سامانه عملی، بهتر است پاسخ‌ها همراه با منبع ثبت‌شده، نوع موضوع و در صورت امکان تاریخ یا مرجع ارائه شوند. همچنین، سامانه باید در پرسش‌های مبهم، حساس یا دارای اطلاعات ناکافی، به‌جای ارائه پاسخ قطعی، کاربر را به بیان جزئیات بیشتر یا مراجعه به منبع تخصصی راهنمایی کند. مدل embedding ریزتنظیم‌شده در این پژوهش می‌تواند کیفیت بازیابی را افزایش دهد، اما جایگزین داوری تخصصی یا اعتبارسنجی محتوای پاسخ‌ها نیست. این تمایز برای تفسیر صحیح نتایج و طراحی مسئولانه سامانه‌های RAG ضروری است."),
        ("۴-۵-۷. جمع‌بندی", "بهبود معیارهای بازیابی نشان می‌دهد مدل BGE-M3 ریزتنظیم‌شده ظرفیت مناسبی برای مرحله retrieval در سامانه‌های فارسی دارد. به‌ویژه، مقدار بالای Recall@5 بیانگر آن است که پاسخ مرجع در اغلب موارد به مجموعه نامزدهای اولیه وارد می‌شود. این ویژگی، زمینه مناسبی برای استفاده از reranker و مدل زبانی بزرگ در یک سامانه RAG فراهم می‌کند. بااین‌حال، نتایج retrieval نباید به‌عنوان سنجش کامل کیفیت RAG تلقی شوند. برای ارزیابی نهایی یک سامانه RAG، باید کیفیت پاسخ تولیدشده، استنادپذیری، صحت محتوایی، رضایت کاربر و رفتار سامانه در پرسش‌های مبهم نیز به‌طور مستقل بررسی شود."),
    )
    for title, description in rag_final_sections:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۴-۶. جمع‌بندی فصل چهارم", size=15, bold=True)
    for paragraph_text in (
        "در این فصل، عملکرد مدل پایه BGE-M3، مدل BGE-M3 ریزتنظیم‌شده، Jina Embeddings v3 و Snowflake Arctic Embed L v2.0 روی مجموعه آزمون ثابت پرسمان بررسی شد. همه مدل‌ها با corpus پاسخ یکسان، پرسش‌های آزمون یکسان و معیارهای مشترک ارزیابی شدند تا مقایسه‌ای منصفانه فراهم شود.",
        "نتایج کمی نشان داد مدل BGE-M3 ریزتنظیم‌شده در هر سه معیار اصلی بهترین عملکرد را به‌دست آورده است. مقدار Recall@1 مدل از ۷۴٫۳۳ درصد در مدل پایه به ۸۱٫۷۰ درصد رسید. مقدار Recall@5 از ۹۱٫۹۴ درصد به ۹۴٫۹۴ درصد افزایش یافت و MRR@10 نیز از ۸۱٫۹۷ درصد به ۸۷٫۶۷ درصد رسید. این نتایج نشان می‌دهد ریزتنظیم دامنه‌ای باعث شده پاسخ صحیح هم بیشتر در میان نتایج ابتدایی حضور داشته باشد و هم در رتبه بالاتری قرار گیرد.",
        "مقایسه با مدل‌های Jina و Snowflake نیز نشان داد مدل ریزتنظیم‌شده در محیط داده پرسمان عملکرد بهتری دارد. این برتری به معنای برتری مطلق BGE-M3 در همه corpusها و زبان‌ها نیست؛ بلکه بیانگر اثر سازگارسازی یک مدل embedding چندزبانه با داده فارسی، ساختار پرسش‌وپاسخ و واژگان تخصصی دامنه هدف است.",
        "تحلیل کیفی دو نمونه نشان داد بهبود مدل فقط به افزایش شباهت واژگانی محدود نیست. در یک پرسش محاوره‌ای و چندقیدی درباره پوشش و ورزش، مدل ریزتنظیم‌شده توانست حکم شرطی و مرتبط را در رتبه اول بازیابی کند. در پرسشی دیگر درباره تفاوت دو اصطلاح فقهی، مدل جدید پاسخ کوتاه و دقیق را از میان متن‌های کلی درباره احکام جدا کرد. این نمونه‌ها با نتایج کمی هم‌راستا هستند و نشان می‌دهند مدل در تشخیص نیاز اطلاعاتی دقیق پرسش بهتر عمل کرده است.",
        "در این فصل همچنین محدودیت‌های ارزیابی بررسی شد. مهم‌ترین محدودیت، احتمال هم‌پوشانی معنایی میان داده آموزش و آزمون است که می‌تواند معیارهای بازیابی را افزایش دهد. محدودیت تک‌پاسخ مرجع، نبود آزمون معناداری آماری و تفاوت میان ارزیابی dense retrieval با عملکرد کامل یک سامانه RAG نیز باید در تفسیر نتایج در نظر گرفته شوند.",
        "به‌طور کلی، نتایج فصل چهارم از فرض اصلی پژوهش پشتیبانی می‌کنند: پالایش داده فارسی، ساخت داده contrastive و ریزتنظیم BGE-M3 با نمونه‌های منفی سخت می‌تواند بازیابی پاسخ‌های فارسی را در corpus پرسمان بهبود دهد. بااین‌حال، پیش از تثبیت نسخه نهایی نتایج، audit کامل نشتی داده و ارزیابی دوباره بر مجموعه آزمون پالایش‌شده ضروری است.",
    ):
        add_text(document, paragraph_text)

    document.add_page_break()
    add_text(document, "فصل پنجم: نتیجه‌گیری و پیشنهادها", size=18, bold=True)
    add_text(document, "۵-۱. نتیجه‌گیری پژوهش", size=15, bold=True)
    for paragraph_text in (
        "هدف این پژوهش، بهبود بازیابی معنایی پاسخ‌های فارسی در مجموعه پرسش‌وپاسخ پرسمان با استفاده از ریزتنظیم مدل embedding چندزبانه BGE-M3 بود. مسئله اصلی این بود که مدل‌های عمومی embedding، با وجود برخورداری از دانش چندزبانه، لزوماً با واژگان، نگارش، سبک پرسش‌ها و ساختار پاسخ‌های یک corpus فارسی دامنه‌ای سازگار نیستند. در نتیجه، ممکن است پاسخ‌های کلی یا دارای هم‌پوشانی واژگانی را پیش از پاسخ دقیق و مرتبط بازیابی کنند.",
        "برای پاسخ به این مسئله، ابتدا داده خام پرسمان پالایش شد. متن‌های غیرپرسشی، پرسش‌های مقاله‌مانند، پاسخ‌های نامتعارف، پاسخ‌های بیش از ۱۰۲۴ token، رکوردهای تکراری و موارد دارای خطای ساختاری کنترل شدند. همچنین، شناسه‌های پایدار برای رکوردها ایجاد شد تا فرایند ساخت داده، تفکیک آموزش و آزمون و ارزیابی مدل قابل ردیابی باشد.",
        "در مرحله بعد، داده آموزش در قالب contrastive آماده شد. برای هر پرسش، یک پاسخ مثبت و چند پاسخ منفی در نظر گرفته شد. پاسخ‌های منفی به‌صورت کاملاً تصادفی انتخاب نشدند، بلکه از میان پاسخ‌های نزدیک بازیابی‌شده و بازرتبه‌بندی‌شده انتخاب شدند. هدف از این تصمیم، آموزش مدل برای تشخیص پاسخ دقیق در میان پاسخ‌های ظاهراً مرتبط بود. این مسئله در داده پرسمان اهمیت زیادی دارد، زیرا بسیاری از پاسخ‌ها از نظر موضوعی به هم نزدیک‌اند، اما فقط برخی از آن‌ها نیاز اطلاعاتی مشخص پرسش را برآورده می‌کنند.",
        "مدل پایه BAAI/bge-m3 با داده آماده‌شده ریزتنظیم شد. نتایج ارزیابی اولیه نشان داد مدل ریزتنظیم‌شده نسبت به مدل پایه و دو مدل embedding عمومی مقایسه‌ای، یعنی Jina Embeddings v3 و Snowflake Arctic Embed L v2.0، در معیارهای Recall@1، Recall@5 و MRR@10 عملکرد بهتری دارد. این نتیجه از فرض اصلی پژوهش پشتیبانی می‌کند: سازگارسازی دامنه‌ای مدل embedding می‌تواند بازیابی پاسخ‌های فارسی را بهبود دهد.",
        "تحلیل کیفی نیز نشان داد مدل ریزتنظیم‌شده در پرسش‌های محاوره‌ای، چندقیدی و اصطلاحی عملکرد بهتری دارد. در این موارد، مدل پایه معمولاً متن‌هایی را بازیابی می‌کرد که از نظر موضوعی نزدیک، اما از نظر پاسخ‌گویی دقیق ناکافی بودند. مدل ریزتنظیم‌شده در موارد بیشتری توانست پاسخ مرتبط را در رتبه‌های بالاتر قرار دهد. این مشاهده نشان می‌دهد بهبود مدل فقط به افزایش شباهت سطحی میان واژه‌ها محدود نیست، بلکه به تشخیص دقیق‌تر نیاز اطلاعاتی پرسش مربوط است.",
        "باوجود این، نتایج پژوهش باید با احتیاط تفسیر شوند. بررسی نمونه‌های کیفی نشان داد که احتمال وجود هم‌پوشانی معنایی میان آموزش و آزمون وجود دارد؛ حتی اگر پرسش یا پاسخ از نظر رشته‌ای دقیقاً یکسان نباشد. چنین هم‌پوشانی‌ای می‌تواند معیارهای ارزیابی را افزایش دهد. بنابراین، مهم‌ترین اقدام پیش از نهایی‌کردن ارقام گزارش‌شده، اجرای audit کامل نشتی داده و ارزیابی دوباره همه مدل‌ها بر مجموعه آزمون پالایش‌شده است.",
        "در مجموع، دستاورد اصلی پژوهش ارائه یک خط لوله قابل بازتولید برای داده فارسی پرسش‌وپاسخ است: پالایش داده، کنترل طول با توکنایزر، حذف تکرارها، ساخت زوج‌های contrastive، انتخاب نمونه‌های منفی سخت، ریزتنظیم مدل BGE-M3 و ارزیابی یکسان چند مدل. این خط لوله می‌تواند برای corpusهای فارسی دیگر نیز استفاده یا با تغییرات محدود سازگار شود.",
        "مدل ریزتنظیم‌شده حاصل می‌تواند به‌عنوان مؤلفه بازیابی در سامانه‌های پرسش‌وپاسخ فارسی و معماری‌های RAG استفاده شود. بااین‌حال، استفاده عملی از آن باید همراه با کنترل کیفیت داده، نمایش منبع پاسخ، مدیریت پرسش‌های مبهم و اعتبارسنجی تخصصی محتوا باشد. مدل embedding کیفیت انتخاب دانش را بهبود می‌دهد، اما جایگزین اعتبارسنجی محتوایی یا داوری تخصصی نیست.",
    ):
        add_text(document, paragraph_text)

    add_text(document, "۵-۲. دستاوردها و نوآوری‌های پژوهش", size=15, bold=True)
    add_text(document, "دستاوردهای این پژوهش فقط به تولید یک مدل ریزتنظیم‌شده محدود نیست. ارزش اصلی کار در طراحی یک مسیر منسجم از داده خام فارسی تا ارزیابی قابل بازتولید مدل embedding است. مهم‌ترین دستاوردها و نوآوری‌ها به شرح زیر هستند.")
    contributions = (
        ("۵-۲-۱. ساخت فرایند پالایش داده برای corpus فارسی واقعی", "داده پرسمان یک corpus واقعی پرسش‌وپاسخ است و از ابتدا برای آموزش مدل بازیابی طراحی نشده بود. به همین دلیل، پیش از استفاده در آموزش و آزمون، وجود متن‌های غیرپرسشی، رکوردهای مقاله‌مانند، تفاوت‌های نگارشی فارسی، پاسخ‌های خالی، پاسخ‌های بسیار بلند و تکرارهای پرسش و پاسخ بررسی شد. نوآوری این مرحله در استفاده از قواعد مرحله‌ای و قابل بازتولید است. به‌جای حذف دستی نمونه‌ها، معیارهای مشخصی برای نرمال‌سازی، تشخیص ساختار پرسشی، شناسایی متن‌های مقاله‌مانند، کنترل طول پاسخ با توکنایزر BGE-M3 و حذف تکرارها تعریف شد. این فرایند می‌تواند برای مجموعه‌های فارسی پرسش‌وپاسخ دیگر نیز استفاده شود."),
        ("۵-۲-۲. کنترل طول پاسخ بر مبنای token مدل", "طول پاسخ‌ها در این پژوهش با شمارش tokenهای توکنایزر BAAI/bge-m3 سنجیده شد، نه صرفاً تعداد کاراکتر یا تعداد واژه. پاسخ‌های دارای بیش از ۱۰۲۴ token از داده نهایی کنار گذاشته شدند. این تصمیم باعث شد طول ورودی‌ها با شرایط واقعی پردازش مدل هماهنگ باشد. استفاده از token به‌جای کاراکتر اهمیت دارد، زیرا یک متن فارسی با تعداد کاراکتر مشابه ممکن است، بسته به ترکیب واژه‌ها و نویسه‌ها، تعداد token متفاوتی داشته باشد. بنابراین، این روش معیار دقیق‌تری برای کنترل هزینه آموزش و جلوگیری از ورود نمونه‌های نامتعارف فراهم می‌کند."),
        ("۵-۲-۳. ساخت داده contrastive با نمونه‌های منفی سخت", "برای هر پرسش آموزشی، یک پاسخ مثبت و هفت پاسخ منفی در نظر گرفته شد. پاسخ‌های منفی با انتخاب تصادفی ساخته نشدند؛ ابتدا ۱۰۰ نامزد نزدیک با مدل BGE-M3 بازیابی شدند، سپس با reranker بازرتبه‌بندی شدند و نامزدهای برتر برای انتخاب منفی بررسی شدند. این روش باعث می‌شود مدل با پاسخ‌هایی آموزش ببیند که از نظر واژگانی یا موضوعی به پرسش نزدیک‌اند، اما پاسخ دقیق آن نیستند. در نتیجه، مدل به‌جای تشخیص ساده تفاوت میان موضوع‌های کاملاً نامرتبط، توانایی تمایز میان پاسخ درست و پاسخ ظاهراً مرتبط را یاد می‌گیرد."),
        ("۵-۲-۴. توجه به خطر منفی کاذب", "در داده‌های پرسش‌وپاسخ فارسی، یک پرسش ممکن است بیش از یک پاسخ قابل‌قبول داشته باشد. بنابراین، هر پاسخ نزدیک نباید به‌صورت خودکار به‌عنوان منفی وارد آموزش شود. در طراحی فرایند ساخت داده، امکان کنترل کیفیت نامزدهای منفی و استفاده از مدل زبانی بزرگ برای تشخیص منفی کاذب در نظر گرفته شد. این توجه، به‌ویژه در حوزه‌های دینی، اخلاقی و مشاوره‌ای مهم است؛ زیرا پاسخ‌های مکمل یا پاسخ‌های دارای تفاوت جزئی ممکن است هر دو برای کاربر مفید باشند. ورود چنین پاسخ‌هایی به‌عنوان منفی می‌تواند به مدل سیگنال آموزشی نادرست دهد."),
        ("۵-۲-۵. مقایسه با چند مدل embedding عمومی", "مدل ریزتنظیم‌شده فقط با مدل پایه خود مقایسه نشد. مدل‌های jinaai/jina-embeddings-v3 و Snowflake/snowflake-arctic-embed-l-v2.0 نیز روی مجموعه آزمون یکسان اجرا شدند. این مقایسه، ارزش ریزتنظیم دامنه‌ای را بهتر نشان می‌دهد؛ زیرا مدل نهایی فقط نسبت به نسخه اولیه BGE-M3 ارزیابی نشده است. نتایج اولیه نشان داد مدل ریزتنظیم‌شده در معیارهای اصلی از هر دو مدل عمومی مقایسه‌ای بهتر عمل کرده است. بااین‌حال، این نتیجه باید پس از audit کامل نشتی داده تثبیت شود."),
        ("۵-۲-۶. ایجاد چارچوب ارزیابی قابل بازتولید", "برای هر مدل، رتبه پاسخ صحیح برای تک‌تک پرسش‌ها ذخیره شد. افزون بر معیارهای کلی، فایل‌های رتبه، گزارش HTML و فایل تحلیل کیفی تولید شدند. این خروجی‌ها امکان بررسی پرسش‌های بهبود‌یافته، پرسش‌های دارای افت، پاسخ‌های بازیابی‌شده و تحلیل خطا را فراهم می‌کنند. این چارچوب، امکان مقایسه مدل‌های جدید یا نسخه‌های ریزتنظیم‌شده بعدی را روی همان داده و تنظیمات فراهم می‌سازد. همچنین، در صورت تغییر داده آزمون پس از اجرای audit نشتی، می‌توان تمام مدل‌ها را با خط لوله ثابت دوباره ارزیابی کرد."),
        ("۵-۲-۷. پیوند مستقیم با کاربردهای RAG فارسی", "خروجی پژوهش فقط یک مدل embedding نیست؛ بلکه یک مؤلفه قابل استفاده در مرحله بازیابی سامانه‌های RAG فارسی است. مدل ریزتنظیم‌شده می‌تواند پرسش کاربر را به embedding تبدیل کند و پاسخ‌ها یا اسناد نزدیک را از corpus بازیابی کند. سپس، این پاسخ‌ها می‌توانند به reranker یا مدل زبانی بزرگ داده شوند. تمرکز بر بازیابی پاسخ‌های فارسی واقعی، کنترل کیفیت داده و بررسی محدودیت‌های پاسخ‌های حساس، باعث می‌شود نتایج پژوهش برای طراحی سامانه‌های پرسش‌وپاسخ فارسی مسئولانه و قابل‌استناد باشد."),
    )
    for title, description in contributions:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۵-۳. پیشنهادهایی برای پژوهش‌های آینده", size=15, bold=True)
    add_text(document, "بر مبنای فرایند اجرا، نتایج اولیه و محدودیت‌های شناسایی‌شده، مسیرهای زیر برای ادامه پژوهش پیشنهاد می‌شوند.")
    future_work = (
        ("۵-۳-۱. اجرای audit کامل نشتی داده و ساخت تفکیک گروهی", "مهم‌ترین پیشنهاد، بازبینی کامل استقلال داده‌های آموزش و آزمون است. این بازبینی باید علاوه بر تطابق دقیق متن، پرسش‌های بازنویسی‌شده، پاسخ‌های بسیار مشابه و زوج‌های پرسش‌وپاسخ هم‌معنا را نیز شناسایی کند. برای این کار می‌توان ابتدا با مدل embedding، نزدیک‌ترین پرسش‌های آموزشی را برای هر پرسش آزمون بازیابی کرد و سپس موارد دارای شباهت بالا را با قواعد خودکار، reranker یا داوری انسانی بررسی نمود. پس از شناسایی گروه‌های مشابه، تمام اعضای هر گروه باید فقط در یکی از بخش‌های آموزش یا آزمون قرار گیرند. این روش که تفکیک گروهی نام دارد، ارزیابی سخت‌گیرانه‌تر و واقع‌بینانه‌تری از توان تعمیم مدل ارائه می‌دهد. سپس، مدل پایه، مدل ریزتنظیم‌شده و مدل‌های مقایسه‌ای باید روی مجموعه آزمون پالایش‌شده دوباره ارزیابی شوند."),
        ("۵-۳-۲. تعریف چند پاسخ مرتبط برای هر پرسش", "ارزیابی فعلی برای هر پرسش فقط یک پاسخ مرجع در نظر می‌گیرد. در عمل، به‌ویژه در پرسش‌های اعتقادی، اخلاقی و مشاوره‌ای، ممکن است چند پاسخ مکمل یا قابل‌قبول وجود داشته باشد. تعریف یک مجموعه پاسخ مرتبط برای هر پرسش می‌تواند ارزیابی را به تجربه واقعی کاربر نزدیک‌تر کند. برای ساخت این برچسب‌ها، می‌توان پاسخ‌های برتر بازیابی‌شده توسط چند مدل را گردآوری کرد و سپس آن‌ها را با داوری انسانی یا مدل زبانی بزرگ بررسی نمود. در این روش، پاسخ‌های مرتبط تنها به یک شناسه محدود نمی‌شوند و معیارهای Recall@k و MRR@k بر مبنای مجموعه‌ای از پاسخ‌های قابل‌قبول محاسبه خواهند شد."),
        ("۵-۳-۳. ارزیابی انسانی و آزمون معناداری آماری", "ارزیابی خودکار برای مقایسه سریع مدل‌ها ضروری است، اما همه جنبه‌های کیفیت پاسخ را پوشش نمی‌دهد. پیشنهاد می‌شود نمونه‌ای تصادفی و نمونه‌ای از پرسش‌های دشوار به ارزیابان انسانی ارائه شود تا ارتباط پاسخ، کفایت، دقت، مفیدبودن و سازگاری با پرسش را امتیازدهی کنند. همچنین، می‌توان با bootstrap روی مجموعه پرسش‌های آزمون، بازه اطمینان برای اختلاف معیارهای مدل‌ها را محاسبه کرد. این کار روشن می‌سازد که بهبود مشاهده‌شده در Recall@1، Recall@5 و MRR@10 در نمونه‌های مختلف آزمون پایدار است یا خیر."),
        ("۵-۳-۴. ریزتنظیم مدل‌های مقایسه‌ای با داده یکسان", "در این پژوهش، Jina Embeddings v3 و Snowflake Arctic Embed L v2.0 در حالت عمومی با مدل BGE-M3 ریزتنظیم‌شده مقایسه شدند. برای مقایسه دقیق‌تر معماری‌ها، می‌توان همین داده آموزش، پاسخ‌های مثبت، نمونه‌های منفی سخت و تنظیمات ارزیابی را برای ریزتنظیم مدل‌های مقایسه‌ای نیز به کار برد. این آزمایش مشخص می‌کند که کدام بخش از بهبود حاصل، ناشی از کیفیت داده و ریزتنظیم دامنه‌ای است و کدام بخش به انتخاب مدل پایه مربوط می‌شود. چنین مقایسه‌ای می‌تواند نتیجه قوی‌تری درباره انتخاب مدل مناسب برای بازیابی فارسی ارائه دهد."),
        ("۵-۳-۵. بهبود انتخاب نمونه‌های منفی", "روش انتخاب نمونه‌های منفی سخت می‌تواند توسعه یابد. در یک مسیر، تعداد نامزدهای بازیابی‌شده افزایش می‌یابد و نمونه‌های منفی از موضوع‌های نزدیک‌تر انتخاب می‌شوند. در مسیر دیگر، پاسخ‌های منفی با استفاده از چند مدل embedding متفاوت بازیابی می‌شوند تا سوگیری یک مدل پایه در انتخاب نامزدها کاهش یابد. همچنین، می‌توان از راهبرد curriculum learning استفاده کرد: آموزش ابتدا با نمونه‌های منفی نسبتاً آسان شروع شود و به‌تدریج نمونه‌های نزدیک‌تر و دشوارتر وارد آموزش شوند. این رویکرد ممکن است پایداری آموزش را افزایش دهد و مدل را برای تشخیص تفاوت‌های ظریف آماده‌تر کند."),
        ("۵-۳-۶. توسعه بازیابی hybrid و reranking", "مدل حاضر بر dense retrieval تمرکز دارد. در آینده می‌توان آن را با بازیابی واژه‌محور مانند BM25 ترکیب کرد. بازیابی hybrid در پرسش‌هایی که دارای نام خاص، عبارت دقیق، شماره آیه، شماره مسئله یا اصطلاح خاص هستند مفید خواهد بود؛ زیرا بازیابی dense و واژه‌محور نقاط قوت متفاوتی دارند. همچنین، می‌توان یک reranker فارسی یا چندزبانه را روی نتایج برتر مدل embedding اعمال کرد. در این معماری، مدل BGE-M3 ریزتنظیم‌شده نامزدهای اولیه را با سرعت بازیابی می‌کند و reranker، ارتباط دقیق پرسش و پاسخ را برای تعداد محدودی از نامزدها بررسی می‌نماید."),
        ("۵-۳-۷. پیاده‌سازی و ارزیابی کامل سامانه RAG", "پژوهش حاضر بر مرحله بازیابی تمرکز دارد. گام بعد، ساخت یک سامانه RAG کامل است که مدل embedding ریزتنظیم‌شده را با reranker، مدل زبانی بزرگ، prompt کنترل‌شده و سازوکار نمایش منبع ترکیب کند. ارزیابی RAG باید فراتر از رتبه بازیابی باشد. معیارهایی مانند صحت پاسخ تولیدشده، میزان اتکا به شواهد، پوشش پاسخ، استنادپذیری، رعایت محدودیت‌های محتوایی و رضایت کاربر باید بررسی شوند. در حوزه پرسمان، نمایش منبع، مرجع و محدودیت پاسخ اهمیت ویژه‌ای دارد."),
        ("۵-۳-۸. گسترش corpus و تحلیل موضوعی", "افزودن منابع فارسی معتبر دیگر، مانند متن‌های آموزشی، پاسخ‌های تخصصی و اسناد دارای مرجع، می‌تواند پوشش موضوعی corpus را افزایش دهد. در این حالت، لازم است کیفیت منبع، مجوز استفاده، تاریخ محتوا و ساختار استنادی نیز کنترل شود. تحلیل موضوعی نیز می‌تواند مشخص کند مدل در کدام حوزه‌ها عملکرد بهتری دارد. برای نمونه، می‌توان نتایج را در گروه‌های احکام، قرآن، اعتقادات، تاریخ، اخلاق و مشاوره جداگانه گزارش کرد. این تحلیل به شناسایی ضعف‌های مدل و اولویت‌بندی داده‌های موردنیاز برای ریزتنظیم بعدی کمک می‌کند."),
        ("۵-۳-۹. جمع‌بندی", "پیشنهادهای ارائه‌شده از سه مسیر اصلی پیروی می‌کنند: افزایش اعتبار ارزیابی، بهبود داده و مدل بازیابی، و توسعه کاربرد عملی در RAG. نخستین اولویت، اجرای audit کامل نشتی داده و تثبیت مجموعه آزمون مستقل است. پس از آن، ارزیابی انسانی، چندپاسخ‌کردن معیارها، مقایسه منصفانه‌تر مدل‌ها و توسعه بازیابی hybrid می‌تواند کیفیت و اعتبار پژوهش را افزایش دهد."),
    )
    for title, description in future_work:
        add_text(document, title, size=14, bold=True)
        add_text(document, description)

    add_text(document, "۵-۴. جمع‌بندی نهایی", size=15, bold=True)
    for paragraph_text in (
        "این پژوهش با هدف بهبود بازیابی معنایی پاسخ‌های فارسی در مجموعه پرسمان انجام شد. مسئله اصلی، فاصله میان قابلیت عمومی مدل‌های embedding چندزبانه و نیازهای خاص یک corpus فارسی پرسش‌وپاسخ بود. پاسخ به این مسئله تنها با انتخاب یک مدل پایه مناسب ممکن نیست؛ کیفیت داده، ساخت نمونه‌های آموزشی، انتخاب پاسخ‌های منفی، استقلال داده آزمون و روش ارزیابی نیز نقش تعیین‌کننده دارند.",
        "در این پژوهش، داده خام پرسمان در چند مرحله پالایش شد و رکوردهای غیرپرسشی، مقاله‌مانند، تکراری یا نامتعارف کنترل شدند. پاسخ‌ها با توکنایزر BGE-M3 بررسی شدند تا نمونه‌های بیش از حد بلند وارد داده نهایی نشوند. سپس، داده آموزشی contrastive با پاسخ مثبت و نمونه‌های منفی سخت ساخته شد و مدل BGE-M3 با این داده ریزتنظیم گردید.",
        "نتایج اولیه ارزیابی نشان داد مدل ریزتنظیم‌شده در مقایسه با مدل پایه و دو مدل embedding عمومی، در معیارهای بازیابی عملکرد بهتری دارد. تحلیل کیفی نیز نشان داد مدل جدید در برخی پرسش‌های محاوره‌ای، چندقیدی و اصطلاحی، پاسخ دقیق‌تری را در رتبه‌های بالا بازیابی می‌کند. این یافته‌ها از این دیدگاه پشتیبانی می‌کنند که ریزتنظیم دامنه‌ای می‌تواند برای بازیابی فارسی مؤثر باشد.",
        "بااین‌حال، بررسی کیفی همچنین اهمیت کنترل نشتی داده را آشکار کرد. وجود پرسش‌ها یا پاسخ‌های بسیار نزدیک در بخش‌های آموزش و آزمون می‌تواند ارزیابی را آسان‌تر و مقدار معیارها را بیش از واقع نشان دهد. بنابراین، خروجی مهم دیگر پژوهش، تأکید بر این نکته است که پالایش داده و تفکیک مستقل آموزش و آزمون بخشی از روش پژوهش هستند، نه صرفاً مرحله‌های مقدماتی آماده‌سازی داده.",
        "در مجموع، این پایان‌نامه یک مسیر قابل بازتولید برای توسعه یک بازیاب فارسی ارائه می‌کند: پالایش corpus، ساخت داده contrastive، انتخاب نمونه‌های منفی سخت، ریزتنظیم embedding model، مقایسه چند مدل و تحلیل کیفی خروجی‌ها. این مسیر می‌تواند پایه‌ای برای توسعه سامانه‌های پرسش‌وپاسخ فارسی، جست‌وجوی معنایی و معماری‌های RAG باشد.",
        "استفاده عملی از مدل ریزتنظیم‌شده باید با ارزیابی تکمیلی، اعتبارسنجی محتوایی پاسخ‌ها، نمایش منبع و کنترل پرسش‌های حساس همراه شود. در گام بعد، اجرای audit کامل نشتی داده، تثبیت مجموعه آزمون و ارزیابی مجدد مدل‌ها، نتیجه‌های پژوهش را از نظر علمی مستحکم‌تر خواهد کرد.",
    ):
        add_text(document, paragraph_text)


def add_ltr_reference(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_ltr_run(run, size=12)
    return paragraph


def add_references(document):
    document.add_page_break()
    add_text(document, "منابع", size=18, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    add_text(document, "منابع فارسی", size=15, bold=True)
    add_text(
        document,
        "سامانه پاسخگویی پرسمان. https://www.porseman.com/",
    )

    add_text(document, "منابع لاتین", size=15, bold=True)
    references = (
        "BAAI. (2024). BAAI/bge-m3 [Model card]. Hugging Face. https://huggingface.co/BAAI/bge-m3",
        "Chen, J., Xiao, S., Zhang, P., Luo, K., Lian, D., & Liu, Z. (2024). BGE M3-Embedding: Multi-lingual, multi-functionality, multi-granularity text embeddings through self-knowledge distillation. arXiv. https://arxiv.org/abs/2402.03216",
        "Jina AI. (2024). jinaai/jina-embeddings-v3 [Model card]. Hugging Face. https://huggingface.co/jinaai/jina-embeddings-v3",
        "Karpukhin, V., Oguz, B., Min, S., Lewis, P., Wu, L., Edunov, S., Chen, D., & Yih, W.-T. (2020). Dense passage retrieval for open-domain question answering. In Proceedings of the 2020 Conference on Empirical Methods in Natural Language Processing (EMNLP) (pp. 6769–6781). Association for Computational Linguistics. https://doi.org/10.18653/v1/2020.emnlp-main.550",
        "Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., Goyal, N., Küttler, H., Lewis, M., Yih, W.-T., Rocktäschel, T., Riedel, S., & Kiela, D. (2020). Retrieval-augmented generation for knowledge-intensive NLP tasks. Advances in Neural Information Processing Systems, 33. https://arxiv.org/abs/2005.11401",
        "Manning, C. D., Raghavan, P., & Schütze, H. (2008). Introduction to information retrieval. Cambridge University Press.",
        "Muennighoff, N., Tazi, N., Magne, L., & Reimers, N. (2023). MTEB: Massive text embedding benchmark. In Proceedings of the 17th Conference of the European Chapter of the Association for Computational Linguistics. https://arxiv.org/abs/2210.07316",
        "QomSSLab. (2026, July 14). cheRAGh-Embedding: Persian RAG embedding benchmark [Interactive benchmark]. Hugging Face Spaces. https://huggingface.co/spaces/QomSSLab/cheRAGh-Embedding",
        "Reimers, N., & Gurevych, I. (2019). Sentence-BERT: Sentence embeddings using Siamese BERT-networks. In Proceedings of the 2019 Conference on Empirical Methods in Natural Language Processing and the 9th International Joint Conference on Natural Language Processing (EMNLP-IJCNLP) (pp. 3982–3992). Association for Computational Linguistics. https://doi.org/10.18653/v1/D19-1410",
        "Snowflake. (2024). Snowflake/snowflake-arctic-embed-l-v2.0 [Model card]. Hugging Face. https://huggingface.co/Snowflake/snowflake-arctic-embed-l-v2.0",
        "Sturua, S., Mohr, I., Akram, M. K., Günther, M., Wang, B., Krimmel, M., Wang, F., Mastrapas, G., Koukounas, A., Wang, N., & Xiao, H. (2024). jina-embeddings-v3: Multilingual embeddings with task LoRA. arXiv. https://arxiv.org/abs/2409.10173",
        "Yu, P., Merrick, L., Nuti, G., & Campos, D. (2024). Arctic-Embed 2.0: Multilingual retrieval without compromise. arXiv. https://arxiv.org/abs/2412.04506",
    )
    for reference in references:
        add_ltr_reference(document, reference)


def save_document(document):
    version_numbers = []
    for path in OUTPUT_PATH.parent.glob(f"{OUTPUT_PATH.stem}_v*.docx"):
        suffix = path.stem.removeprefix(f"{OUTPUT_PATH.stem}_v")
        if suffix.isdigit():
            version_numbers.append(int(suffix))

    next_version = max(version_numbers, default=0) + 1
    versioned_path = OUTPUT_PATH.with_stem(f"{OUTPUT_PATH.stem}_v{next_version}")
    document.save(versioned_path)
    return versioned_path


def main():
    if not LOGO_PATH.is_file():
        raise FileNotFoundError(f"Logo was not found: {LOGO_PATH}")

    document = Document()
    settings = document.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    section = document.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(3)

    normal_style = document.styles["Normal"]
    normal_style.font.name = PERSIAN_FONT
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), PERSIAN_FONT)
    normal_style.font.size = Pt(14)
    normal_properties = normal_style._element.get_or_add_rPr()
    normal_cs_size = normal_properties.find(qn("w:szCs"))
    if normal_cs_size is None:
        normal_cs_size = OxmlElement("w:szCs")
        normal_properties.append(normal_cs_size)
    normal_cs_size.set(qn("w:val"), "28")
    for level in (1, 2, 3):
        heading_style = document.styles[f"Heading {level}"]
        heading_style.font.name = PERSIAN_FONT
        heading_style._element.rPr.rFonts.set(qn("w:eastAsia"), PERSIAN_FONT)

    document.core_properties.title = "پیش‌نویس پایان‌نامه پرسمان"
    document.core_properties.author = "علی احمدی"

    add_title_page(document)
    add_persian_abstract(document)
    add_outline(document)
    add_references(document)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    print(save_document(document))


if __name__ == "__main__":
    main()
